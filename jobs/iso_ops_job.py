"""ISO operations job -- detect, propose floor tasks, evaluate evidence. Never certify.

Reads a plant-profile JSON (the POST /api/detectar body) or a one-row CSV/Excel
of the same fields, plus optional order signals (ciclo / scrap / kg / tkm).
Pandas reads the tabular path; JSON is loaded directly. Neither goes through
``jobs/intake.py`` -- a plant ficha is not demand history.

The engine (``src/iso_catalog.py``) already answers detectar / evidencia_de_orden /
claim_permitido. This job only packages that into a GuidedOutcome the orchestrator
can deliver: OPTIONS when at least one norm is applicable, HANDOFF when the ficha
is empty, never EXECUTED. Kern does not say the plant cumple ISO.
"""

from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Mapping

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from jobs.intake import _norm
from nosilo_core.guided import (
    EXECUTED,
    ExecutionOption,
    HandoffPacket,
    Residual,
    as_handoff,
    as_options,
    confidence_from_coverage,
    verify_guided,
)
from src.deliverable import DataSource, Deliverable, Finding, Kpi
from src.export import write_summary_csv
from src.iso_catalog import (
    AVISO_DETECCION,
    AVISO_ORDEN,
    PerfilInvalido,
    PerfilIso,
    SenalesOrden,
    buscar_norma,
    detectar,
    evidencia_de_orden,
    parse_perfil,
    tarjeta_evidencia,
    texto_prohibido,
)
from src.sanitize import defuse_formula

_LIST_FIELDS = ("sectores", "clientePide", "logistica", "planta", "kern")
_LIST_ALIASES: dict[str, tuple[str, ...]] = {
    "sectores": ("sectores", "sector"),
    "clientePide": ("clientepide", "cliente_pide", "cliente"),
    "logistica": ("logistica", "logistics"),
    "planta": ("planta", "plant"),
    "kern": ("kern", "modulos", "moduloskern"),
}
_SIGNAL_ALIASES: dict[str, tuple[str, ...]] = {
    "ciclo": ("ciclo", "cycle", "tiene_ciclo"),
    "scrap": ("scrap", "tiene_scrap"),
    "kg": ("kg", "tiene_kg"),
    "tkm": ("tkm", "t_km", "tiene_tkm"),
}
_TEXTO_ALIASES = ("texto", "brief", "notas", "notes")
_DECK_BASE_CONFIDENCE = 0.6

_CERTIFY_RESIDUAL = Residual(
    description=(
        "Esto detecta aplicabilidad y potencial de evidencia de piso. "
        "No es un dictamen de conformidad ni una certificacion."
    ),
    risk_if_skipped=(
        "Tratar el resultado como un sello de planta expone al cliente "
        "a un claim que Kern no puede sostener: el dictamen lo da un organismo."
    ),
)

_EMPTY_HANDOFF = HandoffPacket(
    title="Cargar la ficha de planta",
    steps=[
        "Completar sectores, lo que pide el cliente, logistica, planta y modulos KERN.",
        "O adjuntar el JSON/CSV de ficha (tests/fixtures/iso/detectar.ejemplo.json es la forma).",
        "Si hay una orden abierta, marcar ciclo, scrap, kg y t·km cuando existan.",
        "Volver a correr: ficha vacia produce cero normas a proposito, no un catalogo inventado.",
    ],
    data={"detectadas": 0},
    risk_if_skipped=(
        "Sin ficha, Kern no detecta normas. Completarla a ojo despues del entregable "
        "hace pasar un vacio por un diagnostico."
    ),
)


@dataclass(frozen=True)
class PreparedIso:
    """What ``run()`` needs: a validated plant profile plus optional order signals."""

    perfil: PerfilIso
    senales: SenalesOrden
    source: str


@dataclass(frozen=True)
class IsoOpsReport:
    aviso: str
    detectadas: tuple[dict[str, Any], ...]
    orden: dict[str, Any] | None
    outcome: Any
    summary: str
    n_detectadas: int
    n_nativas_orden: int
    n_huecos: int
    perfil_vacio: bool
    source: str


def perfil_presente(params: Mapping[str, Any] | None, data_path: str | None) -> bool:
    """True when the caller sent a ficha, a file, order signals, or free text."""
    if data_path:
        return True
    params = params or {}
    if str(params.get("texto") or "").strip():
        return True
    for key in _LIST_FIELDS:
        raw = params.get(key, params.get(_snake(key), []))
        if isinstance(raw, (list, tuple)) and any(str(x).strip() for x in raw):
            return True
        if isinstance(raw, str) and raw.strip():
            return True
    return any(_as_bool(params.get(k)) for k in ("ciclo", "scrap", "kg", "tkm"))


def prepare(data_path: str | None, params: dict | None = None) -> PreparedIso:
    """Load the ficha from params, JSON, or a one-row CSV/Excel. Params win on overlap."""
    params = dict(params or {})
    from_file: dict[str, Any] = {}
    source = "params"
    if data_path:
        path = Path(data_path)
        if not path.exists():
            raise FileNotFoundError(f"{data_path} does not exist")
        suffix = path.suffix.lower()
        if suffix == ".json":
            from_file = _from_json(path)
            source = "json"
        elif suffix in {".csv", ".xlsx", ".xls"}:
            from_file = _from_table(path)
            source = "csv"
        else:
            raise ValueError(
                "la ficha ISO tiene que ser JSON (cuerpo de /api/detectar) "
                "o CSV/Excel de una fila con sectores/cliente_pide/logistica/planta/kern"
            )

    merged = {**from_file, **{k: v for k, v in params.items() if v not in (None, "", [])}}
    try:
        perfil = parse_perfil(_perfil_body(merged))
    except PerfilInvalido as exc:
        raise ValueError(str(exc)) from exc

    kern_mods = perfil.kern or tuple(
        str(m).lower() for m in (from_file.get("kern") or []) if str(m).strip()
    )
    senales = SenalesOrden(
        tiene_ciclo=_as_bool(merged.get("ciclo")),
        tiene_scrap=_as_bool(merged.get("scrap")),
        tiene_kg=_as_bool(merged.get("kg")),
        tiene_tkm=_as_bool(merged.get("tkm")),
        modulos=kern_mods or ("smartpro", "monitor", "energy", "huella"),
    )
    return PreparedIso(perfil=perfil, senales=senales, source=source)


def run(prepared: PreparedIso, params: dict | None = None) -> IsoOpsReport:
    """Detect applicable norms, attach floor tasks, evaluate order evidence. Never certify."""
    del params
    detection = detectar(prepared.perfil)
    raw_rows = list(detection.get("detectadas") or [])
    if prepared.senales.tiene_tkm:
        for row in raw_rows:
            if row.get("id") == "iso-14083":
                row["huecos"] = []
    detectadas = tuple(_enriquecer(row) for row in raw_rows)
    any_signal = (
        prepared.senales.tiene_ciclo
        or prepared.senales.tiene_scrap
        or prepared.senales.tiene_kg
        or prepared.senales.tiene_tkm
    )
    orden = evidencia_de_orden(prepared.senales) if any_signal else None
    n_nativas = len((orden or {}).get("evidencias") or [])
    n_huecos = sum(1 for row in detectadas if row.get("huecos")) + len((orden or {}).get("huecos") or [])

    residual = _CERTIFY_RESIDUAL
    if prepared.perfil.vacio() and not any_signal:
        outcome = as_handoff(
            AVISO_DETECCION + " Ficha vacia: cero normas detectadas.",
            [_EMPTY_HANDOFF],
            confidence=0.4,
            residuals=[residual],
        )
        summary = "Ficha vacia: cero normas detectadas. No es un dictamen de conformidad."
    else:
        options = _options(detectadas, orden)
        if options:
            summary = (
                f"{detection.get('aviso', AVISO_DETECCION)} "
                f"{len(detectadas)} norma(s) aplicables; "
                f"{n_nativas} evidencia(s) nativa(s) de orden; {n_huecos} hueco(s)."
            )
            outcome = as_options(summary, options, confidence=0.7, residuals=[residual])
        else:
            packet = HandoffPacket(
                title="Ampliar la ficha o las senales de orden",
                steps=[
                    "La ficha no activo ningun trigger de aplicabilidad (>= 0.55).",
                    "Agregar sector, lo que pide el cliente (IATF, PCF, OTIF), logistica o senales de orden.",
                    "No completar el vacio con un catalogo generico: cero normas es el resultado correcto.",
                ],
                data={"detectadas": 0},
                risk_if_skipped=(
                    "Rellenar normas a mano convierte un no-match en un diagnostico falso."
                ),
            )
            summary = (
                f"{detection.get('aviso', AVISO_DETECCION)} "
                "Ninguna norma alcanzo el umbral de aplicabilidad."
            )
            outcome = as_handoff(summary, [packet], confidence=0.5, residuals=[residual])

    return IsoOpsReport(
        aviso=str(detection.get("aviso") or AVISO_DETECCION),
        detectadas=detectadas,
        orden=orden,
        outcome=outcome,
        summary=summary,
        n_detectadas=len(detectadas),
        n_nativas_orden=n_nativas,
        n_huecos=n_huecos,
        perfil_vacio=prepared.perfil.vacio(),
        source=prepared.source,
    )


def verify(report: IsoOpsReport) -> list[str]:
    """QA gate: never-unprotected + never a conformity ruling."""
    issues = list(verify_guided(report.outcome))
    if report.outcome.status == EXECUTED:
        issues.append("outcome must never be EXECUTED: iso_ops detects and proposes, it does not certify")
    blobs = [report.summary, report.aviso, str(report.outcome.summary)]
    for row in report.detectadas:
        blobs.append(str(row.get("claimPermitido", "")))
        blobs.append(str(row.get("limite", "")))
    for opt in report.outcome.options:
        blobs.extend([opt.label, opt.summary, opt.action])
    for packet in report.outcome.handoffs:
        blobs.append(packet.title)
        blobs.extend(packet.steps)
        if not packet.risk_if_skipped.strip():
            issues.append(f"handoff without risk_if_skipped: '{packet.title}'")
    for blob in blobs:
        if texto_prohibido(blob):
            issues.append(f"forbidden conformity claim: {blob[:80]}")
    if report.perfil_vacio and report.n_detectadas != 0:
        issues.append("empty plant profile must not invent detections")
    if report.n_detectadas != len(report.detectadas):
        issues.append("n_detectadas does not match detectadas")
    return issues


def write_operational(report: IsoOpsReport, out_dir: str | Path, client: str = "Client") -> dict[str, Path]:
    """CSV rows for detections and floor tasks, plus a workbook the operator can hand over."""
    del client
    d = Path(out_dir)
    d.mkdir(parents=True, exist_ok=True)

    detect_rows = []
    task_rows = []
    for row in report.detectadas:
        detect_rows.append({
            "id": row.get("id"),
            "codigo": row.get("codigo"),
            "aplicabilidad": row.get("aplicabilidad"),
            "potencial": row.get("potencial"),
            "cobertura": row.get("cobertura"),
            "claimPermitido": row.get("claimPermitido"),
            "limite": row.get("limite"),
            "huecos": " | ".join(row.get("huecos") or []),
        })
        for task in row.get("tareas") or []:
            task_rows.append({
                "norma_id": row.get("id"),
                "codigo": row.get("codigo"),
                "tarea_id": task.get("id"),
                "titulo": task.get("titulo"),
                "detalle": task.get("detalle"),
                "kern": " | ".join(task.get("kern") or []),
            })

    orden_rows = []
    if report.orden:
        for card in report.orden.get("evidencias") or []:
            orden_rows.append({
                "id": card.get("id"),
                "codigo": card.get("codigo"),
                "cobertura": card.get("cobertura"),
                "claimPermitido": card.get("claimPermitido"),
                "limite": card.get("limite"),
                "tipo": "nativa",
            })
        fuera = report.orden.get("fueraDeTramo") or {}
        if fuera:
            orden_rows.append({
                "id": fuera.get("id"),
                "codigo": fuera.get("codigo"),
                "cobertura": "",
                "claimPermitido": fuera.get("claimPermitido"),
                "limite": fuera.get("limite"),
                "tipo": "hueco",
            })

    paths: dict[str, Path] = {
        "csv": write_summary_csv(detect_rows, d / "iso_ops_detectadas.csv"),
        "tareas_csv": write_summary_csv(task_rows, d / "iso_ops_tareas.csv"),
        "excel": _write_xlsx(report, detect_rows, task_rows, orden_rows, d / "iso_ops.xlsx"),
    }
    if orden_rows:
        paths["orden_csv"] = write_summary_csv(orden_rows, d / "iso_ops_orden.csv")
    return paths


def build_deck(
    report: IsoOpsReport,
    *,
    client: str,
    prepared: str | None = None,
    citations: tuple[str, ...] = (),
    confidence: float | None = None,
    lang: str = "en",
) -> Deliverable:
    """Client-facing deck: applicable norms, named gaps, floor tasks. Never a certificate.

    ``confidence`` is derived when not supplied: an empty ficha or named
    huecos are coverage gaps, not silent certainty.
    """
    if confidence is None:
        gaps: list[str] = []
        if report.perfil_vacio or report.n_detectadas == 0:
            gaps.append("ignored_criterion")
        if report.n_huecos:
            gaps.append("ignored_criterion")
        confidence = confidence_from_coverage(gaps, base=_DECK_BASE_CONFIDENCE)
    findings: list[Finding] = []
    if report.detectadas:
        for row in report.detectadas[:6]:
            huecos = " | ".join(row.get("huecos") or []) or "sin hueco nombrado"
            findings.append(Finding(
                f"{row.get('codigo')}: {row.get('claimPermitido')}",
                f"limite: {row.get('limite')} huecos: {huecos}",
                impact=f"aplicabilidad {row.get('aplicabilidad')}, potencial {row.get('potencial')}",
            ))
    elif report.orden and report.orden.get("evidencias"):
        for card in report.orden["evidencias"][:4]:
            findings.append(Finding(
                str(card.get("claimPermitido") or card.get("codigo")),
                f"limite: {card.get('limite')}",
                impact="evidencia de orden, no certificado",
            ))
    else:
        findings.append(Finding(
            "Ficha vacia o sin match: cero normas",
            "Kern no inventa un catalogo cuando la ficha no activa triggers. "
            "Cargar sector, cliente, logistica o senales de orden y volver a correr.",
            impact="HANDOFF: completar ficha",
        ))

    kpis = (
        Kpi("normas detectadas", str(report.n_detectadas),
            rationale="Aplicabilidad >= 0.55 sobre el catalogo iso-operaciones-cadena"),
        Kpi("evidencias nativas de orden", str(report.n_nativas_orden),
            rationale="9001 / 50001 / 14067 solo cuando hay ciclo, scrap o kg"),
        Kpi("huecos nombrados", str(report.n_huecos),
            rationale="Un hueco (p.ej. 14083 sin t·km) se nombra; no se tapa con OEE"),
    )
    recommendations = [
        "Usar el claimPermitido de cada tarjeta; no sustituir al organismo de certificacion.",
        "Cerrar primero los huecos nombrados (t·km del 3PL para 14083 si aplica).",
        "Las tareas de piso salen de norma.tareas del catalogo: son trabajo operativo, no auditoria de certificacion.",
        "Odoo y Excel siguen siendo politica de inventario; este entregable es evidencia de piso, no un sello.",
    ]
    residual = " ".join(
        f"{r.description} {r.risk_if_skipped}".strip() for r in report.outcome.residuals
    ) or _CERTIFY_RESIDUAL.description

    return Deliverable(
        title="ISO operations — floor evidence (not a certificate)",
        client=client,
        summary=report.summary,
        findings=tuple(findings),
        kpis=kpis,
        data_sources=(
            DataSource("Catalogo ISO operaciones/cadena", "knowledge/iso/catalogo.json", "pack version"),
            DataSource("Ficha de planta / senales de orden", report.source, "per run"),
        ),
        recommendations=tuple(recommendations),
        options=tuple(report.outcome.options),
        citations=tuple(citations),
        confidence=confidence,
        residual=residual,
        prepared=prepared or "",
        lang=lang,
    )


def _enriquecer(row: Mapping[str, Any]) -> dict[str, Any]:
    norma = buscar_norma(str(row.get("id") or ""))
    kern = (norma or {}).get("kern") or {}
    tareas = list((norma or {}).get("tareas") or [])
    claim = ""
    if norma is not None:
        card = tarjeta_evidencia(norma)
        claim = str(card.get("claimPermitido") or "")
    return {
        **dict(row),
        "cobertura": kern.get("cobertura") or "",
        "limite": kern.get("limite") or "",
        "claimPermitido": claim,
        "tareas": tareas,
    }


def _options(detectadas: tuple[dict[str, Any], ...], orden: dict[str, Any] | None) -> list[ExecutionOption]:
    options: list[ExecutionOption] = []
    ranked = sorted(
        detectadas,
        key=lambda r: (-float(r.get("aplicabilidad") or 0), -float(r.get("potencial") or 0), str(r.get("id"))),
    )
    for i, row in enumerate(ranked[:8]):
        tareas = row.get("tareas") or []
        first = tareas[0] if tareas else {}
        huecos = row.get("huecos") or []
        if huecos:
            action = f"Nombrar y cerrar el hueco: {huecos[0]}"
        elif first:
            action = f"Tarea de piso: {first.get('titulo')}"
        else:
            action = f"Revisar evidencia de piso para {row.get('codigo')} sin emitir certificado."
        score = round(float(row.get("aplicabilidad") or 0) * 0.6 + float(row.get("potencial") or 0) * 0.4, 4)
        options.append(ExecutionOption(
            label=str(row.get("codigo") or row.get("id")),
            summary=str(row.get("claimPermitido") or AVISO_DETECCION),
            score=score,
            recommended=i == 0,
            action=action,
            tradeoffs=str(row.get("limite") or ""),
        ))
    if not options and orden:
        for i, card in enumerate((orden.get("evidencias") or [])[:4]):
            options.append(ExecutionOption(
                label=str(card.get("codigo") or card.get("id")),
                summary=str(card.get("claimPermitido") or AVISO_ORDEN),
                score=round(0.8 - 0.05 * i, 4),
                recommended=i == 0,
                action="Usar la evidencia de esta orden; no certificar la planta.",
                tradeoffs=str(card.get("limite") or ""),
            ))
    return options


def _perfil_body(merged: Mapping[str, Any]) -> dict[str, Any]:
    body: dict[str, Any] = {"texto": str(merged.get("texto") or "")}
    for key in _LIST_FIELDS:
        raw = merged.get(key, merged.get(_snake(key), []))
        body[key] = _as_str_list(raw)
    return body


def _from_json(path: Path) -> dict[str, Any]:
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError("el JSON de ficha ISO tiene que ser un objeto")
    if isinstance(data.get("perfilEntrada"), dict):
        payload = dict(data["perfilEntrada"])
    elif isinstance(data.get("contratoPost"), dict):
        payload = dict(data["contratoPost"])
    else:
        payload = dict(data)
    senales = data.get("senales") if isinstance(data.get("senales"), dict) else {}
    for key in ("ciclo", "scrap", "kg", "tkm"):
        if key in senales:
            payload[key] = senales[key]
        elif key in data:
            payload[key] = data[key]
    return payload


def _from_table(path: Path) -> dict[str, Any]:
    df = pd.read_excel(path) if path.suffix.lower() in {".xlsx", ".xls"} else pd.read_csv(path)
    if df.empty:
        raise ValueError("the ISO plant-profile file has no rows")
    cols = {_norm(c): c for c in df.columns}
    row = df.iloc[0]
    payload: dict[str, Any] = {}
    for canonical, aliases in _LIST_ALIASES.items():
        header = next((cols[a] for a in aliases if a in cols), None)
        if header is not None:
            payload[canonical] = _as_str_list(row[header])
    texto_header = next((cols[a] for a in _TEXTO_ALIASES if a in cols), None)
    if texto_header is not None and pd.notna(row[texto_header]):
        payload["texto"] = str(row[texto_header]).strip()
    for canonical, aliases in _SIGNAL_ALIASES.items():
        header = next((cols[a] for a in aliases if a in cols), None)
        if header is not None:
            payload[canonical] = _as_bool(row[header])
    return payload


def _write_xlsx(
    report: IsoOpsReport,
    detect_rows: list[dict[str, Any]],
    task_rows: list[dict[str, Any]],
    orden_rows: list[dict[str, Any]],
    path: Path,
) -> Path:
    wb = Workbook()
    aviso = wb.active
    aviso.title = "Aviso"
    aviso["A1"] = report.aviso
    aviso["A2"] = "Kern no certifica. Odoo/Excel no llevan un sello ISO; este libro es evidencia de piso."
    aviso["A3"] = report.summary
    aviso.column_dimensions["A"].width = 88

    _sheet(wb, "Detectadas", detect_rows)
    _sheet(wb, "Tareas", task_rows)
    if orden_rows:
        _sheet(wb, "Orden", orden_rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _sheet(wb: Workbook, title: str, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title)
    if not rows:
        ws["A1"] = "(sin filas)"
        return
    headers = list(rows[0].keys())
    fill = PatternFill("solid", fgColor="1F2A44")
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
    for ri, row in enumerate(rows, 2):
        for ci, h in enumerate(headers, 1):
            ws.cell(row=ri, column=ci, value=defuse_formula(row[h]))
    for ci, h in enumerate(headers, 1):
        width = max(len(str(h)), *(len(str(row[h])) for row in rows))
        ws.column_dimensions[get_column_letter(ci)].width = min(max(width + 2, 10), 40)
    ws.freeze_panes = "A2"


def _as_str_list(value: object) -> list[str]:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return []
    if isinstance(value, (list, tuple)):
        return [str(x).strip() for x in value if str(x).strip()]
    text = str(value).strip()
    if not text or text.lower() in {"nan", "none"}:
        return []
    return [part.strip() for part in text.replace(";", ",").split(",") if part.strip()]


def _as_bool(value: object) -> bool:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return False
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    return str(value).strip().lower() in {"1", "true", "yes", "si", "sí"}


def _snake(key: str) -> str:
    return "cliente_pide" if key == "clientePide" else key
