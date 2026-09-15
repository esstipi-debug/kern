"""Job-layer tests for the E4 SCOR/ISO evidence pack (composition, not a certificate)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from jobs import scor_audit_job
from nosilo_core.guided import HANDOFF
from src.iso_catalog import texto_prohibido

_FIX = Path(__file__).resolve().parent / "fixtures" / "audit" / "scor_pack.csv"

_PARAMS = {
    "as_of": "2026-09-03",
    "dso_days": 45.0,
    "dpo_days": 30.0,
    "produced_at": "2026-09-07T12:00:00Z",
    "run_id": "test-scor-iso",
}


def test_prepare_requires_dataset_column(tmp_path):
    path = tmp_path / "no_tag.csv"
    path.write_text("product_id,cogs,avg_inventory_value\nP1,100,50\n", encoding="utf-8")
    try:
        scor_audit_job.prepare(path)
    except ValueError as exc:
        assert "dataset" in str(exc)
    else:
        raise AssertionError("untagged CSV must be refused")


def test_pack_composes_panels_and_never_claims_conformity(tmp_path):
    prepared = scor_audit_job.prepare(_FIX, _PARAMS)
    assert "financial" in prepared.frames
    assert "orders" in prepared.frames
    report = scor_audit_job.run(prepared, _PARAMS)
    assert report.outcome.status == HANDOFF
    assert scor_audit_job.verify(report) == []
    assert report.n_observed + report.n_gap + report.n_no_evaluable == len(report.checks)
    assert any(m.code == "RL.1.1" for m in report.metrics)
    assert any(m.code == "AM.1.1" for m in report.metrics)
    assert all(c.status in {scor_audit_job.OBSERVED, scor_audit_job.GAP, scor_audit_job.NO_EVALUABLE}
               for c in report.checks)
    assert report.evidence.prepared_by == scor_audit_job.ENGINE_IDENTITY
    blobs = [report.summary, *(c.evidence for c in report.checks), *(m.note for m in report.metrics)]
    for blob in blobs:
        assert not texto_prohibido(blob)
        assert not scor_audit_job._iso_claim_forbidden(blob)

    written = scor_audit_job.write_operational(report, tmp_path, "Acme")
    wb = load_workbook(written["workbook"])
    assert {ws.title for ws in wb.worksheets} == {
        "Scorecard", "ISO checklist", "Coverage", "Lineage",
    }
    lineage = wb["Lineage"]
    reviewed = None
    for row in lineage.iter_rows(min_col=1, max_col=2, values_only=True):
        if row[0] == "Reviewed by":
            reviewed = row[1]
    assert reviewed in (None, "")


def test_missing_orders_panel_is_no_evaluable_not_compliant(tmp_path):
    path = tmp_path / "fin_only.csv"
    path.write_text(
        "dataset,product_id,cogs,avg_inventory_value\nfinancial,P1,50000,12000\n",
        encoding="utf-8",
    )
    report = scor_audit_job.run(scor_audit_job.prepare(path), _PARAMS)
    rl = next(m for m in report.metrics if m.code == "RL.1.1")
    assert rl.status == scor_audit_job.NO_EVALUABLE
    assert rl.value is None
    check = next(c for c in report.checks if "RL.1.1" in c.clause)
    assert check.status == scor_audit_job.NO_EVALUABLE
    assert scor_audit_job.verify(report) == []
