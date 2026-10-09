"""Una negativa del motor no puede salir vestida de numero.

El defecto, medido punta a punta antes de arreglarlo: un SKU con cinco semanas de
ventas REALES (~42 u/semana) y 120 unidades en el estante llegaba al operador como
un evento ``dead_stock`` -- *"120 units on hand with no demand"*. La cadena era
entera y cada eslabon estaba bien escrito:

1. ``jobs/inventory_optimization.py`` se niega a dimensionar con menos de
   ``MIN_PERIODS_FOR_POLICY`` periodos -- correcto, y el comentario lo dice:
   *"Refuse rather than degrade"*.
2. ...pero emitia la negativa con ``forecast=0.0``, ``reorder_point=0.0``,
   ``safety_stock=0.0`` y ``z_factor=0.0``, con ``status="ok"`` y
   ``qa_issues=[]``. El aviso viajaba por otro canal (``policy_kind``, un residual
   de ``scm_agent/tool_options.py``), o sea por la prosa.
3. ``jobs/state_bridge.py::_usable`` existe justamente para omitir esa linea. Su
   docstring lo dice literal: *"un disparador ausente no es cero (pediria siempre)
   y una demanda diaria ausente no es cero (daria cobertura infinita justo sobre el
   SKU del que no se sabe nada)"*. El 0.0 del productor la DERROTABA: 0.0 no es
   ``None``, asi que la linea pasaba.
4. ``src/alerting.py`` ve ``avg_daily_demand == 0`` con stock en mano y emite
   ``dead_stock``.

Un 0 significa "medido, y la respuesta es ninguno". Una negativa no midio nada.

Estos tests fijan las DOS direcciones, porque arreglar solo una cambia el defecto
de lugar: una linea sin politica no puede traer un numero, y una linea con politica
no puede traer un ``None``.
"""

from __future__ import annotations

import types

import pandas as pd
import pytest

from jobs import deliverables, qa
from jobs.inventory_optimization import UNSIZED_POLICY_KINDS, run
from jobs.state_bridge import stock_lines_from_inventory
from src.alerting import detect_events
from src.forecasting import MIN_PERIODS_FOR_POLICY

_PERIOD = "W"
_CAMPOS = ("forecast", "reorder_point", "safety_stock", "z_factor")

#: Lo que hace falta hoy para una negativa ENTERA: menos de `MIN_PERIODS_FOR_COVER`
#: periodos. Desde el 2026-09-20 el motor ya no contesta las dos preguntas con el
#: mismo silencio -- con dos observaciones hay una media y un plazo medidos, y con
#: ellos emite una COBERTURA PROVISIONAL (`tests/test_cobertura_provisional.py`).
#: Con UNA sola no hay promedio que citar: llamar "tasa de demanda" a una
#: observacion es inventar, y ahi la negativa se queda entera. Ese es el sujeto de
#: este archivo, y por eso las series de abajo son dos y no una.
_SERIE_SIN_MEDIA = [40]

#: La serie con la que se REPRODUJO el defecto de cabecera: cinco semanas vendiendo
#: ~42 u/semana. Hoy recibe una cobertura, no una negativa -- asi que sirve para la
#: otra mitad del mismo caso: tampoco puede salir como stock muerto.
_SERIE_CON_COBERTURA = [40, 44, 38, 47, 41]


def _archivo(product_id: str, cantidades, *, unit_cost: float = 10.0,
             on_hand: float | None = None) -> pd.DataFrame:
    fechas = pd.date_range("2026-01-04", periods=len(cantidades), freq=_PERIOD)
    filas = []
    for f, q in zip(fechas, cantidades, strict=True):
        fila = {"date": f, "product_id": product_id, "quantity": float(q),
                "unit_cost": unit_cost, "lead_time_days": 14.0}
        if on_hand is not None:
            fila["on_hand"] = on_hand
        filas.append(fila)
    return pd.DataFrame(filas)


def _rec(df: pd.DataFrame, pid: str, *, on_hand: float | None = None, **kw):
    """``run()`` recibe la posicion por argumento, no por columna del frame: la
    columna la lee ``jobs/intake.py::stock_position`` en el ``prepare`` del tool."""
    if on_hand is not None:
        kw["on_hand_by_sku"] = {pid: on_hand}
    reporte = run(df, period_freq=_PERIOD, **kw)
    return reporte, {r.product_id: r for r in reporte.recommendations}[pid]


# ---------------------------------------------------------------------------
# 1. El productor: ausencia es None, medicion es 0
# ---------------------------------------------------------------------------


def test_historia_insuficiente_anula_los_cuatro_campos_dimensionados():
    """La negativa por conocimiento: el SKU vende, no sabemos NI CUANTO NI EN PROMEDIO."""
    _, r = _rec(_archivo("NUEVO", _SERIE_SIN_MEDIA), "NUEVO")

    assert r.policy_kind == "insufficient_history"
    assert r.n_periods < MIN_PERIODS_FOR_POLICY
    for campo in _CAMPOS:
        assert getattr(r, campo) is None, f"{campo} salio {getattr(r, campo)!r}, no None"
    # Los HECHOS medidos no se anulan: solo lo que el motor no calculo.
    assert r.n_periods == 1
    assert r.unit_cost == 10.0
    assert r.investment == 0.0   # no se estaciona plata: eso SI es un cero medido
    # Y no hay media que publicar: con una observacion, "el promedio" es esa
    # observacion. El campo de la cobertura tambien queda ausente.
    assert r.observed_mean_per_period is None
    assert r.cover_periods is None


def test_sin_costo_usable_anula_lo_mismo_porque_la_demanda_es_real():
    """La otra negativa: hay demanda de sobra, lo que falta es el costo.

    Un 0 aca diria que el SKU no vende, y vende todas las semanas.
    """
    _, r = _rec(_archivo("SIN-COSTO", [30] * 12, unit_cost=0.0), "SIN-COSTO")

    assert r.policy_kind == "skipped_zero_cost"
    for campo in _CAMPOS:
        assert getattr(r, campo) is None, f"{campo} salio {getattr(r, campo)!r}, no None"


def test_demanda_medida_en_cero_CONSERVA_sus_ceros():
    """La mitad de la regla que no cambia, y el caso que prueba que la regla es
    sobre la ausencia y no sobre la rama.

    Aca la demanda agregada SI se midio y dio cero. "El pronostico es 0 y el
    disparador es 0" es una medicion -- exactamente lo que un 0 debe significar.
    """
    _, r = _rec(_archivo("SIN-VENTA", [0] * 12), "SIN-VENTA")

    assert r.policy_kind == "skipped_zero_demand"
    for campo in _CAMPOS:
        assert getattr(r, campo) == 0.0, f"{campo} salio {getattr(r, campo)!r}, no 0.0"


# ---------------------------------------------------------------------------
# 2. La regresion punta a punta: el falso dead_stock
# ---------------------------------------------------------------------------


def _eventos_del_tower(reporte):
    lineas = stock_lines_from_inventory(reporte)
    return lineas, detect_events([
        {"product_id": ln.product_id, "on_hand": ln.on_hand,
         "reorder_point": ln.reorder_point, "avg_daily_demand": ln.avg_daily_demand}
        for ln in lineas
    ])


@pytest.mark.parametrize("serie,kind", [
    (_SERIE_SIN_MEDIA, "insufficient_history"),
    (_SERIE_CON_COBERTURA, "provisional_cover"),
])
def test_un_sku_que_vende_todas_las_semanas_NO_sale_como_stock_muerto(serie, kind):
    """El caso de cabecera, en sus DOS formas.

    La segunda es el archivo original con el que se reprodujo el defecto: cinco
    semanas a ~42 u/semana. Hoy sale con una cobertura provisional en vez de una
    negativa, y el falso `dead_stock` tiene que seguir sin salir -- por otra razon,
    no por la misma. Si el escalon de la cobertura hubiera devuelto un `forecast`
    numerico para que el Tower lo tomara, el defecto volvia por esa puerta.
    """
    reporte, r = _rec(_archivo("NUEVO", serie), "NUEVO", on_hand=120.0)
    assert r.policy_kind == kind
    assert r.on_hand == 120.0, "el estante es un hecho medido y tiene que seguir ahi"

    lineas, eventos = _eventos_del_tower(reporte)
    assert lineas == [], "un SKU sin tasa de demanda no se puede monitorear; se omite"
    assert [e.kind for e in eventos] == []


def test_la_cobertura_deja_al_Tower_CALLADO_y_no_es_un_descuido():
    """Por que la cobertura no entra al sensado, dicho donde se nota.

    Tiene disparador (la cobertura) y tiene una tasa medida
    (`observed_mean_per_period`), asi que tecnicamente se podria monitorear. No se
    hace, y la razon es la misma que hace honesta a la cobertura: no tiene buffer,
    o sea que se espera que el estante caiga por debajo la MITAD de las veces. Un
    `rop_breach` que suena la mitad de las semanas sobre cada SKU nuevo no es una
    alarma, es ruido -- y el Tower deja de leerse. `forecast` en `None` es lo que
    hoy lo mantiene afuera; este test fija que salga por esa puerta a proposito.
    """
    reporte, r = _rec(_archivo("NUEVO", _SERIE_CON_COBERTURA), "NUEVO", on_hand=120.0)
    assert r.policy_kind == "provisional_cover"
    assert r.reorder_point is not None and r.observed_mean_per_period is not None
    assert r.forecast is None
    assert stock_lines_from_inventory(reporte) == []


def test_el_dead_stock_VERDADERO_sigue_saliendo():
    """El freno del arreglo anterior: anular de mas apagaria una alarma correcta.

    Un SKU con demanda medida en cero y stock parado SI es stock muerto, y ese
    evento tiene que sobrevivir -- por eso `skipped_zero_demand` conserva sus ceros.
    """
    reporte, r = _rec(_archivo("SIN-VENTA", [0] * 12), "SIN-VENTA", on_hand=90.0)
    assert r.policy_kind == "skipped_zero_demand"

    lineas, eventos = _eventos_del_tower(reporte)
    assert len(lineas) == 1
    assert [e.kind for e in eventos] == ["dead_stock"]


# ---------------------------------------------------------------------------
# 3. La compuerta de QA, en las dos direcciones
# ---------------------------------------------------------------------------


def test_qa_pasa_sobre_una_negativa_bien_declarada():
    reporte, _ = _rec(_archivo("NUEVO", _SERIE_SIN_MEDIA), "NUEVO")
    assert qa.verify(reporte) == []


@pytest.mark.parametrize("campo", _CAMPOS)
def test_qa_rechaza_una_negativa_que_trae_un_numero(campo):
    """Por mutacion: se vuelve a poner el 0.0 de antes, campo por campo."""
    from dataclasses import replace

    reporte, r = _rec(_archivo("NUEVO", _SERIE_SIN_MEDIA), "NUEVO")
    mutado = replace(reporte, recommendations=[replace(r, **{campo: 0.0})])

    issues = qa.verify(mutado)
    assert any("declined to size" in i for i in issues), issues
    assert any(campo in i for i in issues), issues


@pytest.mark.parametrize("campo", _CAMPOS)
def test_qa_rechaza_una_politica_con_precio_a_la_que_le_falta_un_numero(campo):
    """La direccion contraria: anularlo TODO tambien seria mentir.

    Sin este test, "arreglar" el defecto poniendo None en todas partes pasaria la
    compuerta y dejaria al operador sin politica en los SKU que si se dimensionaron.
    """
    from dataclasses import replace

    reporte, r = _rec(_archivo("LARGO", [20 + (i % 4) for i in range(24)]), "LARGO")
    assert r.policy_kind in {"(s, Q)", "(R, S)"}
    mutado = replace(reporte, recommendations=[replace(r, **{campo: None})])

    issues = qa.verify(mutado)
    assert any(f"no {campo}" in i for i in issues), issues


def test_las_TRES_listas_de_politicas_sin_dimensionar_coinciden():
    """La misma lista estaba escrita a mano en tres lugares y las tres no coincidian.

    `jobs.qa` tenia las tres clases; `src.verify.coherencia` y
    `scm_agent.package_specs` tenian solo dos, olvidando `insufficient_history`.
    Cada copia perdida costaba algo distinto: en coherencia, un SKU sin politica
    pasaba toda invariante trivialmente; en package_specs -- la que mas dolia --
    ese SKU entraba a la PLANILLA DE PEDIDOS del cliente como `kern_optimized` con
    disparador cero, pisando el numero propio del cliente y suprimiendo una
    reposicion real.

    Ahora `jobs.inventory_optimization.UNSIZED_POLICY_KINDS` es la definicion y las
    otras la leen o la repiten; este test es el que impide que se vuelvan a separar.

    Desde la cobertura provisional son DOS listas y no una, y la diferencia importa:
    `UNSIZED_POLICY_KINDS` es "el motor no puso ningun numero" (y por eso `jobs.qa`
    exige los cuatro campos en `None`), mientras que `NOT_FULLY_SIZED_POLICY_KINDS`
    agrega la cobertura -- que SI trae un disparador, pero sin buffer ni z, asi que
    las invariantes por SKU de coherencia y la planilla del cliente la tienen que
    excluir igual. Colapsarlas en una sola volveria a romper algo en los dos
    sentidos: con la lista chica, un SKU con `safety_stock=None` reventaria
    coherencia; con la grande, `jobs.qa` exigiria `reorder_point=None` sobre una
    cobertura que existe.
    """
    from jobs.inventory_optimization import (
        NOT_FULLY_SIZED_POLICY_KINDS,
        PROVISIONAL_POLICY_KINDS,
    )
    from src.verify.coherencia import POLITICAS_SALTADAS

    assert qa.UNPRICED_POLICY_KINDS == UNSIZED_POLICY_KINDS
    assert POLITICAS_SALTADAS == NOT_FULLY_SIZED_POLICY_KINDS
    # Y que las listas sean las que el productor de verdad emite, no unas escritas
    # al lado.
    assert UNSIZED_POLICY_KINDS == {
        "insufficient_history", "skipped_zero_cost", "skipped_zero_demand",
    }
    assert PROVISIONAL_POLICY_KINDS == {"provisional_cover"}
    assert NOT_FULLY_SIZED_POLICY_KINDS == UNSIZED_POLICY_KINDS | PROVISIONAL_POLICY_KINDS
    assert not (UNSIZED_POLICY_KINDS & PROVISIONAL_POLICY_KINDS), (
        "una cobertura no es una negativa: `jobs.qa` exige formas OPUESTAS sobre "
        "las dos, asi que si una kind cae en las dos listas la compuerta se "
        "contradice a si misma"
    )

    # Y que sea el productor el que lo demuestre, no solo la constante: las tres
    # formas salen de tres archivos reales.
    emitidas = set()
    for serie, costo in ((_SERIE_SIN_MEDIA, 10.0), ([0] * 12, 10.0),
                         ([30] * 12, 0.0), (_SERIE_CON_COBERTURA, 10.0)):
        _, r = _rec(_archivo("X", serie, unit_cost=costo), "X")
        emitidas.add(r.policy_kind)
    assert emitidas == NOT_FULLY_SIZED_POLICY_KINDS


def test_un_sku_sin_politica_no_entra_a_la_planilla_de_pedidos_del_cliente():
    """El consumidor mas consecuente de todos, y el que la lista incompleta dejaba pasar.

    `scm_agent/package_specs.py::_optimized_targets_from_inventory` enhebra la
    politica del optimizador dentro de `excel_replenishment` como target preferido, y
    esa planilla es lo que el cliente usa para pedir. Un SKU sin politica ahi se
    renderiza `source="kern_optimized"` con `trigger=0.0`: no solo no dispara, ademas
    PISA el numero propio del cliente y lo hace en nombre de Kern.

    Van las DOS formas sin politica en la misma cartera. La cobertura es la que
    podria colarse ahora: a diferencia de la negativa, SI trae un `reorder_point`
    positivo, asi que un filtro escrito como "excluir los que traen cero" la dejaria
    pasar. Y pisar el numero del cliente con una cobertura sin buffer es peor que
    pisarlo con un cero, porque parece una politica.
    """
    from scm_agent.package_specs import _optimized_targets_from_inventory

    fechas = pd.date_range("2026-01-04", periods=24, freq=_PERIOD)
    filas = [{"date": f, "product_id": "LARGO", "quantity": 20.0 + (i % 4),
              "unit_cost": 40.0} for i, f in enumerate(fechas)]
    filas += [{"date": fechas[-1], "product_id": "NUEVO", "quantity": 42.0,
               "unit_cost": 10.0}]
    filas += [{"date": fechas[-i - 1], "product_id": "CUBIERTO", "quantity": 42.0,
               "unit_cost": 10.0} for i in range(5)]
    reporte = run(pd.DataFrame(filas), period_freq=_PERIOD)
    kinds = {r.product_id: r.policy_kind for r in reporte.recommendations}
    assert kinds == {"LARGO": "(s, Q)", "NUEVO": "insufficient_history",
                     "CUBIERTO": "provisional_cover"}
    cubierto = {r.product_id: r for r in reporte.recommendations}["CUBIERTO"]
    assert cubierto.reorder_point > 0, (
        "el test no probaria nada si la cobertura saliera en cero: el filtro viejo "
        "la habria excluido por accidente"
    )

    targets = _optimized_targets_from_inventory(
        {"inventory_optimization": reporte}
    )["optimized_targets"]

    for pid in ("NUEVO", "CUBIERTO"):
        assert pid not in targets, (
            f"{pid} no puede viajar a la planilla del cliente: degrada al numero "
            "propio del cliente justamente por estar ausente del dict"
        )
    assert "LARGO" in targets and targets["LARGO"]["reorder_point"] > 0


# ---------------------------------------------------------------------------
# 4. Lo que el operador lee
# ---------------------------------------------------------------------------


def test_la_planilla_deja_la_celda_vacia_y_el_markdown_una_raya():
    _, r = _rec(_archivo("NUEVO", _SERIE_SIN_MEDIA), "NUEVO")
    fila = deliverables._row(r)

    assert fila["reorder_point_s"] is None
    assert fila["safety_stock"] is None
    assert fila["forecast_per_period"] is None
    assert deliverables._celda(None, 0) == "—"
    assert deliverables._celda(12.0, 0) == "12"


def test_el_markdown_no_imprime_un_cero_en_el_disparador(tmp_path):
    reporte, _ = _rec(_archivo("NUEVO", _SERIE_SIN_MEDIA), "NUEVO")
    destino = deliverables.write_report_md(reporte, tmp_path / "r.md")
    fila = [ln for ln in destino.read_text(encoding="utf-8").splitlines()
            if ln.startswith("| NUEVO |")]

    assert fila, "el SKU refusado tiene que aparecer en la tabla, no desaparecer"
    assert "insufficient_history" in fila[0]
    assert "—" in fila[0]
    assert " 0 |" not in fila[0] and " 0.0 |" not in fila[0]


def test_el_resumen_nombra_la_negativa_en_vez_de_recomendar_cero(tmp_path):
    """Punta a punta por el orquestador, que es como lo lee el operador."""
    from scm_agent.orchestrator import Orchestrator

    destino = tmp_path / "v.csv"
    _archivo("NUEVO", _SERIE_SIN_MEDIA, on_hand=120.0).to_csv(destino, index=False)
    res = Orchestrator(clients_root=None).run(
        "punto de reorden y stock de seguridad por SKU",
        data_path=str(destino), job_type="inventory_optimization",
        out_dir=str(tmp_path / "out"),
    )

    assert res.status == "ok"
    assert "No policy was sized" in res.summary
    assert "NOT a recommendation" in res.summary
    assert "recommended inventory investment $0" not in res.summary


# ---------------------------------------------------------------------------
# 5. El marcador: la cantidad emitida no cambia, la trazabilidad si
# ---------------------------------------------------------------------------


def test_el_corredor_omite_el_sku_sin_politica_en_vez_de_ponerle_cero(monkeypatch):
    from src.simworld import runners

    rec_ok = types.SimpleNamespace(product_id="A", reorder_point=30.0, forecast=6.0,
                                   order_up_to=50.0, order_quantity=20.0)
    rec_no = types.SimpleNamespace(product_id="B", reorder_point=None, forecast=None,
                                   order_up_to=None, order_quantity=None)
    resultado = types.SimpleNamespace(
        status="ok", summary="s", qa_issues=[], guided=None, confidence=1.0,
        report=types.SimpleNamespace(recommendations=[rec_ok, rec_no],
                                     intake_quality=None, intake_quality_threshold=0.2),
    )
    monkeypatch.setattr(runners, "export_history", lambda *a, **k: a[2])
    monkeypatch.setattr(runners, "Orchestrator",
                        lambda **k: types.SimpleNamespace(run=lambda *a, **k2: resultado))

    pol = runners.kern_policy(None, [], level=0, seed=0, out_dir="/tmp/x", forced=True)

    assert "A" in pol.reorder_point
    assert "B" not in pol.reorder_point, "un SKU sin politica no entra al plan con un 0"
    assert "B" not in pol.order_up_to
    assert "B" not in pol.forecast


def test_omitir_y_poner_cero_emiten_la_MISMA_orden():
    """Por que el arreglo no mueve la captura del marcador, dicho como test.

    `apply_policy` ya trata un SKU ausente igual que uno con disparador 0: no pide.
    O sea que el numero publicado (-40,6% limpia) no estaba inflado por esto -- el
    carril se comportaba bien y lo REPORTABA mal. Lo que cambia es que ahora se
    puede distinguir "me negue" de "decidi no reponer".
    """
    import numpy as np

    from src.simworld.runners import Policy, apply_policy
    from src.simworld.world import WorldState, generate_world

    mundo = generate_world(seed=7, n_skus=3, n_suppliers=1, weeks=6)
    estado = WorldState(week=0, on_hand=np.array([5.0, 5.0, 5.0]), in_transit=())
    ausente = Policy(reorder_point={}, order_up_to={})
    en_cero = Policy(reorder_point={s.sku: 0.0 for s in mundo.skus},
                     order_up_to={s.sku: 0.0 for s in mundo.skus})

    assert apply_policy(mundo, estado, ausente) == apply_policy(mundo, estado, en_cero)


# ---------------------------------------------------------------------------
# 6. El consumidor que el mapa encontro y la suite no: la pantalla de brecha
# ---------------------------------------------------------------------------
#
# `webapp/demo_gap_value.py` RESTA dos dimensionamientos del mismo SKU (el del ERP
# contra el de la realidad) y arma `SkuGap`, un dataclass de floats. No filtraba por
# `policy_kind`: filtraba por tener plazo declarado, observado y dispersion. Un SKU
# con recepciones de sobra pero poca historia de demanda pasaba ese filtro y llegaba
# con los dos buffers en 0.0 -> `delta_value` 0,00, o sea "el ERP y la realidad
# coinciden en este SKU" sobre uno que nadie dimensiono. Con `None` ese mismo camino
# revienta con TypeError en una ruta de produccion (`webapp/app.py`), asi que la
# unica salida honesta es excluirlo Y CONTARLO.

_RECEIPTS = "data/sample_receipts_portfolio.csv"
_DEMAND = "data/sample_demand_portfolio.csv"


def _con_sku_sin_dimensionar(rec: pd.DataFrame, dem: pd.DataFrame):
    """Un SKU con recepciones de sobra y menos historia que el minimo.

    Proveedor PROPIO a proposito: colgarlo de uno existente mueve el scorecard de
    plazo de ese proveedor y con el las brechas de todos sus SKU, y entonces el test
    no podria afirmar que el `net` no se movio.
    """
    nuevas_rec = pd.DataFrame([
        {"po_id": f"OC-N{i}", "supplier": "Proveedor Nuevo SA", "product_id": "SKU-SIN-DIM",
         "order_date": f"2024-0{(i % 9) + 1}-01", "receipt_date": f"2024-0{(i % 9) + 1}-1{i % 8}",
         "declared_lead_time_days": 20}
        for i in range(6)
    ])
    ultimas = sorted(dem["date"].unique())[-4:]   # 4 periodos < MIN_PERIODS_FOR_POLICY
    nuevas_dem = pd.DataFrame([
        {"date": f, "product_id": "SKU-SIN-DIM", "quantity": 42.0,
         "unit_cost": 10.0, "lead_time_days": 20}
        for f in ultimas
    ])
    return (pd.concat([rec, nuevas_rec], ignore_index=True),
            pd.concat([dem, nuevas_dem], ignore_index=True))


def test_la_pantalla_de_brecha_excluye_y_CUENTA_el_sku_sin_dimensionar():
    pytest.importorskip("openpyxl")
    from webapp.demo_gap_value import run_gap_value

    rec, dem = pd.read_csv(_RECEIPTS), pd.read_csv(_DEMAND)
    base = run_gap_value(rec, dem)
    con = run_gap_value(*_con_sku_sin_dimensionar(rec, dem))

    # No revienta -- es una ruta de produccion del webapp.
    assert con.n_unsized_skus == 1
    assert "SKU-SIN-DIM" not in {g.product_id for g in con.gaps}
    # Y no se cuela como una brecha de cero, que era el sintoma viejo.
    assert len(con.gaps) == len(base.gaps)
    assert con.net == pytest.approx(base.net)
    assert base.n_unsized_skus == 0


def test_el_libro_de_brecha_declara_la_exclusion_en_el_embudo(tmp_path):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    from webapp.demo_gap_value import run_gap_value
    from webapp.demo_gap_workbook import write_gap_value_workbook

    rec, dem = pd.read_csv(_RECEIPTS), pd.read_csv(_DEMAND)

    def texto(valor, nombre):
        ruta = write_gap_value_workbook(valor, tmp_path / nombre,
                                        dataset_label="d.csv", ts="T")
        wb = load_workbook(ruta)
        return "\n".join(
            str(c) for hoja in wb.sheetnames
            for fila in wb[hoja].iter_rows(values_only=True) for c in fila if c
        )

    con = texto(run_gap_value(*_con_sku_sin_dimensionar(rec, dem)), "con.xlsx")
    assert "NO los dimensiono" in con

    # Y una corrida sin ninguno no gana un renglon que no dice nada.
    sin = texto(run_gap_value(rec, dem), "sin.xlsx")
    assert "NO los dimensiono" not in sin


# ---------------------------------------------------------------------------
# 7. La premisa del carve-out, hecha cierta por construccion
# ---------------------------------------------------------------------------
#
# `skipped_zero_demand` conserva sus ceros porque "la demanda se midio y dio cero".
# La rama dispara con `annual_demand <= 0`, que mezcla la medicion (exactamente 0)
# con un pronostico que salio BAJO cero -- `to_engine_inputs` multiplica por los
# periodos del anio y no acota por abajo. Un negativo no es la medicion de nada, y
# publicar 0.0 ahi diria "medimos cero" sobre un SKU que puede estar vendiendo.
#
# Al escribir el arreglo se creyo que era hipotetico: con series decrecientes CORTAS
# (seis a ocho puntos) el motor devuelve valores positivos muy por encima de la
# ultima observacion. Es falso, y lo encontro la revision adversarial: `auto_ets`
# extrapola bajo cero sobre una serie decreciente LARGA. El primer test de abajo usa
# esa serie real -- catorce semanas, 1.440 unidades vendidas -- y sin el arreglo el
# SKU saldria como `dead_stock`. El segundo forza el caso por monkeypatch para fijar
# el borde exacto (0.0 contra negativo) sin depender de que metodo elija el motor.


#: Serie decreciente larga con la que `auto_ets` extrapola bajo cero. Catorce
#: semanas, 1.440 unidades vendidas: un SKU que claramente NO es stock muerto.
_SERIE_QUE_PRONOSTICA_NEGATIVO = [200, 185, 170, 155, 140, 125, 110, 95, 80, 65,
                                  50, 35, 20, 10]


def test_una_serie_REAL_que_pronostica_negativo_no_sale_como_stock_muerto():
    """El caso alcanzable, sin monkeypatch, con el motor de pronostico de verdad.

    Si esto se pone verde porque el motor dejo de pronosticar negativo, la rama de
    `annual_demand < 0` deja de estar ejercida por un camino real y hay que
    revisarla -- por eso el test afirma primero que el pronostico ES negativo.
    """
    import numpy as np

    from src.forecasting import forecast_with_oos_sigma

    crudo = forecast_with_oos_sigma(
        np.array([float(x) for x in _SERIE_QUE_PRONOSTICA_NEGATIVO])
    )
    assert crudo.forecast < 0, (
        f"el motor ya no pronostica negativo aqui ({crudo.forecast}); la rama de "
        "annual_demand < 0 quedo sin camino real que la ejerza"
    )

    reporte, r = _rec(_archivo("BAJANDO", _SERIE_QUE_PRONOSTICA_NEGATIVO),
                      "BAJANDO", on_hand=120.0)
    assert r.policy_kind == "skipped_zero_demand"
    for campo in _CAMPOS:
        assert getattr(r, campo) is None, f"{campo} salio {getattr(r, campo)!r}"
    assert qa.verify(reporte) == []

    lineas, eventos = _eventos_del_tower(reporte)
    assert lineas == [] and eventos == [], (
        "un SKU que vendio 1.440 unidades en catorce semanas no puede salir como "
        "dead_stock por un pronostico que extrapolo bajo cero"
    )


def _forecast_negativo(monkeypatch, valor: float):
    from jobs import inventory_optimization as job

    real = job.forecast_with_oos_sigma

    def falso(series, **kw):
        import dataclasses
        return dataclasses.replace(real(series, **kw), forecast=valor)

    monkeypatch.setattr(job, "forecast_with_oos_sigma", falso)


def test_una_demanda_anual_NEGATIVA_no_publica_ceros(monkeypatch):
    _forecast_negativo(monkeypatch, -3.0)
    reporte, r = _rec(_archivo("BAJANDO", [40, 44, 38, 47, 41, 43, 39, 45]), "BAJANDO",
                      on_hand=120.0)

    assert r.policy_kind == "skipped_zero_demand"
    for campo in _CAMPOS:
        assert getattr(r, campo) is None, f"{campo} salio {getattr(r, campo)!r}"
    assert qa.verify(reporte) == []
    # Y por lo tanto el Tower NO lo llama stock muerto.
    lineas, eventos = _eventos_del_tower(reporte)
    assert lineas == [] and eventos == []


def test_una_demanda_anual_de_CERO_exacto_sigue_publicando_ceros(monkeypatch):
    """El otro lado del mismo borde: 0.0 es una medicion y no se toca."""
    _forecast_negativo(monkeypatch, 0.0)
    reporte, r = _rec(_archivo("PARADO", [40, 44, 38, 47, 41, 43, 39, 45]), "PARADO",
                      on_hand=120.0)

    assert r.policy_kind == "skipped_zero_demand"
    for campo in _CAMPOS:
        assert getattr(r, campo) == 0.0
    assert qa.verify(reporte) == []
    _, eventos = _eventos_del_tower(reporte)
    assert [e.kind for e in eventos] == ["dead_stock"]


def test_qa_rechaza_una_linea_de_demanda_cero_que_MEZCLA_ceros_con_nulos():
    """Las dos formas son legitimas; media y media no, porque entonces un consumidor
    lee un 0 medido al lado de una ausencia y no puede saber cual es cual."""
    from dataclasses import replace

    reporte, r = _rec(_archivo("SIN-VENTA", [0] * 12), "SIN-VENTA")
    assert r.policy_kind == "skipped_zero_demand"
    mezclado = replace(reporte, recommendations=[replace(r, reorder_point=None)])

    issues = qa.verify(mezclado)
    assert any("mixes measured zeros with nulls" in i for i in issues), issues


def test_qa_rechaza_una_linea_de_demanda_cero_con_un_numero_positivo():
    from dataclasses import replace

    reporte, r = _rec(_archivo("SIN-VENTA", [0] * 12), "SIN-VENTA")
    mutado = replace(reporte, recommendations=[replace(r, reorder_point=25.0)])

    issues = qa.verify(mutado)
    assert any("has a zero policy, not a positive one" in i for i in issues), issues


# ---------------------------------------------------------------------------
# 8. La cartera MIXTA: el caso que no es ni todo refusado ni nada refusado
# ---------------------------------------------------------------------------


def _mixta() -> pd.DataFrame:
    """Un SKU dimensionado y uno declinado. `NUEVO` trae UN periodo a proposito: con
    dos ya recibiria una cobertura, y entonces esta cartera no tendria ninguna
    negativa que contar y el test dejaria de medir lo que dice medir."""
    fechas = pd.date_range("2026-01-04", periods=24, freq=_PERIOD)
    filas = [{"date": f, "product_id": "LARGO", "quantity": 20.0 + (i % 4),
              "unit_cost": 40.0} for i, f in enumerate(fechas)]
    filas += [{"date": fechas[-1], "product_id": "NUEVO", "quantity": 42.0,
               "unit_cost": 10.0}]
    return pd.DataFrame(filas)


def test_el_resumen_del_agente_excluye_lo_no_dimensionado_de_la_cifra(tmp_path):
    """La mitad MIXTA: la inversion que se publica no cubre a todos los SKU, y el
    resumen tiene que decirlo o el lector supone que si."""
    from scm_agent.orchestrator import Orchestrator

    destino = tmp_path / "v.csv"
    _mixta().to_csv(destino, index=False)
    res = Orchestrator(clients_root=None).run(
        "punto de reorden y stock de seguridad por SKU",
        data_path=str(destino), job_type="inventory_optimization",
        out_dir=str(tmp_path / "out"),
    )

    assert res.status == "ok"
    assert "recommended inventory investment" in res.summary   # sigue siendo un consejo
    assert "1 SKU(s) carry no sized policy" in res.summary
    assert "excluded from that figure" in res.summary


def test_el_deck_que_lee_el_CLIENTE_dice_lo_mismo_que_el_resumen():
    """El deck arma su propio encabezado; arreglar solo el del agente dejaba al
    cliente leyendo "recommended $0 inventory investment" sobre una corrida en la
    que no se dimensiono nada."""
    from jobs.inventory_deliverable import build

    todo_refusado, _ = _rec(_archivo("NUEVO", _SERIE_SIN_MEDIA), "NUEVO")
    deck = build(todo_refusado, client="C")
    assert "No policy was sized" in deck.summary
    assert "NOT a recommendation" in deck.summary
    assert "recommended $0" not in deck.summary

    mixto = run(_mixta(), period_freq=_PERIOD)
    deck_mixto = build(mixto, client="C")
    assert "recommended $" in deck_mixto.summary
    assert "1 SKU(s) carry no sized policy" in deck_mixto.summary


# ---------------------------------------------------------------------------
# 9. La prosa tampoco puede conflacionar las dos causas
# ---------------------------------------------------------------------------
#
# Este es un defecto que introdujo el PRIMER arreglo, y lo encontro la revision
# adversarial: el resumen de "no se dimensiono nada" decia, para TODA cartera sin
# politica, "No reorder point ... was computed - this is NOT a recommendation to
# hold nothing". Sobre una cartera de demanda medida en cero eso contradice el
# propio registro (el reorder point SI se computo, y vale 0) y niega un hallazgo
# real: no hay nada que estacionar. Separar las dos causas en los CAMPOS y
# mezclarlas en la FRASE deja el mismo defecto un piso mas arriba.


def _nota(df: pd.DataFrame):
    from jobs.inventory_optimization import unsized_note
    return unsized_note(run(df, period_freq=_PERIOD))


def test_una_cartera_de_demanda_medida_en_cero_NO_se_llama_negativa():
    titular, cola = _nota(_archivo("SIN-VENTA", [0] * 12))
    assert titular is not None and cola == ""
    assert "zero measured demand" in titular
    assert "a measurement, not a refusal" in titular
    # Y NO puede decir que no se computo, porque se computo.
    assert "no reorder point" not in titular.lower()
    assert "NOT a recommendation to hold nothing" not in titular


def test_una_cartera_toda_declinada_SI_se_llama_negativa():
    titular, _ = _nota(_archivo("NUEVO", _SERIE_SIN_MEDIA))
    assert titular is not None
    assert "declined all 1 SKU(s)" in titular
    assert "too little demand history" in titular
    assert "NOT a recommendation to hold nothing" in titular


def test_las_dos_causas_juntas_se_nombran_POR_SEPARADO():
    fechas = pd.date_range("2026-01-04", periods=24, freq=_PERIOD)
    df = pd.DataFrame(
        [{"date": fechas[-1], "product_id": "NUEVO", "quantity": 42.0,
          "unit_cost": 10.0}]
        + [{"date": d, "product_id": "PARADO", "quantity": 0.0, "unit_cost": 10.0}
           for d in fechas[:12]]
    )
    titular, _ = _nota(df)
    assert titular is not None
    assert "declined 1" in titular
    assert "1 have zero measured demand" in titular
    assert "Where it declined" in titular, (
        "la salvedad tiene que alcanzar SOLO a los declinados, no a los medidos"
    )


def test_la_cola_del_caso_mixto_tambien_distingue():
    fechas = pd.date_range("2026-01-04", periods=24, freq=_PERIOD)
    df = pd.DataFrame(
        [{"date": d, "product_id": "LARGO", "quantity": 20.0 + (i % 4),
          "unit_cost": 40.0} for i, d in enumerate(fechas)]
        + [{"date": fechas[-1], "product_id": "NUEVO", "quantity": 42.0,
            "unit_cost": 10.0}]
        + [{"date": d, "product_id": "PARADO", "quantity": 0.0, "unit_cost": 10.0}
           for d in fechas[:12]]
    )
    titular, cola = _nota(df)
    assert titular is None, "con una politica dimensionada el titular normal se queda"
    assert "1 the engine declined to size" in cola
    assert "1 with zero measured demand" in cola
    assert "excluded from that figure" in cola


def test_una_cartera_TODA_cubierta_no_publica_el_encabezado_de_cero(tmp_path):
    """El mismo defecto un piso mas arriba, y el que trajo la cobertura provisional.

    El titular reemplazaba al encabezado cuando habia SKU declinados. Una cartera
    entera de coberturas no tiene ninguno -- no hubo negativa, hubo un escalon --,
    asi que caia al encabezado normal: *"recommended inventory investment $0 at 95%
    service level"* sobre una corrida en la que no se dimensiono ninguna politica y
    en la que ADEMAS hay plata que gastar. Un lector que se queda en la primera
    frase lee exactamente lo contrario de lo que paso, que es la definicion del
    defecto que este archivo persigue.

    Lo que decide es cuantos SKU quedaron DIMENSIONADOS, no cuantos quedaron sin
    nada.
    """
    from scm_agent.orchestrator import Orchestrator

    reporte, r = _rec(_archivo("NUEVO", _SERIE_CON_COBERTURA), "NUEVO")
    assert r.policy_kind == "provisional_cover"
    assert reporte.requested_investment == 0.0
    assert reporte.provisional_investment > 0.0

    titular, cola = _nota(_archivo("NUEVO", _SERIE_CON_COBERTURA))
    assert titular is not None and cola == ""
    assert "No policy was sized for any of the 1 SKU(s)" in titular
    assert "PROVISIONAL COVER" in titular
    assert "fall short about half the time" in titular
    assert "$0" not in titular, (
        "el encabezado de inversion no puede aparecer sobre una corrida sin "
        "ninguna politica dimensionada"
    )
    # Y tampoco puede decir "counted separately": no hay de que separarse.
    assert "counted separately" not in titular
    assert "more SKU(s)" not in titular

    destino = tmp_path / "v.csv"
    _archivo("NUEVO", _SERIE_CON_COBERTURA).to_csv(destino, index=False)
    res = Orchestrator(clients_root=None).run(
        "reorder point por SKU", data_path=str(destino),
        job_type="inventory_optimization", out_dir=str(tmp_path / "o"),
    )
    assert res.status == "ok"
    assert res.summary == titular


def test_una_cartera_MIXTA_con_cobertura_conserva_su_encabezado():
    """La direccion contraria, que es la que hace no-trivial al test de arriba:
    con una politica dimensionada la cifra de inversion es real, el encabezado se
    queda, y la cobertura viaja en la cola diciendo que su plata va aparte."""
    fechas = pd.date_range("2026-01-04", periods=24, freq=_PERIOD)
    df = pd.DataFrame(
        [{"date": d, "product_id": "LARGO", "quantity": 20.0 + (i % 4),
          "unit_cost": 40.0} for i, d in enumerate(fechas)]
        + [{"date": fechas[-i - 1], "product_id": "CUBIERTO", "quantity": 42.0,
            "unit_cost": 10.0} for i in range(5)]
    )
    titular, cola = _nota(df)

    assert titular is None
    assert "1 more SKU(s) got a PROVISIONAL COVER" in cola
    assert "counted separately" in cola


def test_el_agente_y_el_deck_dicen_LO_MISMO(tmp_path):
    """Las dos superficies escribian su propia frase y por eso una quedo sin
    arreglar. Ahora salen de la misma funcion, y esto lo fija."""
    from jobs.inventory_deliverable import build
    from scm_agent.orchestrator import Orchestrator

    destino = tmp_path / "v.csv"
    _archivo("NUEVO", _SERIE_SIN_MEDIA).to_csv(destino, index=False)
    res = Orchestrator(clients_root=None).run(
        "reorder point por SKU", data_path=str(destino),
        job_type="inventory_optimization", out_dir=str(tmp_path / "o"),
    )
    reporte, _ = _rec(_archivo("NUEVO", _SERIE_SIN_MEDIA), "NUEVO")
    deck = build(reporte, client="C")

    titular, _ = _nota(_archivo("NUEVO", _SERIE_SIN_MEDIA))
    assert res.summary == titular
    assert deck.summary.startswith(titular)
