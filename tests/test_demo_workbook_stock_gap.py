"""`escaneo_stock.xlsx` y `brecha_en_plata.xlsx`: los otros dos libros de /demo.

Los dos comparten primitivas con `plazo_real.xlsx` (webapp/demo_workbook_base.py),
asi que lo que se prueba aca es lo que NO comparten: la aritmetica propia de cada
uno, que es donde un entregable de plata miente sin que nadie lo note.

Los dos defectos que estos tests existen para impedir, ambos verificados sobre
datos reales antes de escribirse:

* **stock** -- `excess_value` es solo el excedente sobre el objetivo, mientras que
  `on_hand_value` es todo. La resta `total - dead - excess` sobreestima el valor
  sano en exactamente la parte de los excedentes que si cubre demanda, y poner
  plata al lado de un conteo compara poblaciones distintas en la misma fila.
* **brecha** -- `recurring_saving` se calcula sobre la caja liberable y
  `net_carrying_change` sobre el neto, que normalmente es POSITIVO. Sin el
  termino del medio en la hoja, "ahorro" y "cambio +" se leen como una
  contradiccion.
"""
from __future__ import annotations

import math

import pytest

pytest.importorskip("openpyxl")
pytest.importorskip("pandas")

import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from webapp.demo_gap_value import HOLDING_RATE, run_gap_value  # noqa: E402
from webapp.demo_gap_workbook import write_gap_value_workbook  # noqa: E402
from webapp.demo_scan import run_demo_scan  # noqa: E402
from webapp.demo_stock_workbook import (  # noqa: E402
    ABC_THRESHOLDS,
    _Derivadas,
    write_stock_workbook,
)
from webapp.demo_workbook_base import NO_MEDIDO, rows_below  # noqa: E402

STOCK_CSV = "data/sample_stock_snapshot.csv"
RECEIPTS_CSV = "data/sample_receipts_portfolio.csv"
DEMAND_CSV = "data/sample_demand_portfolio.csv"


@pytest.fixture()
def stock():
    return run_demo_scan(pd.read_csv(STOCK_CSV))


@pytest.fixture()
def libro_stock(tmp_path, stock):
    return load_workbook(
        write_stock_workbook(stock, tmp_path / "s.xlsx", dataset_label="stock.csv", ts="T")
    )


@pytest.fixture()
def gap():
    return run_gap_value(pd.read_csv(RECEIPTS_CSV), pd.read_csv(DEMAND_CSV))


@pytest.fixture()
def libro_gap(tmp_path, gap):
    return load_workbook(
        write_gap_value_workbook(gap, tmp_path / "g.xlsx", dataset_label="dos.csv", ts="T")
    )


# --- escaneo_stock: las dos poblaciones -------------------------------------


def test_el_valor_por_clase_cierra_en_el_total(stock):
    """La columna de plata suma el total del inventario. La otra no, y por eso van separadas."""
    d = _Derivadas(stock)
    suma = d.v_healthy + d.v_excess_total + d.v_sin + d.v_idle
    assert suma == pytest.approx(stock.eo.total_on_hand_value)


def test_la_formula_ingenua_sobreestima_el_valor_sano(stock):
    """`total - dead_value - excess_value` absorbe lo que los excedentes SI cubren.

    No es un matiz: sobre el archivo de ejemplo son 1.350 de mas atribuidos a una
    fila que dice "sano" al lado de un conteo de SKU sanos.
    """
    d = _Derivadas(stock)
    eo = stock.eo
    ingenua = eo.total_on_hand_value - eo.dead_value - eo.excess_value
    assert ingenua > d.v_healthy
    assert ingenua - d.v_healthy == pytest.approx(d.v_excess_within)


def test_la_columna_de_riesgo_cierra_en_la_plata_detenida(stock):
    d = _Derivadas(stock)
    assert stock.eo.excess_value + d.r_sin + d.r_idle == pytest.approx(stock.eo.eo_value)


def test_las_dos_columnas_del_resumen_tienen_su_propio_total(libro_stock, stock):
    ws = libro_stock["Resumen"]
    fila_total = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Total")
    assert ws.cell(fila_total, 2).value == pytest.approx(stock.eo.total_on_hand_value)
    assert ws.cell(fila_total, 4).value == pytest.approx(stock.eo.eo_value)
    # Y las filas de arriba suman cada columna por separado.
    valor = sum(ws.cell(r, 2).value for r in range(12, fila_total))
    riesgo = sum(ws.cell(r, 4).value for r in range(12, fila_total))
    assert valor == pytest.approx(stock.eo.total_on_hand_value)
    assert riesgo == pytest.approx(stock.eo.eo_value)


def test_el_balde_muerto_se_parte_por_la_rama_que_el_dato_prueba(stock):
    """`classify_sku` marca DEAD por demanda<=0 O por antiguedad.

    Decir "sin movimiento hace N dias" sobre todo el balde afirma algo que el
    archivo puede no haber traido nunca: la fecha de ultima venta es opcional.
    """
    d = _Derivadas(stock)
    assert len(d.dead_sin) + len(d.dead_idle) == stock.eo.n_dead
    # La cobertura es infinita si y solo si la demanda diaria es <= 0.
    assert all(not math.isfinite(ln.days_of_cover) for ln in d.dead_sin)
    assert all(math.isfinite(ln.days_of_cover) for ln in d.dead_idle)


def test_cobertura_infinita_se_escribe_no_medido_y_no_la_palabra_inf(libro_stock):
    ws = libro_stock["Plata en riesgo"]
    valores = [ws.cell(r, 7).value for r in range(5, ws.max_row + 1)]
    assert NO_MEDIDO in valores, "algun SKU sin demanda tiene que decirlo"
    assert not any(str(v).lower() == "inf" for v in valores)


def test_el_grafico_de_cobertura_no_puede_recibir_un_no_finito(libro_stock):
    """Se excluyen por construccion, no se renderizan como texto en un eje numerico."""
    ws = libro_stock["Cobertura"]
    for r in range(6, ws.max_row + 1):
        v = ws.cell(r, 2).value
        if v is None:
            continue
        assert isinstance(v, (int, float)) and math.isfinite(v) and v >= 0


def test_gmroi_no_aparece_en_ninguna_celda(libro_stock, stock):
    """`gmroi` no es None aqui: es 0.0, porque el margen se fija en cero.

    Un None no puede escribirse por accidente como numero; un 0,0 si, y
    "GMROI 0,00" afirma que el inventario del cliente no rinde nada.
    """
    assert stock.fin.gmroi == 0.0
    texto = " ".join(
        str(c.value).lower() for ws in libro_stock for row in ws.iter_rows()
        for c in row if c.value is not None
    )
    assert "gmroi" in texto, "tiene que estar NOMBRADO en 'Que NO se midio'"
    hojas_datos = [ws for ws in libro_stock if ws.title != "Como leer esto"]
    for ws in hojas_datos:
        for row in ws.iter_rows():
            for c in row:
                assert "gmroi" not in str(c.value).lower()


def test_los_cortes_abc_siguen_siendo_los_del_motor():
    """Si el motor cambia su default, la hoja deja de describir lo que hizo."""
    import inspect

    from jobs import abc_xyz_job

    assert inspect.signature(abc_xyz_job.run).parameters["abc_thresholds"].default == ABC_THRESHOLDS


# --- brecha_en_plata: los signos y el embudo --------------------------------


def test_las_dos_monedas_se_reportan_positivas_y_el_neto_firmado(gap):
    assert gap.cash_release >= 0 and gap.exposure >= 0
    assert gap.exposure - gap.cash_release == pytest.approx(gap.net)


def test_las_tres_filas_anuales_cierran_en_la_hoja(libro_gap, gap):
    """El termino del medio no sale del motor. Sin el, la hoja se contradice."""
    ws = libro_gap["Resumen"]
    fila = next(r for r in range(1, ws.max_row + 1)
                if ws.cell(r, 1).value == "Ahorro, solo sobre la caja que se libera")
    ahorro = ws.cell(fila, 2).value
    costo = ws.cell(fila + 1, 2).value
    neto = ws.cell(fila + 2, 2).value
    assert ahorro == pytest.approx(gap.recurring_saving)
    assert costo == pytest.approx(gap.exposure * HOLDING_RATE)
    assert costo - ahorro == pytest.approx(neto)
    assert neto == pytest.approx(gap.net_carrying_change)


def test_ningun_grafico_de_plata_tiene_una_sola_serie(libro_gap):
    """Reportar solo la caja liberable es 'reduce inventario', la frase prohibida."""
    for titulo in ("Resumen",):
        for ch in libro_gap[titulo]._charts:
            assert len(ch.series) == 2, f"{ch.title} tiene {len(ch.series)} serie(s)"


def test_ningun_bloque_de_grafico_lleva_un_negativo(libro_gap):
    """`_bar` fija el eje en cero: una barra bajo cero desaparece sin error."""
    for ws in libro_gap:
        for row in ws.iter_rows():
            for c in row:
                if c.column >= 13 and isinstance(c.value, (int, float)):
                    assert c.value >= 0, f"{ws.title}!{c.coordinate} = {c.value}"


def test_el_embudo_declara_el_denominador(libro_gap, gap):
    ws = libro_gap["Calidad de los datos"]
    etiquetas = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(1, ws.max_row + 1)}
    assert etiquetas["1. SKU distintos en tu archivo de movimientos"] == gap.n_demand_skus
    assert etiquetas["3. SKU valorizados: los que cumplen las tres a la vez"] == len(gap.gaps)
    # Y el denominador viaja tambien a la primera pantalla.
    resumen = libro_gap["Resumen"]
    assert str(gap.n_demand_skus) in str(resumen["C4"].value)


def test_los_tres_requisitos_no_se_presentan_como_encadenados(libro_gap):
    """Declarado, observado y variacion son pruebas PARALELAS.

    Encadenarlas afirma una relacion de subconjunto que no se sostiene, y un
    embudo dibujado con ellas puede ir hacia arriba.
    """
    ws = libro_gap["Calidad de los datos"]
    texto = " ".join(str(ws.cell(r, 1).value) for r in range(1, ws.max_row + 1))
    assert "EN PARALELO" in texto


def test_el_embudo_del_grafico_es_monotono(libro_gap):
    ws = libro_gap["Calidad de los datos"]
    serie = [ws.cell(r, 7).value for r in range(3, 6)]
    assert serie == sorted(serie, reverse=True), serie


def test_la_reconciliacion_de_recepciones_cuadra(libro_gap):
    ws = libro_gap["Calidad de los datos"]
    v = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(1, ws.max_row + 1)}
    assert v["Total fuera de los numeros"] + v["Recepciones usadas"] == v["Filas del archivo"]


def test_calidad_nombra_como_lo_leyo_kern(libro_gap):
    """El mapeo vive en Calidad de los datos, no en los hallazgos: no es un defecto."""
    ws = libro_gap["Calidad de los datos"]
    etiquetas = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
    assert "Como lo leyo Kern" in etiquetas
    assert "Donde miro" in etiquetas
    assert "Mapeo de columnas" in etiquetas
    mapeo = next(
        ws.cell(r, 2).value
        for r in range(1, ws.max_row + 1)
        if ws.cell(r, 1).value == "Mapeo de columnas"
    )
    assert "product_id <- 'product_id'" in mapeo
    assert "date <- 'date'" in mapeo
    lugar = next(
        ws.cell(r, 2).value
        for r in range(1, ws.max_row + 1)
        if ws.cell(r, 1).value == "Donde miro"
    )
    assert "CSV" in lugar

    """`_days_to_periods` convierte adentro; exponer un periodo es el error de x7."""
    ws = libro_gap["Como leer esto"]
    texto = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "ninguna cifra de este archivo esta en semanas" in texto.lower()


# --- los dos, contra el contrato comun --------------------------------------


@pytest.mark.parametrize("cual", ["stock", "gap"])
def test_con_qa_fallado_no_hay_libro(tmp_path, cual, stock, gap):
    import dataclasses

    obj = stock if cual == "stock" else gap
    roto = dataclasses.replace(obj, qa_issues=("forzado",), findings=())
    escribir = write_stock_workbook if cual == "stock" else write_gap_value_workbook
    with pytest.raises(ValueError, match="QA fallado"):
        escribir(roto, tmp_path / "x.xlsx", dataset_label="d", ts="T")


@pytest.mark.parametrize("cual", ["stock", "gap"])
def test_dos_corridas_dan_el_mismo_libro(tmp_path, cual, stock, gap):
    obj = stock if cual == "stock" else gap
    escribir = write_stock_workbook if cual == "stock" else write_gap_value_workbook
    kw = {"dataset_label": "d.csv", "ts": "2026-08-08T00:00:00Z"}

    def celdas(path):
        wb = load_workbook(path)
        return [(ws.title, c.coordinate, c.value)
                for ws in wb for row in ws.iter_rows() for c in row]

    assert celdas(escribir(obj, tmp_path / "a.xlsx", **kw)) == celdas(
        escribir(obj, tmp_path / "b.xlsx", **kw)
    )


@pytest.mark.parametrize("cual", ["stock", "gap"])
def test_inyeccion_de_formula_neutralizada(tmp_path, cual):
    """El SKU y el proveedor salen de un archivo subido por un anonimo."""
    payload = "=cmd|'/c calc'!A1"
    if cual == "stock":
        df = pd.read_csv(STOCK_CSV)
        df.loc[0, df.columns[0]] = payload
        obj, escribir = run_demo_scan(df), write_stock_workbook
    else:
        rec = pd.read_csv(RECEIPTS_CSV)
        rec["supplier"] = rec["supplier"].replace({rec["supplier"].iloc[0]: payload})
        obj, escribir = run_gap_value(rec, pd.read_csv(DEMAND_CSV)), write_gap_value_workbook
    if not obj.ok:
        pytest.skip("el archivo doctorado no pasa el QA; la inyeccion se prueba en el otro caso")

    wb = load_workbook(escribir(obj, tmp_path / "x.xlsx", dataset_label="d", ts="T"))
    encontrada = False
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if c.value and "cmd" in str(c.value):
                    encontrada = True
                    assert str(c.value).startswith("'")
                    assert c.data_type != "f"
    assert encontrada


@pytest.mark.parametrize("cual", ["stock", "gap"])
def test_los_graficos_son_nativos_y_sobreviven_al_autofiltro(tmp_path, cual, stock, gap):
    import zipfile

    obj = stock if cual == "stock" else gap
    escribir = write_stock_workbook if cual == "stock" else write_gap_value_workbook
    path = escribir(obj, tmp_path / "x.xlsx", dataset_label="d", ts="T")
    with zipfile.ZipFile(path) as z:
        charts = [n for n in z.namelist() if n.startswith("xl/charts/chart")]
        assert charts, "sin graficos no es el entregable"
        for name in charts:
            assert 'plotVisOnly val="0"' in z.read(name).decode(), name
    assert not any(getattr(ws, "_images", []) for ws in load_workbook(path))


def test_rows_below_reproduce_el_calculo_a_mano_del_primer_libro():
    """La constante de 0,529 cm por fila es la que ya usaba `plazo_real`."""
    assert rows_below(9.5) == 20
