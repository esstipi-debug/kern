"""La planilla real del cliente: totales, celdas combinadas y filas de relleno.

``jobs/intake.py`` ya resuelve dos defectos de un libro real -- encuentra la hoja
y la fila de encabezado por evidencia, y parsea separadores de miles LATAM y
anglosajones. Quedaban tres que nadie miraba, y son peores que los resueltos
porque **no fallan**: producen un numero confiado y equivocado.

1. **La fila de totales.** Toda planilla hecha a mano termina en ``TOTAL``. Hoy
   entra como una fila de demanda mas, con la suma entera del periodo cargada a
   un solo dia. Eso no rompe el pronostico: lo corre. Y el error viaja hasta el
   stock de seguridad, que se dimensiona sobre una dispersion inventada.
2. **Las celdas combinadas.** Escribir el SKU una vez y combinar cinco filas es
   como se ve una planilla ordenada. ``pd.read_excel`` deja cuatro vacias, y esas
   cuatro filas de venta real se caen del intake por "sin producto".
3. **Las filas y columnas de relleno.** Separadores en blanco entre bloques y
   columnas ``Unnamed`` sin nada adentro.

La regla que gobierna el modulo entero: **cada fila descartada se cuenta y se
declara**. Un limpiador silencioso es peor que ninguno -- convierte "tu planilla
tenia 12 filas de basura" en "tu demanda es 12 filas mas chica", y nadie lo ve.
"""

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

from src.excel_hygiene import clean_frame, is_total_label, read_clean_table


def _libro(tmp_path: Path, filas: list[list], *, combinar: str | None = None) -> Path:
    ruta = tmp_path / "cliente.xlsx"
    wb = Workbook()
    ws = wb.active
    for fila in filas:
        ws.append(fila)
    if combinar:
        ws.merge_cells(combinar)
    wb.save(ruta)
    return ruta


_CABECERA = ["fecha", "sku", "cantidad"]


# -- la fila de totales --------------------------------------------------------


def test_reconoce_las_etiquetas_de_total_en_los_dos_idiomas():
    for etiqueta in ("TOTAL", "Total", "total general", "Subtotal", "SUMA", "Grand Total", "totales"):
        assert is_total_label(etiqueta), etiqueta


def test_no_confunde_un_producto_que_se_llama_total():
    """``TOTAL-500`` es un SKU. La palabra sola es una etiqueta; dentro de un codigo, no."""
    for sku in ("TOTAL-500", "SUBTOTAL-X1", "Kit Total Care", "totalizador"):
        assert not is_total_label(sku)


def test_una_fila_de_total_no_entra_como_demanda():
    frame = pd.DataFrame(
        {
            "fecha": ["2026-01-05", "2026-01-06", ""],
            "sku": ["SKU-1", "SKU-2", "TOTAL"],
            "cantidad": [10.0, 5.0, 15.0],
        }
    )

    limpio = clean_frame(frame, key_column="sku")

    assert list(limpio.frame["sku"]) == ["SKU-1", "SKU-2"]
    assert limpio.dropped["fila_de_total"] == 1


def test_detecta_la_fila_de_total_aunque_no_diga_total():
    """El caso peligroso: la etiqueta esta en blanco y solo cuadra la suma."""
    frame = pd.DataFrame(
        {
            "fecha": ["2026-01-05", "2026-01-06", ""],
            "sku": ["SKU-1", "SKU-2", ""],
            "cantidad": [10.0, 5.0, 15.0],
        }
    )

    limpio = clean_frame(frame, key_column="sku", numeric_columns=("cantidad",))

    assert len(limpio.frame) == 2
    assert limpio.dropped["fila_de_total"] == 1


def test_una_coincidencia_de_suma_en_el_medio_no_es_un_total():
    """Solo la ULTIMA fila del bloque puede ser el total. Una del medio es venta."""
    frame = pd.DataFrame(
        {
            "fecha": ["2026-01-05", "2026-01-06", "2026-01-07"],
            "sku": ["SKU-1", "SKU-2", "SKU-3"],
            "cantidad": [10.0, 10.0, 4.0],
        }
    )

    limpio = clean_frame(frame, key_column="sku", numeric_columns=("cantidad",))

    assert len(limpio.frame) == 3
    assert limpio.dropped["fila_de_total"] == 0


def test_el_total_cuadra_la_cantidad_y_deja_el_precio_en_blanco():
    """La forma REAL de una fila de total, y la que se escapaba.

    Nadie suma precios unitarios: quien cierra una planilla a mano totaliza la
    cantidad y el importe, y deja la columna de precio vacia. Exigir que
    **todas** las columnas numericas cuadren dejaba pasar justo ese caso -- el
    mas comun de todos -- porque el precio en blanco abortaba la comprobacion
    antes de llegar a la cantidad, que si cuadraba.
    """
    frame = pd.DataFrame(
        {
            "fecha": ["2026-01-05", "2026-01-06", ""],
            "sku": ["SKU-1", "SKU-2", ""],
            "cantidad": [10.0, 5.0, 15.0],
            "precio": [3.0, 7.0, None],
        }
    )

    limpio = clean_frame(frame, key_column="sku")

    assert limpio.dropped["fila_de_total"] == 1
    assert len(limpio.frame) == 2


def test_una_columna_numerica_que_CONTRADICE_la_suma_salva_a_la_fila():
    """Saltear los blancos no es saltear los desacuerdos.

    Si una columna numerica trae un numero y ese numero NO es la suma, la fila
    no es un total: es demanda que casualmente cuadra en otra columna. Este es
    el limite de la regla de arriba, y sin el, aflojarla descartaria ventas.
    """
    frame = pd.DataFrame(
        {
            "fecha": ["2026-01-05", "2026-01-06", ""],
            "sku": ["SKU-1", "SKU-2", ""],
            "cantidad": [10.0, 5.0, 15.0],
            "importe": [30.0, 35.0, 999.0],
        }
    )

    limpio = clean_frame(frame, key_column="sku")

    assert limpio.dropped["fila_de_total"] == 0
    assert len(limpio.frame) == 3


def test_una_ultima_fila_sin_ningun_numero_no_es_un_total():
    """Con todas las numericas en blanco no hay NADA que cuadre.

    Sin esta guarda, saltear los blancos volveria la comprobacion vacuamente
    verdadera y se llevaria puesta cualquier fila final incompleta.
    """
    frame = pd.DataFrame(
        {
            "fecha": ["2026-01-05", "2026-01-06", "2026-01-07"],
            "sku": ["SKU-1", "SKU-2", ""],
            "cantidad": [10.0, 5.0, None],
            "precio": [3.0, 7.0, None],
        }
    )

    limpio = clean_frame(frame, key_column="sku")

    assert limpio.dropped["fila_de_total"] == 0
    assert len(limpio.frame) == 3


def test_una_venta_que_casualmente_cuadra_con_la_suma_pero_tiene_sku_se_conserva():
    """Sin etiqueta de total Y con producto y fecha propios, es una fila de datos."""
    frame = pd.DataFrame(
        {
            "fecha": ["2026-01-05", "2026-01-06", "2026-01-07"],
            "sku": ["SKU-1", "SKU-2", "SKU-3"],
            "cantidad": [10.0, 5.0, 15.0],
        }
    )

    limpio = clean_frame(frame, key_column="sku", numeric_columns=("cantidad",))

    assert len(limpio.frame) == 3


# -- celdas combinadas ---------------------------------------------------------


def test_una_celda_combinada_reparte_su_valor_a_las_filas_que_cubre(tmp_path):
    ruta = _libro(
        tmp_path,
        [
            _CABECERA,
            ["2026-01-05", "SKU-1", 10],
            ["2026-01-06", None, 4],
            ["2026-01-07", None, 6],
            ["2026-01-08", "SKU-2", 3],
        ],
        combinar="B2:B4",
    )

    lectura = read_clean_table(ruta, key_column="sku")

    assert list(lectura.frame["sku"]) == ["SKU-1", "SKU-1", "SKU-1", "SKU-2"]
    assert lectura.filled["sku"] == 2


def test_una_celda_vacia_que_NO_estaba_combinada_no_se_rellena(tmp_path):
    """La diferencia entre leer el libro y adivinar: el hueco real queda hueco."""
    ruta = _libro(
        tmp_path,
        [
            _CABECERA,
            ["2026-01-05", "SKU-1", 10],
            ["2026-01-06", None, 4],
        ],
    )

    lectura = read_clean_table(ruta, key_column="sku")

    # pandas devuelve NaN por la celda vacia: quedo vacia, que es el punto
    assert pd.isna(lectura.frame["sku"].tolist()[1])
    assert lectura.filled.get("sku", 0) == 0


# -- relleno -------------------------------------------------------------------


def test_las_filas_enteramente_vacias_se_descartan_y_se_cuentan():
    frame = pd.DataFrame(
        {
            "fecha": ["2026-01-05", None, "2026-01-07"],
            "sku": ["SKU-1", None, "SKU-2"],
            "cantidad": [10.0, None, 3.0],
        }
    )

    limpio = clean_frame(frame, key_column="sku")

    assert len(limpio.frame) == 2
    assert limpio.dropped["fila_vacia"] == 1


def test_una_columna_sin_nombre_y_sin_datos_se_va():
    frame = pd.DataFrame(
        {"fecha": ["2026-01-05"], "sku": ["SKU-1"], "cantidad": [10.0], "Unnamed: 3": [None]}
    )

    limpio = clean_frame(frame, key_column="sku")

    assert "Unnamed: 3" not in limpio.frame.columns
    assert limpio.dropped_columns == ("Unnamed: 3",)


def test_una_columna_sin_nombre_PERO_con_datos_se_conserva():
    """Un encabezado faltante no autoriza a tirar la columna: puede ser el precio."""
    frame = pd.DataFrame(
        {"fecha": ["2026-01-05"], "sku": ["SKU-1"], "cantidad": [10.0], "Unnamed: 3": [99.0]}
    )

    limpio = clean_frame(frame, key_column="sku")

    assert "Unnamed: 3" in limpio.frame.columns
    assert limpio.dropped_columns == ()


# -- la regla que gobierna -----------------------------------------------------


def test_nada_se_descarta_en_silencio():
    frame = pd.DataFrame(
        {
            "fecha": ["2026-01-05", None, "", "2026-01-08"],
            "sku": ["SKU-1", None, "TOTAL", "SKU-2"],
            "cantidad": [10.0, None, 10.0, 7.0],
            "Unnamed: 4": [None, None, None, None],
        }
    )

    limpio = clean_frame(frame, key_column="sku")

    # 4 filas entran, 2 salen, y las 2 que faltan estan explicadas una por una
    assert len(limpio.frame) == 2
    assert sum(limpio.dropped.values()) == 2
    assert limpio.dropped["fila_vacia"] == 1
    assert limpio.dropped["fila_de_total"] == 1
    assert limpio.dropped_columns == ("Unnamed: 4",)
    assert limpio.rows_in == 4


def test_un_frame_limpio_pasa_intacto_y_lo_dice():
    frame = pd.DataFrame(
        {"fecha": ["2026-01-05", "2026-01-06"], "sku": ["SKU-1", "SKU-2"], "cantidad": [10.0, 5.0]}
    )

    limpio = clean_frame(frame, key_column="sku")

    assert len(limpio.frame) == 2
    assert sum(limpio.dropped.values()) == 0
    assert limpio.clean is True


def test_un_libro_que_es_puro_relleno_no_devuelve_un_frame_vacio_en_silencio(tmp_path):
    ruta = _libro(tmp_path, [_CABECERA, [None, None, None], ["", "TOTAL", 0]])

    lectura = read_clean_table(ruta, key_column="sku")

    assert lectura.frame.empty
    assert lectura.clean is False
    assert sum(lectura.dropped.values()) == 2


# -- de punta a punta ----------------------------------------------------------


def test_la_planilla_de_un_cliente_de_verdad(tmp_path):
    """Título arriba, celda combinada, separador en blanco y total al final."""
    ruta = tmp_path / "ventas.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Reporte de Ventas 2026"])
    ws.append([])
    ws.append(_CABECERA)
    ws.append(["2026-01-05", "SKU-1", 10])
    ws.append(["2026-01-06", None, 4])
    ws.append([])
    ws.append(["2026-01-07", "SKU-2", 6])
    ws.append(["", "TOTAL", 20])
    ws.merge_cells("B4:B5")
    wb.save(ruta)

    lectura = read_clean_table(ruta, key_column="sku", header_row=2)

    assert list(lectura.frame["sku"]) == ["SKU-1", "SKU-1", "SKU-2"]
    assert lectura.frame["cantidad"].sum() == pytest.approx(20.0)
    assert lectura.dropped["fila_vacia"] == 1
    assert lectura.dropped["fila_de_total"] == 1
    assert lectura.filled["sku"] == 1


# -- conectado al unico punto de entrada de archivos de cliente ----------------


def test_load_table_limpia_el_excel_del_cliente_y_deja_el_informe(tmp_path):
    """``load_table`` lo usan el intake, pricing y la subida del webapp: uno solo."""
    from jobs.intake import HYGIENE_ATTR, load_table

    ruta = tmp_path / "cliente.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Reporte 2026"])
    ws.append(["fecha", "sku", "cantidad"])
    ws.append(["2026-01-05", "SKU-1", 10])
    ws.append(["2026-01-06", None, 4])
    ws.append([])
    ws.append(["", "TOTAL", 14])
    ws.merge_cells("B3:B4")
    wb.save(ruta)

    frame = load_table(ruta)

    assert list(frame["sku"]) == ["SKU-1", "SKU-1"]
    informe = frame.attrs[HYGIENE_ATTR]
    assert informe.dropped["fila_de_total"] == 1
    assert informe.dropped["fila_vacia"] == 1
    assert informe.filled["sku"] == 1
    assert informe.clean is False


def test_un_csv_no_pasa_por_la_higiene_de_excel(tmp_path):
    """El camino CSV queda igual: ahi un titulo arriba ya rompe con un error claro."""
    from jobs.intake import HYGIENE_ATTR, load_table

    ruta = tmp_path / "datos.csv"
    ruta.write_text("fecha,sku,cantidad\n2026-01-05,SKU-1,10\n", encoding="utf-8")

    frame = load_table(ruta)

    assert len(frame) == 1
    assert HYGIENE_ATTR not in frame.attrs


def test_un_excel_sin_columna_de_producto_reconocible_pasa_crudo(tmp_path):
    """Sin saber cual es el SKU no hay donde buscar el total: mejor no adivinar."""
    from jobs.intake import HYGIENE_ATTR, load_table

    ruta = tmp_path / "raro.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["columna_a", "columna_b"])
    ws.append([1, 2])
    wb.save(ruta)

    frame = load_table(ruta)

    assert HYGIENE_ATTR not in frame.attrs
    assert len(frame) == 1
