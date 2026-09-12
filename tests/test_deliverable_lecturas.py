"""Las lecturas salen en el entregable RIO ABAJO del gate de citas, y no tocan nada mas.

Tres cosas se fijan aca, y las tres son de producto antes que de codigo:

1. **Nada de lo que ya se emite cambia.** La seccion nueva es opcional y condicional, asi
   que un entregable sin lecturas tiene que renderizar BYTE A BYTE lo mismo que antes. No
   se promete: los dos golden de abajo (en/es) se capturaron con ``to_markdown()`` sobre
   el commit anterior al cambio y se comparan con ``==``, no con ``in``.

2. **Un pasaje no es una cita.** ``readings`` y ``citations`` son dos campos, dos
   secciones y dos listas: un pasaje del indice de lectura no tiene ``node_id`` y
   ``filter_citations`` lo rechazaria con razon. Poblar uno no puede mover el otro.

3. **La pagina es la FISICA del PDF.** ``page_label`` viaja TEXTUAL ("PDF p.83-84") y la
   intro lo dice en la misma linea en que dice que no son citas. El folio impreso del
   libro arranca despues (Vandeput: 26 paginas de corrimiento; Liu: 12), asi que un
   cliente que lee "p.83" y abre el libro impreso en la 83 no encuentra nada -- y una
   cita que el propio cliente falsea es peor que no citar.
"""
from __future__ import annotations

import pytest
from openpyxl import load_workbook

from nosilo_core.guided import ExecutionOption
from src import i18n
from src.connectors import lectura
from src.connectors.lectura import MAX_EXCERPT_CHARS, Reading
from src.deliverable import Branding, DataSource, Deliverable, Finding, Kpi

# ---- el entregable de referencia, con TODAS las secciones pobladas -------------------


def _poblado(lang: str = "en", **extra) -> Deliverable:
    campos = dict(
        title="Inventory Optimization",
        client="Acme DTC",
        summary="Analyzed 320 SKUs; recommended $48,000 inventory investment at 95% service level.",
        findings=(
            Finding("12 SKUs at stockout risk", "Reorder now to hold 95% service", impact="$9,200 protected"),
            Finding("18 SKUs overstocked", "Excess vs target", impact="$14,500 cash to release"),
        ),
        kpis=(
            Kpi("Inventory turns", "6.2", "8-12", "Cash efficiency of stock"),
            Kpi("Fill rate", "94%", "95%+", "Service the customer sees"),
        ),
        data_sources=(
            DataSource("Demand history", "shopify_orders.csv", "weekly"),
            DataSource("On-hand stock", "Cin7 export", "daily"),
        ),
        recommendations=("Issue POs for the 12 at-risk SKUs", "Mark down the 18 overstocked SKUs"),
        options=(
            ExecutionOption(label="Do X", summary="best move", recommended=True,
                            action="apply X", tradeoffs="fast"),
            ExecutionOption(label="Do Y", summary="the alternative", action="apply Y", tradeoffs="slow"),
        ),
        citations=("Reorder Point - vandeput.pdf Ch.4 -> src/eoq.py:L12",),
        confidence=0.9,
        residual="Approve the drafted POs and confirm supplier lead times before issuing.",
        prepared="2026-06-23",
        lang=lang,
        branding=Branding(name="Acme Consulting", logo_url="https://acme.example/logo.png"),
    )
    campos.update(extra)  # `extra` PISA, no duplica: varios tests apagan `citations`
    return Deliverable(**campos)


_LECTURAS = (
    Reading(source="vandeput-inventory-optimization",
            page_label="PDF p.83-84",
            excerpt="Safety stock sizing depends on the risk period, not on the lead time alone."),
    Reading(source="liu-supply-chain-analytics",
            page_label="PDF p.12",
            excerpt="Analytics without a decision attached is reporting."),
)


# ---- 1. byte a byte: lo que ya se emitia no se movio ---------------------------------

#: Capturado con ``to_markdown()`` ANTES de agregar el campo ``readings`` (BASE
#: d210995). Se compara con ``==`` a proposito: un ``in`` no ve un espacio de mas.
GOLDEN_EN = """![Acme Consulting](https://acme.example/logo.png)

# Inventory Optimization - Acme DTC
*Prepared 2026-06-23*

## Executive summary

Analyzed 320 SKUs; recommended $48,000 inventory investment at 95% service level.

## Key findings

- **12 SKUs at stockout risk** - Reorder now to hold 95% service _(impact: $9,200 protected)_
- **18 SKUs overstocked** - Excess vs target _(impact: $14,500 cash to release)_

## Recommendations

1. Issue POs for the 12 at-risk SKUs
2. Mark down the 18 overstocked SKUs

## Options to act

Ranked, executable options - the recommended default is marked:

1. **Do X** _(recommended)_ - best move
   - Action: apply X
   - Trade-off: fast
2. **Do Y** - the alternative
   - Action: apply Y
   - Trade-off: slow

## KPIs

| KPI | Value | Target | Why it matters |
|---|---|---|---|
| Inventory turns | 6.2 | 8-12 | Cash efficiency of stock |
| Fill rate | 94% | 95%+ | Service the customer sees |

## Data sources

Every figure above is traceable to its origin:

| Metric | Source | Refresh |
|---|---|---|
| Demand history | shopify_orders.csv | weekly |
| On-hand stock | Cin7 export | daily |

## Methodology & grounding

Grounded in the L3 supply-chain knowledge base:

- Reorder Point - vandeput.pdf Ch.4 -> src/eoq.py:L12

## Coverage & handoff

Confidence: **90%**.
Approve the drafted POs and confirm supplier lead times before issuing.

---
_Prepared by Acme Consulting_
"""

GOLDEN_ES = """![Acme Consulting](https://acme.example/logo.png)

# Inventory Optimization - Acme DTC
*Preparado 2026-06-23*

## Resumen ejecutivo

Analyzed 320 SKUs; recommended $48,000 inventory investment at 95% service level.

## Hallazgos principales

- **12 SKUs at stockout risk** - Reorder now to hold 95% service _(impact: $9,200 protected)_
- **18 SKUs overstocked** - Excess vs target _(impact: $14,500 cash to release)_

## Recomendaciones

1. Issue POs for the 12 at-risk SKUs
2. Mark down the 18 overstocked SKUs

## Opciones para actuar

Opciones priorizadas y ejecutables - la recomendada esta marcada:

1. **Do X** _(recomendada)_ - best move
   - Accion: apply X
   - Contrapartida: fast
2. **Do Y** - the alternative
   - Accion: apply Y
   - Contrapartida: slow

## KPIs

| KPI | Valor | Objetivo | Por que importa |
|---|---|---|---|
| Inventory turns | 6.2 | 8-12 | Cash efficiency of stock |
| Fill rate | 94% | 95%+ | Service the customer sees |

## Fuentes de datos

Cada cifra de arriba es trazable a su origen:

| Metrica | Fuente | Actualizacion |
|---|---|---|
| Demand history | shopify_orders.csv | weekly |
| On-hand stock | Cin7 export | daily |

## Metodologia y fundamento

Fundamentado en la base de conocimiento L3 de cadena de suministro:

- Reorder Point - vandeput.pdf Ch.4 -> src/eoq.py:L12

## Cobertura y transferencia

Confianza: **90%**.
Approve the drafted POs and confirm supplier lead times before issuing.

---
_Preparado por Acme Consulting_
"""


def test_un_entregable_sin_lecturas_renderiza_byte_a_byte_lo_mismo_que_antes_en():
    assert _poblado("en").to_markdown() == GOLDEN_EN


def test_un_entregable_sin_lecturas_renderiza_byte_a_byte_lo_mismo_que_antes_es():
    assert _poblado("es").to_markdown() == GOLDEN_ES


def test_readings_por_defecto_esta_vacio():
    assert Deliverable(title="T", client="C", summary="s").readings == ()


def test_sin_lecturas_no_hay_seccion_de_lecturas():
    md = _poblado("en").to_markdown()
    assert i18n.label("hdr_readings", "en") not in md
    assert i18n.label("hdr_readings", "es") not in md


# ---- 2. la seccion propia, y donde va -----------------------------------------------


def test_con_lecturas_aparece_la_seccion_entre_metodologia_y_cobertura():
    md = _poblado("en", readings=_LECTURAS).to_markdown()
    metodologia = md.index(f"## {i18n.label('hdr_methodology', 'en')}")
    lecturas = md.index(f"## {i18n.label('hdr_readings', 'en')}")
    cobertura = md.index(f"## {i18n.label('hdr_coverage_handoff', 'en')}")
    assert metodologia < lecturas < cobertura


def test_la_seccion_sigue_yendo_antes_de_cobertura_aunque_no_haya_citas():
    # Sin citas no hay seccion de Metodologia; la de lecturas no puede caerse con ella
    # ni saltar detras del bloque de cobertura, que es el cierre del documento.
    md = _poblado("en", citations=(), readings=_LECTURAS).to_markdown()
    assert f"## {i18n.label('hdr_methodology', 'en')}" not in md
    assert md.index(f"## {i18n.label('hdr_readings', 'en')}") < md.index(
        f"## {i18n.label('hdr_coverage_handoff', 'en')}"
    )


def test_cada_linea_lleva_libro_pagina_textual_y_extracto_entre_comillas():
    md = _poblado("en", readings=_LECTURAS).to_markdown()
    assert ('- **vandeput-inventory-optimization** PDF p.83-84 - "Safety stock sizing '
            'depends on the risk period, not on the lead time alone."') in md
    assert ('- **liu-supply-chain-analytics** PDF p.12 - '
            '"Analytics without a decision attached is reporting."') in md


def test_la_pagina_viaja_textual_y_nunca_se_convierte_en_folio_pelado():
    md = _poblado("en", readings=(Reading("vandeput", "PDF p.83-84", "x"),)).to_markdown()
    assert "PDF p.83-84" in md
    assert "p.83-84 -" in md  # el rotulo entero, no "p.83" suelto
    assert " p.83 " not in md


def test_una_lectura_sin_pagina_igual_dice_que_no_la_tiene():
    md = _poblado("en", readings=(Reading("vandeput", "sin pagina", "x"),)).to_markdown()
    assert '- **vandeput** sin pagina - "x"' in md


def test_una_lectura_sin_extracto_no_inventa_un_par_de_comillas_vacias():
    # El conector permite una entrada sin cuerpo (libro + pagina ya es procedencia
    # completa). Renderizar `- **x** PDF p.5 - ""` afirma una cita textual que nadie leyo.
    md = _poblado("en", readings=(Reading("vandeput", "PDF p.5", ""),)).to_markdown()
    assert "- **vandeput** PDF p.5" in md
    assert '""' not in md


def test_el_markdown_con_lecturas_sigue_siendo_ascii_y_cp1252():
    md = _poblado("es", readings=_LECTURAS).to_markdown()
    assert md.isascii()
    md.encode("cp1252")


# ---- 3. la intro: no son citas, y la pagina es la FISICA -----------------------------


def test_la_intro_dice_que_no_son_las_fuentes_citadas_del_metodo():
    for lang, palabra in (("en", "cited"), ("es", "citadas")):
        intro = i18n.label("readings_intro", lang)
        assert palabra in intro.lower()
        assert intro in _poblado(lang, readings=_LECTURAS).to_markdown()


def test_la_intro_dice_que_la_pagina_es_la_fisica_del_pdf_no_el_folio_impreso():
    # Esta mitad es la que evita que el cliente falsee la cita solo: abre el libro
    # impreso en la 83, no encuentra nada, y la cita queda desacreditada.
    for lang, fisica, impreso in (("en", "physical", "printed"), ("es", "fisica", "impreso")):
        intro = i18n.label("readings_intro", lang).lower()
        assert fisica in intro and impreso in intro
        assert "pdf" in intro


def test_la_intro_es_una_sola_linea():
    for lang in ("en", "es"):
        assert "\n" not in i18n.label("readings_intro", lang)


def test_el_header_y_la_intro_salen_de_i18n_en_los_dos_idiomas():
    for lang in ("en", "es"):
        md = _poblado(lang, readings=_LECTURAS).to_markdown()
        assert f"## {i18n.label('hdr_readings', lang)}" in md
        assert i18n.label("readings_intro", lang) in md
    assert i18n.label("hdr_readings", "en") != i18n.label("hdr_readings", "es")


# ---- 4. readings y citations no se mezclan JAMAS -------------------------------------


def test_poblar_readings_deja_citations_intacto():
    base = _poblado("en")
    con_lecturas = _poblado("en", readings=_LECTURAS)
    assert con_lecturas.citations == base.citations
    seccion = con_lecturas.to_markdown().split(f"## {i18n.label('hdr_methodology', 'en')}")[1]
    seccion = seccion.split(f"## {i18n.label('hdr_readings', 'en')}")[0]
    assert "vandeput-inventory-optimization" not in seccion  # ningun pasaje entro a Metodologia


def test_poblar_citations_deja_readings_vacio_y_sin_seccion():
    d = _poblado("en", citations=("Otra cita - x.pdf Ch.1",))
    assert d.readings == ()
    assert f"## {i18n.label('hdr_readings', 'en')}" not in d.to_markdown()


def test_una_lectura_nunca_aparece_como_vinneta_de_la_seccion_de_citas():
    md = _poblado("en", citations=(), readings=_LECTURAS).to_markdown()
    assert f"## {i18n.label('hdr_methodology', 'en')}" not in md
    assert i18n.label("methodology_intro", "en") not in md


# ---- 5. el xlsx: misma pagina textual, mismo aviso -----------------------------------


def test_el_xlsx_no_gana_ninguna_hoja_cuando_no_hay_lecturas(tmp_path):
    wb = load_workbook(_poblado("en").to_excel(tmp_path / "d.xlsx"))
    assert i18n.label("sheet_readings", "en") not in wb.sheetnames


def test_el_xlsx_gana_la_hoja_de_lecturas_con_el_aviso_y_la_pagina_textual(tmp_path):
    wb = load_workbook(_poblado("en", readings=_LECTURAS).to_excel(tmp_path / "d.xlsx"))
    hoja = wb[i18n.label("sheet_readings", "en")]
    filas = [tuple(r) for r in hoja.iter_rows(values_only=True)]
    # El aviso viaja PEGADO a la pagina: una columna "Page" suelta en una planilla, sin
    # decir que es la fisica del PDF, es exactamente el modo de falla que esto evita.
    assert filas[0][0] == i18n.label("readings_intro", "en")
    assert filas[1] == (i18n.label("readings_col_book", "en"),
                        i18n.label("readings_col_page", "en"),
                        i18n.label("readings_col_excerpt", "en"))
    assert filas[2] == ("vandeput-inventory-optimization", "PDF p.83-84", _LECTURAS[0].excerpt)


def test_la_hoja_de_lecturas_no_toca_la_de_citas(tmp_path):
    wb = load_workbook(_poblado("en", readings=_LECTURAS).to_excel(tmp_path / "d.xlsx"))
    citas = [tuple(r) for r in wb[i18n.label("sheet_citations", "en")].iter_rows(values_only=True)]
    assert citas == [(i18n.label("col_source", "en"),),
                     ("Reorder Point - vandeput.pdf Ch.4 -> src/eoq.py:L12",)]


# ---- 6. el tope del extracto es UNO solo, y es el del conector -----------------------


def test_el_tope_del_conector_es_el_que_llega_a_la_pagina_y_al_xlsx(tmp_path):
    # El tope de copyright vive en el conector (`MAX_EXCERPT_CHARS`), no se vuelve a
    # aplicar aca: un segundo tope seria una segunda verdad. Lo que se fija es que el
    # unico que existe efectivamente llega a los dos formatos.
    largo = "palabra " * 200
    (lectura_larga,) = lectura.parse_readings(f"1. **vandeput** PDF p.9\n   {largo}\n")
    assert len(lectura_larga.excerpt) <= MAX_EXCERPT_CHARS

    d = _poblado("en", readings=(lectura_larga,))
    assert lectura_larga.excerpt in d.to_markdown()
    hoja = load_workbook(d.to_excel(tmp_path / "d.xlsx"))[i18n.label("sheet_readings", "en")]
    assert [tuple(r) for r in hoja.iter_rows(values_only=True)][2][2] == lectura_larga.excerpt


def test_el_render_no_puede_recibir_un_extracto_por_encima_del_tope():
    """El render no recorta, y no hace falta: el valor invalido no se puede construir.

    Este es el agujero que tenia la primera version. `to_markdown()` imprime el extracto
    tal como llega, asi que `Reading("libro", "PDF p.1", "x" * 5000)` publicaba 5.000
    caracteres de un PDF ajeno entre comillas en el entregable del cliente. El arreglo NO
    es un segundo recorte aca -- eso serian dos verdades sobre cuanto se puede copiar --
    sino que el tipo rechace el valor en el borde donde se construye.
    """
    with pytest.raises(ValueError):
        Reading("vandeput", "PDF p.1", "x" * (MAX_EXCERPT_CHARS + 1))


def test_el_render_no_puede_recibir_un_extracto_partido_en_dos_lineas():
    """Un salto de linea desprende la cola del extracto del libro que la respalda.

    La vinneta es `- **libro** PDF p.83 - "texto"`: con un salto adentro, la segunda
    mitad sale como un renglon suelto de prosa que ya no cuelga de ningun libro ni de
    ninguna pagina, y con la comilla de cierre colgando. Tampoco se sanea aca, por la
    misma razon: el valor no llega a existir.
    """
    for campo in range(3):
        args = ["vandeput", "PDF p.83", "cuerpo"]
        args[campo] = args[campo] + chr(10) + "cola desprendida"
        with pytest.raises(ValueError):
            Reading(*args)


def test_lo_que_el_conector_produce_siempre_se_puede_renderizar(tmp_path):
    """La cara util de la misma invariante: el productor real nunca queda del lado malo.

    Si la validacion del tipo fuera mas estricta que lo que `parse_readings` produce,
    esto seria un entregable caido en vez de una seccion de procedencia. Se le pasa el
    peor cable posible y el entregable tiene que salir en los dos formatos.
    """
    largo = ("palabra " * 400).strip()
    crudo = chr(10).join(["1. **vandeput** PDF p.9", "   " + largo,
                          "   segunda linea del cuerpo", ""])
    lecturas = lectura.parse_readings(crudo)
    d = _poblado("en", readings=lecturas)
    assert lecturas[0].excerpt in d.to_markdown()
    d.to_excel(tmp_path / "d.xlsx")


# ---- 7. el xlsx no puede COMERSE el pasaje ------------------------------------------
#
# El requisito es que el extracto llegue igual a los dos formatos. openpyxl rompe eso en
# silencio para una clase de extracto: todo str que arranca con "=" se escribe como
# FORMULA (`data_type='f'`), asi que el markdown conserva el texto y en Excel el cliente
# ve `#NAME?` donde tenia que ver el pasaje. Y es alcanzable desde el productor real: un
# libro de inventario esta lleno de renglones que sobreviven `_sanear` arrancando con
# "=" o "-" ("= 1.65 x sigma x raiz L"). La casa ya contesta esto en todos los demas
# escritores de Excel con `src.sanitize.defuse_formula`; esta era la primera celda de
# `to_excel` alimentada por texto de otro proceso.

_PELIGROSOS = ("=SUM(A1:A9)+cita del libro", "= 1.65 x sigma x raiz L",
               "+1 vuelta de inventario", "-30% de sobrestock", "@sku critico")


def _hoja_lecturas(d, tmp_path, nombre="d.xlsx"):
    wb = load_workbook(d.to_excel(tmp_path / nombre))
    return [tuple(r) for r in wb[i18n.label("sheet_readings", "en")].iter_rows(values_only=True)]


@pytest.mark.parametrize("texto", _PELIGROSOS)
def test_un_extracto_que_arranca_como_formula_sigue_siendo_texto_en_el_xlsx(texto, tmp_path):
    d = _poblado("en", readings=(Reading("vandeput", "PDF p.9", texto),))
    wb = load_workbook(d.to_excel(tmp_path / "d.xlsx"))
    celda = wb[i18n.label("sheet_readings", "en")]["C3"]
    assert celda.data_type == "s", "el pasaje se escribio como formula: el cliente ve #NAME?"
    assert texto in str(celda.value)


@pytest.mark.parametrize("texto", _PELIGROSOS)
def test_el_libro_y_la_pagina_tampoco_se_evaporan_como_formula(texto, tmp_path):
    """Las tres celdas vienen del mismo texto ajeno; una sola defendida no alcanza."""
    d = _poblado("en", readings=(Reading(texto[:120], texto, "cuerpo"),))
    wb = load_workbook(d.to_excel(tmp_path / "d.xlsx"))
    hoja = wb[i18n.label("sheet_readings", "en")]
    assert hoja["A3"].data_type == "s" and hoja["B3"].data_type == "s"
    assert texto in str(hoja["B3"].value)


def test_el_pasaje_sigue_entero_en_el_markdown_sin_apostrofe(tmp_path):
    """El precio del arreglo se paga SOLO en el xlsx, no en el markdown.

    `defuse_formula` antepone un apostrofe, y openpyxl lo escribe como texto literal (no
    como el estilo `quotePrefix`), asi que en Excel se ve. Es el costo aceptado: un
    apostrofe al lado del pasaje es mejor que el pasaje ausente. Lo que NO se acepta es
    que ese apostrofe se filtre al markdown, que no tiene el problema que lo motiva.
    """
    texto = "= 1.65 x sigma x raiz L"
    d = _poblado("en", readings=(Reading("vandeput", "PDF p.9", texto),))
    assert f'- **vandeput** PDF p.9 - "{texto}"' in d.to_markdown()
    assert "'=" not in d.to_markdown()


def test_un_extracto_inocente_llega_identico_a_las_dos_caras(tmp_path):
    """El defuse no puede cobrarle nada al 99% de los pasajes, que no arrancan con eso."""
    d = _poblado("en", readings=_LECTURAS)
    filas = _hoja_lecturas(d, tmp_path)
    assert filas[2] == ("vandeput-inventory-optimization", "PDF p.83-84", _LECTURAS[0].excerpt)
    assert filas[3] == ("liu-supply-chain-analytics", "PDF p.12", _LECTURAS[1].excerpt)


def test_lo_que_el_conector_produce_arrancando_con_formula_sobrevive(tmp_path):
    """La misma invariante, pero entrando por el productor real en vez de a mano."""
    crudo = "1. **vandeput** PDF p.9\n   = 1.65 x sigma x raiz de L dias\n"
    (leida,) = lectura.parse_readings(crudo)
    assert leida.excerpt.startswith("=")  # el saneo no toca el primer caracter
    d = _poblado("en", readings=(leida,))
    assert leida.excerpt in d.to_markdown()
    assert leida.excerpt in str(_hoja_lecturas(d, tmp_path)[2][2])
# ---- 8. lo ilegible no llega a la pagina del cliente, por ninguna de las dos caras ---
#
# El piso de legibilidad vive en el conector (`lectura.MAX_RELLENO_EXCERPT`), pero la
# unica razon por la que existe es esta: `to_markdown` y `to_excel` imprimen el extracto
# SIN filtro, y son la ultima milla antes del cliente. Estos dos tests entran por el
# productor real, no construyendo `Reading` a mano, porque lo que se afirma es que el
# camino COMPLETO -- cable, parser, deck, render -- no publica un pasaje que nadie puede
# leer bajo un libro y una pagina que si son reales.

# Textual del indice vivo: vandeput-data-science-supply-chain-forecasting, PDF p.298-299.
# El U+00B7 es lo que manda el PDF con el glyph-mapping roto donde iba una letra.
_CABLE_ILEGIBLE = (
    "1. **vandeput-data-science-supply-chain-forecasting** PDF p.298-299\n"
    "   y\u00b7 \u00b7 \u00b7 \u00b7up\u00b7lk \u00b7h\u00b7l\u00b7m\u00b7k\u00b7gbl \u00b7f\u00b7j is one of the most used "
    "libraries in \u00b7ython (it was created by \u00b7es \u00b7c\u00b7inney in 2008). \u00b7he name "
    "comes from h\u00b7f el \u00b7 ta, because it helps to order data into ta\u00b7 bles.\n"
)


def test_un_pasaje_ilegible_llega_al_markdown_como_libro_y_pagina_sin_comillas():
    (leida,) = lectura.parse_readings(_CABLE_ILEGIBLE)
    md = _poblado("en", readings=(leida,)).to_markdown()

    assert "vandeput-data-science-supply-chain-forecasting" in md
    assert "PDF p.298-299" in md
    assert "-ython" not in md and "-c-inney" not in md  # la basura no viaja
    linea = next(ln for ln in md.splitlines() if "PDF p.298-299" in ln)
    assert '"' not in linea  # y no queda un par de comillas afirmando una cita


def test_un_pasaje_ilegible_tampoco_entra_al_xlsx(tmp_path):
    """La otra cara. El xlsx es la que el cliente filtra, ordena y pega en un mail."""
    (leida,) = lectura.parse_readings(_CABLE_ILEGIBLE)
    filas = _hoja_lecturas(_poblado("en", readings=(leida,)), tmp_path, "ilegible.xlsx")

    assert filas[2][:2] == ("vandeput-data-science-supply-chain-forecasting",
                            "PDF p.298-299")
    assert not (filas[2][2] or "")
