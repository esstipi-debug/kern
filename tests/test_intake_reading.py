"""How Kern read the file has to reach the deliverable the client opens.

``src.excel_hygiene`` already counted every discarded row, and ``load_table``
left the report in ``attrs`` 'so the deliverable can tell the client'. Nothing
consumed it. Column mapping died inside ``prepare_tracked``. The client who
opens the Excel could not reconcile our numbers against their own file.

Same placement as the delivery-date caveat: the map of sources, not the
findings. A finding is a problem of their operation; this is how we read.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

from jobs import deliverables, inventory_deliverable
from jobs.intake import (
    HEADER_ATTR,
    HYGIENE_ATTR,
    SHEET_ATTR,
    detect_columns,
    load_table,
    normalize_tracked,
    prepare_tracked,
)
from jobs.inventory_optimization import run as run_inventory

PORTFOLIO = "data/sample_demand_portfolio.csv"


def _excel_con_portada_total_y_combinada(ruta: Path) -> Path:
    """The shape a real LatAm export takes: title row, merged SKU, TOTAL at the end.

    Eight demand days so ``run_inventory`` has enough history to size a policy
    -- the reading has to survive a real JobReport, not only ``prepare_tracked``.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Ventas"
    ws.append(["Reporte de ventas 2026"])
    ws.append(["fecha", "sku", "cantidad", "costo_unitario"])
    ws.append(["2026-01-01", "SKU-1", 10, 5])
    for day in range(2, 9):
        ws.append([f"2026-01-{day:02d}", None, 10, 5])
    ws.append([])
    ws.append(["", "TOTAL", 80, ""])
    ws.merge_cells("B3:B10")
    wb.save(ruta)
    return ruta


def test_load_table_deja_hoja_y_fila_en_attrs(tmp_path: Path):
    ruta = _excel_con_portada_total_y_combinada(tmp_path / "cliente.xlsx")
    frame = load_table(ruta)

    assert frame.attrs[SHEET_ATTR] == "Ventas"
    assert frame.attrs[HEADER_ATTR] == 1  # title on row 0; pandas 0-based
    assert HYGIENE_ATTR in frame.attrs


def test_un_csv_no_finge_una_hoja_de_excel(tmp_path: Path):
    ruta = tmp_path / "datos.csv"
    ruta.write_text("fecha,sku,cantidad\n2026-01-05,SKU-1,10\n", encoding="utf-8")

    frame = load_table(ruta)

    assert SHEET_ATTR not in frame.attrs
    assert HEADER_ATTR not in frame.attrs


def test_prepare_tracked_lleva_el_mapeo_y_la_higiene(tmp_path: Path):
    ruta = _excel_con_portada_total_y_combinada(tmp_path / "cliente.xlsx")
    _demand, quality = prepare_tracked(ruta, period="D")

    reading = quality.reading
    assert reading is not None
    assert reading.source_kind == "excel"
    assert reading.sheet == "Ventas"
    assert reading.header_row == 1
    assert dict(reading.mapping)["product_id"] == "sku"
    assert dict(reading.mapping)["date"] == "fecha"
    assert dict(reading.mapping)["quantity"] == "cantidad"
    assert reading.hygiene is not None
    assert reading.hygiene.dropped["fila_de_total"] == 1
    assert reading.hygiene.dropped["fila_vacia"] == 1
    assert reading.hygiene.filled["sku"] == 7
    assert reading.noteworthy_hygiene()
    assert "header on row 2" in reading.location_text()
    assert "sku" in reading.mapping_text()
    assert "total/summary" in reading.hygiene_text()


def test_prepare_tracked_sobre_el_sample_nombra_las_columnas_canonicas():
    _demand, quality = prepare_tracked(PORTFOLIO)

    reading = quality.reading
    assert reading is not None
    assert reading.source_kind == "csv"
    assert "CSV" in reading.location_text()
    mapeo = dict(reading.mapping)
    assert mapeo["date"] == "date"
    assert mapeo["product_id"] == "product_id"
    assert mapeo["quantity"] == "quantity"
    assert not reading.noteworthy_hygiene()


def test_normalize_tracked_sobre_un_frame_en_memoria_igual_deja_el_mapeo():
    raw = pd.DataFrame(
        {"InvoiceDate": ["2024-01-01"], "StockCode": ["A"], "Quantity": [10]}
    )
    mapping = detect_columns(raw)
    _frame, quality = normalize_tracked(raw, mapping, period="W")

    assert quality.reading is not None
    mapeo = dict(quality.reading.mapping)
    assert mapeo["date"] == "InvoiceDate"
    assert mapeo["product_id"] == "StockCode"
    assert quality.reading.source_kind == "csv"


def test_el_reporte_lleva_la_lectura_hasta_el_jobreport(tmp_path: Path):
    ruta = _excel_con_portada_total_y_combinada(tmp_path / "cliente.xlsx")
    demand, quality = prepare_tracked(ruta, period="D")
    reporte = run_inventory(demand, period_freq="D", intake_quality=quality)

    assert reporte.intake_quality is quality
    assert reporte.intake_quality.reading is not None
    assert dict(reporte.intake_quality.reading.mapping)["product_id"] == "sku"


def test_el_deck_publica_la_lectura_como_procedencia_y_no_como_hallazgo(tmp_path: Path):
    ruta = _excel_con_portada_total_y_combinada(tmp_path / "cliente.xlsx")
    demand, quality = prepare_tracked(ruta, period="D")
    reporte = run_inventory(demand, period_freq="D", intake_quality=quality)
    deck = inventory_deliverable.build(reporte, client="ACME")

    fuentes = {d.field: d.source for d in deck.data_sources}
    assert "How Kern read the file" in fuentes
    assert "Ventas" in fuentes["How Kern read the file"]
    assert "header on row 2" in fuentes["How Kern read the file"]
    assert "sku" in fuentes["Column mapping"]
    assert "total/summary" in fuentes["What was taken off the spreadsheet"]
    assert not any("Ventas" in f.detail for f in deck.findings)
    assert not any("total/summary" in f.detail for f in deck.findings)


def test_sin_lectura_no_se_agrega_una_fila_vacia_al_mapa_de_fuentes():
    demand, quality = prepare_tracked(PORTFOLIO)
    reporte = run_inventory(demand, intake_quality=quality)
    # Force the untracked shape the older callers still produce.
    reporte_sin = run_inventory(demand)
    deck = inventory_deliverable.build(reporte_sin, client="ACME")

    assert all(d.source.strip() for d in deck.data_sources)
    assert not any(d.field == "How Kern read the file" for d in deck.data_sources)
    # And the tracked path DID add it, so the skip is not vacuously true.
    deck_con = inventory_deliverable.build(reporte, client="ACME")
    assert any(d.field == "How Kern read the file" for d in deck_con.data_sources)


def test_la_lectura_llega_al_archivo_que_abre_el_cliente(tmp_path: Path):
    ruta = _excel_con_portada_total_y_combinada(tmp_path / "cliente.xlsx")
    demand, quality = prepare_tracked(ruta, period="D")
    reporte = run_inventory(demand, period_freq="D", intake_quality=quality)
    deck = inventory_deliverable.build(reporte, client="ACME")
    escritos = deck.write_all(tmp_path / "deck")

    markdown = Path(escritos["report"]).read_text(encoding="utf-8")
    assert "How Kern read the file" in markdown
    assert "sku" in markdown
    assert "Ventas" in markdown

    wb = load_workbook(escritos["workbook"])
    fuentes = wb["Data Sources"]
    celdas = " ".join(
        str(c.value) for row in fuentes.iter_rows() for c in row if c.value
    )
    assert "sku" in celdas
    assert "Ventas" in celdas


def test_el_xlsx_y_el_md_del_playbook_tambien_lo_dicen(tmp_path: Path):
    """jobs/deliverables.py is what examples/run_inventory_job.py writes."""
    ruta = _excel_con_portada_total_y_combinada(tmp_path / "cliente.xlsx")
    demand, quality = prepare_tracked(ruta, period="D")
    reporte = run_inventory(demand, period_freq="D", intake_quality=quality)
    written = deliverables.write_all(reporte, tmp_path / "out", client="ACME")

    md = written["report"].read_text(encoding="utf-8")
    assert "## How Kern read this file" in md
    assert "sku" in md
    assert "Ventas" in md

    wb = load_workbook(written["excel"])
    assert "How Kern read this" in wb.sheetnames
    hoja = wb["How Kern read this"]
    texto = " ".join(
        str(c.value) for row in hoja.iter_rows() for c in row if c.value
    )
    assert "sku" in texto
    assert "total/summary" in texto


def test_untracked_playbook_no_inventa_la_hoja(tmp_path: Path):
    """The older prepare() path still writes a two-sheet workbook."""
    demand = pd.read_csv(PORTFOLIO)
    from jobs.intake import detect_columns, normalize
    reporte = run_inventory(normalize(demand, detect_columns(demand)))
    written = deliverables.write_all(reporte, tmp_path, client="Acme")
    wb = load_workbook(written["excel"])
    assert "How Kern read this" not in wb.sheetnames
    assert "How Kern read this file" not in written["report"].read_text(encoding="utf-8")
