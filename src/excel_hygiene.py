"""Los defectos de una planilla real que NO fallan, y por eso son los peores.

``jobs/intake.py`` ya resuelve dos: encuentra la hoja y la fila de encabezado por
evidencia (``detect_sheet_and_header``) y parsea separadores de miles LATAM y
anglosajones (``_parse_quantities``). Los tres que quedaban comparten una
propiedad que los otros no tienen: **no producen un error, producen un numero**.

* La fila ``TOTAL`` del final entra como una venta mas, con la suma del periodo
  entero cargada a un solo dia. No rompe el pronostico: lo corre. Y el desvio
  inventado que genera viaja hasta el stock de seguridad, que se dimensiona
  sobre una dispersion que nunca existio.
* Las celdas combinadas -- escribir el SKU una vez y combinar cinco filas -- son
  como se ve una planilla ordenada. ``pd.read_excel`` devuelve cuatro vacias, y
  esas cuatro filas de venta real se caen del intake por "sin producto". Menos
  demanda de la que hubo, en silencio.
* Los separadores en blanco y las columnas ``Unnamed`` sin datos.

Por que esto es producto y no higiene: un embudo de autoservicio a precio bajo
solo cierra si el archivo del cliente entra sin que nadie lo toque. Cada planilla
que hay que arreglar a mano es margen quemado, y quien exporta a Excel lo hace,
por revelacion de preferencia, para no pagar.

**La regla que gobierna el modulo: cada fila descartada se cuenta y se declara.**
Un limpiador silencioso es peor que ninguno -- convierte "tu planilla traia doce
filas de basura" en "tu demanda es doce filas mas chica", y eso no se ve en la
forma de la salida. Es el mismo motivo por el que ``src/order_fulfillment.py``
reporta ``not_measured`` en vez de un cero.
"""

from __future__ import annotations

import math
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd

#: Etiquetas que marcan una fila de cierre, sin acentos y en minuscula. Cubren
#: castellano e ingles porque una planilla LATAM mezcla los dos sin avisar.
_TOTAL_LABELS = frozenset(
    {
        "total",
        "totales",
        "total general",
        "gran total",
        "grand total",
        "subtotal",
        "sub total",
        "sub-total",
        "suma",
        "sumatoria",
        "acumulado",
        "resumen",
    }
)

#: Cuanto puede alejarse una fila de la suma de las de arriba y seguir siendo su
#: total. No es cero por los redondeos de una planilla que muestra dos decimales
#: y guarda quince.
_SUM_TOLERANCE = 0.005

_UNNAMED = re.compile(r"^unnamed:?\s*\d*$", re.IGNORECASE)


def _fold(value: object) -> str:
    """Minuscula, sin acentos y sin espacios de sobra. ``"Total "`` -> ``"total"``."""
    texto = str(value if value is not None else "").strip().lower()
    sin_tildes = unicodedata.normalize("NFKD", texto)
    return "".join(c for c in sin_tildes if not unicodedata.combining(c))


def is_total_label(value: object) -> bool:
    """Si esta celda es la ETIQUETA de una fila de cierre.

    Coincidencia exacta contra el vocabulario y no por subcadena, a proposito:
    ``TOTAL-500`` es un SKU perfectamente valido, y ``Kit Total Care`` es un
    producto. Descartar filas porque su nombre CONTIENE "total" borraria ventas
    reales -- y lo haria justo en los catalogos que mas usan la palabra.
    """
    return _fold(value) in _TOTAL_LABELS


def _is_blank(value: object) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    return str(value).strip() == ""


def _numeric_columns(frame: pd.DataFrame) -> tuple[str, ...]:
    return tuple(c for c in frame.columns if pd.api.types.is_numeric_dtype(frame[c]))


def _closes_the_block(frame: pd.DataFrame, idx: int, columns: tuple[str, ...]) -> bool:
    """Si la ULTIMA fila cuadra con la suma de todas las de arriba.

    El caso peligroso de verdad: la etiqueta quedo en blanco y lo unico que
    delata a la fila es que su numero es la suma. Se exige que sea la ultima fila
    -- una coincidencia en el medio es una venta que casualmente cuadra, y
    descartarla seria inventar un total donde habia demanda.
    """
    if idx != len(frame) - 1 or len(frame) < 2:
        return False
    previas = frame.iloc[:idx]
    for columna in columns:
        valor = frame.iloc[idx][columna]
        if _is_blank(valor):
            return False
        try:
            objetivo = float(valor)
            suma = float(previas[columna].astype(float).sum())
        except (TypeError, ValueError):
            return False
        if suma == 0 or abs(objetivo - suma) > _SUM_TOLERANCE * max(abs(suma), 1.0):
            return False
    return True


@dataclass(frozen=True)
class HygieneReport:
    """Lo que se le saco a la planilla, SIN la planilla.

    Existe separado de ``CleanTable`` por una razon mecanica y no estetica:
    este informe viaja en ``DataFrame.attrs``, y pandas propaga y compara
    ``attrs`` al operar sobre el frame. Un informe que contuviera el frame
    haria que el frame se contenga a si mismo, y cualquier operacion despues
    revienta con un ``RecursionError`` a novecientos cuadros de profundidad --
    lejisimos de la linea que lo causo.
    """

    dropped: dict = field(default_factory=dict)
    dropped_columns: tuple[str, ...] = ()
    filled: dict = field(default_factory=dict)
    rows_in: int = 0

    @property
    def clean(self) -> bool:
        return not any(self.dropped.values()) and not self.dropped_columns

    @property
    def rows_dropped(self) -> int:
        return sum(self.dropped.values())


@dataclass(frozen=True)
class CleanTable:
    """La tabla limpia, mas el inventario completo de lo que se le saco."""

    frame: pd.DataFrame
    #: Motivo -> cuantas filas. Las claves existen siempre, aunque valgan 0.
    dropped: dict = field(default_factory=dict)
    #: Columnas eliminadas por no tener ni nombre ni datos.
    dropped_columns: tuple[str, ...] = ()
    #: Columna -> cuantas celdas se completaron desde una celda COMBINADA del
    #: libro. Nunca desde una suposicion: ver ``read_clean_table``.
    filled: dict = field(default_factory=dict)
    #: Cuantas filas traia el archivo antes de tocarlo.
    rows_in: int = 0

    @property
    def clean(self) -> bool:
        """Si el archivo entro sin que hubiera que sacarle nada.

        Un ``False`` no es un error: es lo que el entregable tiene que decirle al
        cliente sobre su propia planilla.
        """
        return not any(self.dropped.values()) and not self.dropped_columns

    @property
    def report(self) -> HygieneReport:
        """El informe suelto, que es lo unico que puede viajar en ``attrs``."""
        return HygieneReport(
            dropped=dict(self.dropped),
            dropped_columns=self.dropped_columns,
            filled=dict(self.filled),
            rows_in=self.rows_in,
        )


def clean_frame(
    frame: pd.DataFrame,
    *,
    key_column: str,
    numeric_columns: tuple[str, ...] | None = None,
) -> CleanTable:
    """Saca relleno, separadores y filas de cierre. Cuenta todo lo que saca.

    ``key_column`` es la columna que identifica la fila (el SKU, normalmente):
    es donde se busca la etiqueta de total y lo que decide si una fila es un
    separador. ``numeric_columns`` acota contra que columnas se prueba el cuadre
    de la suma; por defecto, todas las numericas.
    """
    descartes = {"fila_vacia": 0, "fila_de_total": 0}
    trabajo = frame.copy()

    vacias = [c for c in trabajo.columns if _UNNAMED.match(str(c)) and trabajo[c].isna().all()]
    if vacias:
        trabajo = trabajo.drop(columns=vacias)

    numericas = numeric_columns if numeric_columns is not None else _numeric_columns(trabajo)
    conservar: list[int] = []
    for posicion in range(len(trabajo)):
        fila = trabajo.iloc[posicion]
        if all(_is_blank(v) for v in fila):
            descartes["fila_vacia"] += 1
            continue
        if is_total_label(fila.get(key_column)):
            descartes["fila_de_total"] += 1
            continue
        if _is_blank(fila.get(key_column)) and _closes_the_block(trabajo, posicion, numericas):
            descartes["fila_de_total"] += 1
            continue
        conservar.append(posicion)

    limpio = trabajo.iloc[conservar].reset_index(drop=True)
    return CleanTable(
        frame=limpio,
        dropped=descartes,
        dropped_columns=tuple(str(c) for c in vacias),
        filled={},
        rows_in=len(frame),
    )


def _merged_fills(path: str | Path, sheet: object, header_row: int) -> dict[int, dict[int, object]]:
    """Fila -> columna -> valor, SOLO para celdas que el libro declara combinadas.

    Se lee del propio libro (``ws.merged_cells.ranges``) y no se adivina
    rellenando hacia abajo. La diferencia importa: un ``ffill`` ciego convierte
    todo hueco en una repeticion del valor de arriba, y un hueco real -- una fila
    a la que de verdad le falta el producto -- pasaria a tener uno inventado.
    Aca solo se completa donde el archivo dice que hubo una combinacion.
    """
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=False, data_only=True)
    try:
        ws = wb[sheet] if isinstance(sheet, str) else wb.worksheets[int(sheet or 0)]
        salida: dict[int, dict[int, object]] = {}
        for rango in list(ws.merged_cells.ranges):
            ancla = ws.cell(row=rango.min_row, column=rango.min_col).value
            if ancla is None:
                continue
            for fila_xl in range(rango.min_row, rango.max_row + 1):
                for col_xl in range(rango.min_col, rango.max_col + 1):
                    if fila_xl == rango.min_row and col_xl == rango.min_col:
                        continue
                    # -1 por el encabezado, -1 mas porque openpyxl cuenta desde 1
                    idx = fila_xl - header_row - 2
                    if idx >= 0:
                        salida.setdefault(idx, {})[col_xl - 1] = ancla
        return salida
    finally:
        wb.close()


def read_clean_table(
    path: str | Path,
    *,
    key_column: str,
    sheet: object = 0,
    header_row: int = 0,
    numeric_columns: tuple[str, ...] | None = None,
) -> CleanTable:
    """Lee un ``.xlsx`` resolviendo combinadas, y despues lo limpia.

    El orden no es casual: las combinadas se resuelven ANTES de descartar filas.
    Una fila cuyo unico defecto era heredar el SKU de una celda combinada tiene
    que sobrevivir, y si se limpiara primero se iria por "vacia" o por "sin
    producto" -- que es exactamente la venta real que hoy se pierde.
    """
    frame = pd.read_excel(path, sheet_name=sheet, header=header_row)
    if isinstance(frame, dict):  # sheet_name=None devuelve todas
        frame = next(iter(frame.values()))

    completados: dict[str, int] = {}
    for idx, columnas in _merged_fills(path, sheet, header_row).items():
        if idx >= len(frame):
            continue
        for col_idx, valor in columnas.items():
            if col_idx >= len(frame.columns):
                continue
            nombre = str(frame.columns[col_idx])
            if _is_blank(frame.iat[idx, col_idx]):
                frame.iat[idx, col_idx] = valor
                completados[nombre] = completados.get(nombre, 0) + 1

    limpio = clean_frame(frame, key_column=key_column, numeric_columns=numeric_columns)
    return CleanTable(
        frame=limpio.frame,
        dropped=limpio.dropped,
        dropped_columns=limpio.dropped_columns,
        filled=completados,
        rows_in=limpio.rows_in,
    )
