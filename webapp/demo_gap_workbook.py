"""``brecha_en_plata.xlsx`` -- el entregable de /api/demo-gap-value.

Tercero de los tres libros de /demo. Las cinco reglas comunes estan en el
docstring de :mod:`webapp.demo_workbook`; las primitivas en
:mod:`webapp.demo_workbook_base`.

Tres cosas propias de este libro, y ninguna es de estilo:

1. **Las dos monedas van juntas o no van.** El docstring de
   :mod:`webapp.demo_gap_value` prohibe reportar solo la caja liberable: leido
   solo, esto es "reduce inventario", que es la frase que operaciones entiende
   como *"quieres que me quiebre"*. Ningun grafico de plata de este libro tiene
   una sola serie.

2. **Las dos cifras anuales se contienen.** ``recurring_saving`` se calcula sobre
   la caja liberable; ``net_carrying_change`` es el neto por la misma tasa, y el
   neto normalmente es POSITIVO. Un lector que ve "ahorro" y "cambio +" a tres
   filas de distancia lee una contradiccion, salvo que el termino del medio --
   el costo de cubrir lo que falta -- este en la hoja. No sale del motor: lo
   calcula esta planilla, y lo dice.

       exposure x tasa  -  ahorro  =  cambio neto del costo de mantener

3. **Nunca se escribe un negativo en un bloque de grafico.** ``_bar`` fija
   ``y_axis.scaling.min = 0``, asi que una barra bajo cero no se dibuja: se
   desvanece sin error. Las dos mitades van como dos series de magnitudes
   positivas, distinguidas por color y leyenda.

Y una regla de unidades: se lee **solo** ``SkuGap`` y ``GapValue``, nunca un
``JobReport``. ``SkuGap`` guarda los plazos en DIAS -- se llenan antes de la
conversion a periodos -- mientras que ``SkuRecommendation`` expone campos en
periodos. Mezclarlos es el error invisible de factor 7.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import Reference
from openpyxl.styles import Font

from webapp.demo_gap_value import CTA_PATH, HOLDING_RATE, GapValue
from webapp.demo_lead_time_scan import MIN_DELIVERIES
from webapp.demo_stock_workbook import NOTA_MONEDA
from webapp.demo_workbook_base import (
    _BASE_URL,
    _FILL_CAJA,
    _FILL_TARDE,
    _MID,
    _NOTE,
    _SECTION,
    _TITLE,
    _WRAP,
    C_ALERTA,
    C_CAJA,
    C_DECLARADO,
    C_OBSERVADO,
    FMT_COUNT,
    FMT_DAYS_1,
    FMT_MONEY,
    FMT_MONEY_2,
    FMT_MONEY_SIGNED,
    FMT_PCT,
    FMT_PCT_1,
    FMT_UNITS_1,
    NO_MEDIDO,
    TOP_N_SKU,
    _bar,
    _cap,
    _cell,
    _chart_block,
    _headers,
    _isnumber_rule,
    _link_row,
    _widths,
    rows_below,
)
from webapp.operator_profile import get_operator_profile


def _sheet_resumen(ws, value: GapValue, *, dataset_label: str, ts: str) -> None:
    _widths(ws, {"A": 40, "B": 22, "C": 30, "D": 3})
    # El termino que falta para que las tres filas anuales cierren. No existe en
    # GapValue: lo calcula esta planilla, y la hoja lo declara.
    carry_exposure = value.exposure * HOLDING_RATE

    _cell(ws, 1, 1,
          "La diferencia entre el plazo que promete tu sistema y el real, en plata", font=_TITLE)
    _cell(ws, 2, 1, "Archivos analizados:")
    _cell(ws, 2, 2, dataset_label)
    _cell(ws, 3, 1, "Fecha del analisis:")
    _cell(ws, 3, 2, ts)
    _cell(ws, 4, 1, "SKU valorizados:")
    _cell(ws, 4, 2, len(value.gaps), fmt=FMT_COUNT, font=_MID)
    # El denominador va al lado del numero, en la primera pantalla.
    _cell(ws, 4, 3, f"de {value.n_demand_skus} SKU del archivo de movimientos")
    _cell(ws, 5, 1, "Costo de mantener usado:")
    _cell(ws, 5, 2, HOLDING_RATE, fmt=FMT_PCT)
    _cell(ws, 5, 3, "anual -- supuesto por defecto del escaneo, no medido en tu operacion")

    _cell(ws, 7, 1, "Las dos monedas", font=_SECTION)
    _headers(ws, 8, ["Concepto", "Monto", "A quien le importa"])
    _cell(ws, 9, 1, "Caja que se libera, una sola vez")
    _cell(ws, 9, 2, value.cash_release, fmt=FMT_MONEY, font=Font(bold=True, size=16, color=C_CAJA))
    _cell(ws, 9, 3, "Gerencia general / finanzas")
    _cell(ws, 10, 1, "Capital que hay que poner, una sola vez")
    _cell(ws, 10, 2, value.exposure, fmt=FMT_MONEY, font=Font(bold=True, size=16, color=C_ALERTA))
    _cell(ws, 10, 3, "Operaciones")
    _cell(ws, 11, 1, "Cambio neto del inventario")
    _cell(ws, 11, 2, value.net, fmt=FMT_MONEY_SIGNED, font=Font(bold=True, size=14))
    _cell(ws, 11, 3, "Ambos")

    header_row, last_row = _chart_block(
        ws, 1, "Datos del grafico 1 de esta hoja: las dos monedas",
        ["", "Caja que se libera", "Capital que hay que poner"],
        [["Al balance, una sola vez", value.cash_release, value.exposure]],
        start_col=14, fmt=FMT_MONEY,
    )
    ws.add_chart(_bar(
        data=Reference(ws, min_col=15, max_col=16, min_row=header_row, max_row=last_row),
        cats=Reference(ws, min_col=14, min_row=header_row + 1, max_row=last_row),
        title="Las dos monedas", y_title="Monto, en la moneda de tu archivo",
        colors=(C_CAJA, C_ALERTA), width=19, height=8, gap=60,
    ), "A13")

    row = 13 + rows_below(8.0)
    _cell(ws, row, 1, "El costo de mantener, al anio", font=_SECTION)
    _headers(ws, row + 1, ["Concepto", "Monto al anio", "Sobre que se calcula"])
    pct = f"{HOLDING_RATE * 100:.0f}%"
    anuales = [
        ("Ahorro, solo sobre la caja que se libera", value.recurring_saving,
         f"El {pct} de la caja que se libera.", FMT_MONEY, None),
        ("Costo de cubrir lo que hoy falta", carry_exposure,
         f"El {pct} del capital que hay que poner. Lo calcula esta planilla; no sale del motor.",
         FMT_MONEY, None),
        ("Cambio neto del costo de mantener", value.net_carrying_change,
         "La resta de las dos filas de arriba.", FMT_MONEY_SIGNED, Font(bold=True, size=12)),
    ]
    for i, (etiqueta, monto, sobre, fmt, font) in enumerate(anuales):
        r = row + 2 + i
        _cell(ws, r, 1, etiqueta)
        _cell(ws, r, 2, monto, fmt=fmt, font=font)
        _cell(ws, r, 3, sobre, align=_WRAP)
    r = row + 5
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    _cell(ws, r, 1,
          "La ultima fila YA incluye el ahorro de la primera: es el costo de cubrir lo que falta "
          "menos ese ahorro. Positivo significa que, despues de corregir, mantener el inventario "
          "cuesta mas. Se muestran las dos mitades porque son dos decisiones distintas: bajar el "
          "buffer inflado no depende de que hagas nada con los SKU que estan cortos.",
          font=_NOTE, align=_WRAP)
    ws.row_dimensions[r].height = 44

    header_row, last_row = _chart_block(
        ws, 1, "Datos del grafico 2 de esta hoja: el costo de mantener, al anio",
        ["", "Ahorro sobre la caja que se libera", "Costo de cubrir lo que falta"],
        [["Al anio", value.recurring_saving, carry_exposure]],
        start_col=20, fmt=FMT_MONEY,
    )
    chart2_row = r + 2
    ws.add_chart(_bar(
        data=Reference(ws, min_col=21, max_col=22, min_row=header_row, max_row=last_row),
        cats=Reference(ws, min_col=20, min_row=header_row + 1, max_row=last_row),
        title=f"El costo de mantener, al anio (a {pct})", y_title="Monto al anio",
        colors=(C_CAJA, C_ALERTA), width=19, height=7, gap=60,
    ), f"A{chart2_row}")

    row = chart2_row + rows_below(7.0)
    _cell(ws, row, 1, "Que dice esto", font=_SECTION)
    row += 1
    for finding in value.findings:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        _cell(ws, row, 1, finding, align=_WRAP)
        ws.row_dimensions[row].height = 42
        row += 1
    if not value.cash_release:
        # `_findings` emite su parrafo de caja liberable sin condicion, asi que
        # con cero pega "$0 ... es ejecutable por si solo". Se corrige al lado.
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        _cell(ws, row, 1,
              "En este archivo no hay caja liberable: ningun proveedor tiene el plazo declarado "
              "inflado respecto de lo que realmente demora. Todo el movimiento esta del lado del "
              "capital que falta.", font=_NOTE, align=_WRAP)
        ws.row_dimensions[row].height = 42
        row += 1

    q = value.intake_quality
    syn = dict(q.synthesized) if q else {}
    n_cost_syn = syn.get("unit_cost", 0)
    if q and 0 < n_cost_syn < q.n_canonical_rows:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        _cell(ws, row, 1,
              f"Atencion: {n_cost_syn} de {q.n_canonical_rows} filas agregadas no traian costo "
              "unitario. Ver 'Calidad de los datos'.",
              font=Font(bold=True, color=C_ALERTA), align=_WRAP)
        row += 1

    row += 1
    _cell(ws, row, 1, "Siguiente paso", font=_SECTION)
    row += 1
    _link_row(ws, row, 1, "Diagnostico de Arranque (abrir)", f"{_BASE_URL}{CTA_PATH}")
    row += 1
    operator = get_operator_profile()
    if operator.email:
        _cell(ws, row, 1, "Responde a:")
        _cell(ws, row, 2, f"{operator.name} <{operator.email}>")
        row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    _cell(ws, row, 1,
          "Este archivo lo genero un escaneo automatico sobre los que subiste, en solo lectura: no "
          "se escribio nada en tus sistemas. Los dos montos de una sola vez y los dos montos al "
          f"anio no se suman entre si: son unidades distintas. {NOTA_MONEDA}",
          font=_NOTE, align=_WRAP)
    ws.row_dimensions[row].height = 60


def _sheet_skus(ws, value: GapValue) -> None:
    h = value.headline
    _cell(ws, 1, 1, "SKU por SKU: donde esta cada peso", font=_SECTION)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=5)
    _cell(ws, 2, 1,
          f"El SKU que mas mueve la aguja es {h['worst_sku']}: {h['worst_value']:,.0f}. Negativo es "
          "caja que se libera; positivo es capital que hay que poner.".replace(",", "."),
          font=_NOTE, align=_WRAP)

    ordenados = sorted(value.gaps, key=lambda g: g.delta_value)
    rows = _cap(list(ordenados), ws, 3, 1, len(value.gaps), "plata")
    _headers(ws, 4, [
        "SKU", "Proveedor", "Lo que dice tu sistema (d)", "Lo que dicen tus recepciones (d)",
        "Variacion entre recepciones (d)", "Stock de seguridad que calcula tu sistema (u)",
        "Stock de seguridad contra el plazo real (u)", "Diferencia (u)", "Costo unitario",
        "Plata", "Lado",
    ])
    for r, g in enumerate(rows, 5):
        lado = ("caja que se libera" if g.delta_value < 0
                else "capital que hay que poner" if g.delta_value > 0 else "sin cambio")
        _cell(ws, r, 1, g.product_id)
        _cell(ws, r, 2, g.supplier)
        _cell(ws, r, 3, g.declared_lt_days, fmt=FMT_DAYS_1)
        _cell(ws, r, 4, g.observed_lt_days, fmt=FMT_DAYS_1)
        _cell(ws, r, 5, g.observed_lt_std_days, fmt=FMT_DAYS_1)
        _cell(ws, r, 6, g.ss_erp, fmt=FMT_UNITS_1)
        _cell(ws, r, 7, g.ss_real, fmt=FMT_UNITS_1)
        _cell(ws, r, 8, g.delta_units, fmt=FMT_UNITS_1)
        _cell(ws, r, 9, g.unit_cost, fmt=FMT_MONEY_2)
        _cell(ws, r, 10, g.delta_value, fmt=FMT_MONEY_SIGNED)
        _cell(ws, r, 11, lado)
    last = 4 + len(rows)
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:K{last}"
    if rows:
        for ref, rule in (
            _isnumber_rule("J", 5, last, "<0", _FILL_CAJA),
            _isnumber_rule("J", 5, last, ">0", _FILL_TARDE),
        ):
            ws.conditional_formatting.add(ref, rule)

    n_zero = sum(1 for g in value.gaps if g.delta_value == 0.0)
    n_declared_zero = sum(1 for g in value.gaps if g.declared_lt_days <= 0)
    leyendas = [
        "Fondo verde en 'Plata': el plazo real no justifica ese stock de seguridad. Es caja que se libera.",
        "Fondo rojo en 'Plata': falta stock de seguridad. Es capital que hay que poner para no quebrar.",
        "La columna 'Stock de seguridad que calcula tu sistema' sale de dimensionar contra el plazo "
        "declarado, que es lo que hace un ERP.",
        "Un SKU que se le compra a dos proveedores queda asignado al que mas veces se lo entrego. El "
        "plazo mezclado ponderado por volumen es trabajo del Diagnostico.",
    ]
    if n_declared_zero:
        leyendas.append(
            f"{n_declared_zero} SKU tienen plazo declarado en cero en tu sistema. Un cero ahi suele "
            "ser un campo sin configurar, no una promesa de entrega inmediata: para esos SKU la "
            "columna 'Plata' es todo el stock de seguridad, no una diferencia.")
    if n_zero:
        leyendas.append(
            f"{n_zero} SKU no cambian: no registraron demanda ni costo en el periodo, asi que el "
            "calculo da el mismo stock de seguridad en los dos escenarios. Estan contados en los "
            "SKU valorizados y en ninguna de las dos monedas.")
    for i, texto in enumerate(leyendas):
        _cell(ws, last + 2 + i, 1, texto, font=_NOTE)

    top = sorted(value.gaps, key=lambda g: abs(g.delta_value), reverse=True)[:TOP_N_SKU]
    top = sorted(top, key=lambda g: g.delta_value)
    k = len(top)
    if k:
        # Cada SKU lleva valor en UNA serie y None en la otra: magnitudes positivas
        # a los dos lados, porque _bar fija el eje en cero.
        hr1, lr1 = _chart_block(
            ws, 1, f"Datos del grafico: los {k} SKU que mas mueven la aguja",
            ["SKU", "Caja que se libera", "Capital que hay que poner"],
            [[g.product_id,
              abs(g.delta_value) if g.delta_value < 0 else None,
              g.delta_value if g.delta_value > 0 else None] for g in top],
            start_col=13, fmt=FMT_MONEY,
        )
        alto = min(4 + 0.6 * k, 16)
        ws.add_chart(_bar(
            data=Reference(ws, min_col=14, max_col=15, min_row=hr1, max_row=lr1),
            cats=Reference(ws, min_col=13, min_row=hr1 + 1, max_row=lr1),
            title=f"Los {k} SKU que mas mueven la aguja", kind="bar",
            colors=(C_CAJA, C_ALERTA), width=18, height=alto,
        ), "Q2")

        hr2, lr2 = _chart_block(
            ws, lr1 + 3, "Datos del grafico: el mismo calculo, con el plazo real",
            ["SKU", "Lo que calcula tu sistema (u)", "Contra el plazo real (u)"],
            [[g.product_id, round(g.ss_erp, 1), round(g.ss_real, 1)] for g in top],
            start_col=13, fmt=FMT_UNITS_1,
        )
        # El ancla sale del ALTO RENDERIZADO del primero, jamas de la fila de su
        # bloque: el extent de un ancla es EMU absolutos.
        ws.add_chart(_bar(
            data=Reference(ws, min_col=14, max_col=15, min_row=hr2, max_row=lr2),
            cats=Reference(ws, min_col=13, min_row=hr2 + 1, max_row=lr2),
            title="Por que cambia el numero: el mismo calculo, con el plazo real",
            y_title="Unidades", colors=(C_DECLARADO, C_OBSERVADO), width=18, height=alto, gap=60,
        ), f"Q{2 + rows_below(alto)}")

    _widths(ws, {"A": 22, "B": 26, "C": 24, "D": 26, "E": 24, "F": 30, "G": 30,
                 "H": 16, "I": 14, "J": 14, "K": 24})


def _sheet_calidad(ws, value: GapValue) -> None:
    _widths(ws, {"A": 52, "B": 16, "C": 18, "D": 44, "E": 3})
    _cell(ws, 1, 1, "De donde salieron estos numeros", font=_SECTION)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=4)
    _cell(ws, 2, 1,
          "Una cifra de plata construida sobre una parte del archivo tiene que decir que parte era.",
          font=_NOTE, align=_WRAP)

    def share(n: int) -> float | str:
        return n / value.n_demand_skus if value.n_demand_skus else NO_MEDIDO

    _cell(ws, 4, 1, "Cuantos SKU quedaron detras de la cifra", font=Font(bold=True))
    _headers(ws, 5, ["Paso", "SKU", "Parte del archivo de movimientos"])
    bold = Font(bold=True)
    filas = [
        ("1. SKU distintos en tu archivo de movimientos", value.n_demand_skus, None),
        ("2. De esos, los que tus recepciones identifican con un proveedor",
         value.n_matched_skus, None),
    ]
    row = 6
    for etiqueta, n, font in filas:
        _cell(ws, row, 1, etiqueta, font=font)
        _cell(ws, row, 2, n, fmt=FMT_COUNT, font=font)
        _cell(ws, row, 3, share(n), fmt=FMT_PCT_1, font=font)
        row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    _cell(ws, row, 1,
          "De esos, cuantos cumplen cada requisito. Son tres pruebas EN PARALELO, no pasos "
          "encadenados: un SKU puede tener plazo observado sin tener declarado.",
          font=_NOTE, align=_WRAP)
    row += 1
    for etiqueta, n in (
        ("Con plazo DECLARADO en el archivo de recepciones", value.n_with_declared),
        (f"Con plazo OBSERVADO (proveedor con {MIN_DELIVERIES}+ recepciones con fecha)",
         value.n_with_observed),
        ("Con VARIACION observada (proveedor con 2+ recepciones con fecha)", value.n_with_std),
    ):
        _cell(ws, row, 1, etiqueta)
        _cell(ws, row, 2, n, fmt=FMT_COUNT)
        _cell(ws, row, 3, share(n), fmt=FMT_PCT_1)
        row += 1
    _cell(ws, row, 1, "3. SKU valorizados: los que cumplen las tres a la vez", font=bold)
    _cell(ws, row, 2, len(value.gaps), fmt=FMT_COUNT, font=bold)
    _cell(ws, row, 3, share(len(value.gaps)), fmt=FMT_PCT_1, font=bold)
    row += 2
    # Fila propia: NO es un subconjunto del archivo de movimientos, y meterla en el
    # embudo dibujaria una barra mas alta que la primera.
    _cell(ws, row, 1, "SKU en tus recepciones que NO aparecen en el archivo de movimientos")
    _cell(ws, row, 2, value.n_receipt_skus - value.n_matched_skus, fmt=FMT_COUNT)
    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    _cell(ws, row, 1,
          "Suelen ser codigos distintos entre compras y ventas. El cruce de este escaneo es por "
          "codigo exacto; la homologacion la hace el Diagnostico.", font=_NOTE, align=_WRAP)
    row += 2

    hr, lr = _chart_block(
        ws, 1, "Datos del grafico: el embudo de SKU", ["Paso", "SKU"],
        [["En el archivo de movimientos", value.n_demand_skus],
         ["Con proveedor en las recepciones", value.n_matched_skus],
         ["Valorizados", len(value.gaps)]],
        start_col=6, fmt=FMT_COUNT,
    )
    ws.add_chart(_bar(
        data=Reference(ws, min_col=7, min_row=hr, max_row=lr),
        cats=Reference(ws, min_col=6, min_row=hr + 1, max_row=lr),
        title="De cuantos SKU salio la cifra de plata", kind="bar",
        colors=(C_OBSERVADO,), legend=None, width=17, height=7,
    ), "I1")

    q = value.intake_quality
    if q is not None:
        _cell(ws, row, 1, "Archivo de movimientos (demanda)", font=bold)
        row += 1
        _headers(ws, row, ["Concepto", "Filas", "Parte del archivo"])
        row += 1
        usadas = q.n_raw - q.n_dropped
        for etiqueta, n, parte in (
            ("Filas en el archivo", q.n_raw, 1.0 if q.n_raw else NO_MEDIDO),
            ("Filas usadas", usadas, usadas / q.n_raw if q.n_raw else NO_MEDIDO),
            ("Descartadas: sin fecha usable", q.n_dropped_bad_date,
             q.n_dropped_bad_date / q.n_raw if q.n_raw else NO_MEDIDO),
            ("Descartadas: cantidad ilegible", q.n_dropped_bad_quantity,
             q.n_dropped_bad_quantity / q.n_raw if q.n_raw else NO_MEDIDO),
            ("Descartadas: cantidad negativa", q.n_dropped_negative_quantity,
             q.n_dropped_negative_quantity / q.n_raw if q.n_raw else NO_MEDIDO),
            # Poblacion DISTINTA, no una parte de la de arriba.
            ("Filas agregadas por SKU y semana", q.n_canonical_rows, NO_MEDIDO),
        ):
            _cell(ws, row, 1, etiqueta)
            _cell(ws, row, 2, n, fmt=FMT_COUNT)
            _cell(ws, row, 3, parte, fmt=FMT_PCT_1)
            row += 1
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        _cell(ws, row, 1,
              "El motor agrupa los movimientos por SKU y semana antes de calcular. Las columnas "
              "rellenadas de abajo se cuentan sobre estas filas agregadas, no sobre las del archivo.",
              font=_NOTE, align=_WRAP)
        row += 2

        _cell(ws, row, 1, "Columnas que el motor relleno con un valor por defecto", font=bold)
        row += 1
        _headers(ws, row, ["Columna", "Celdas rellenadas", "Parte de las filas agregadas",
                           "Afecta la cifra de plata"])
        row += 1
        veredictos = {
            "lead_time_days": "No. Para los SKU valorizados el plazo sale de tus recepciones y "
                              "reemplaza este campo.",
            "lead_time_std": "No. Un 0 aqui es exactamente lo que hace tu ERP: es el escenario "
                             "contra el que se compara.",
            "unit_cost": "SI. Cada celda rellenada vale 1,0 y multiplica directo la columna 'Plata'.",
        }
        if q.synthesized:
            for nombre, n in q.synthesized:
                _cell(ws, row, 1, nombre)
                _cell(ws, row, 2, n, fmt=FMT_COUNT)
                _cell(ws, row, 3, n / q.n_canonical_rows if q.n_canonical_rows else NO_MEDIDO,
                      fmt=FMT_PCT_1)
                _cell(ws, row, 4, veredictos.get(nombre, "No afecta la cifra de plata de este escaneo."),
                      align=_WRAP)
                row += 1
        else:
            _cell(ws, row, 1, "Ninguna: todas las columnas venian en tu archivo")
            _cell(ws, row, 2, 0, fmt=FMT_COUNT)
            _cell(ws, row, 3, 0.0, fmt=FMT_PCT_1)
            _cell(ws, row, 4, "--")
            row += 1
        n_cost = dict(q.synthesized).get("unit_cost", 0)
        if 0 < n_cost < q.n_canonical_rows:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            _cell(ws, row, 1,
                  f"Atencion: {n_cost} de {q.n_canonical_rows} filas agregadas no traian costo "
                  "unitario y el motor las valorizo a 1,0. Los SKU afectados tienen una cifra de "
                  "'Plata' que no sale de tu maestro de articulos.",
                  font=Font(bold=True, color=C_ALERTA), align=_WRAP)
            row += 1
        row += 1

        reading = q.reading
        if reading is not None:
            _cell(ws, row, 1, "Como lo leyo Kern", font=bold)
            row += 1
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            _cell(ws, row, 1,
                  "El mapeo y lo que se le saco a tu planilla, para que puedas abrir tu archivo "
                  "y confirmar. No es un hallazgo de tu operacion: es como se leyo el dato.",
                  font=_NOTE, align=_WRAP)
            row += 1
            if reading.source_kind == "excel":
                lugar = (
                    f"Excel, hoja {reading.sheet}, encabezado en la fila "
                    f"{(reading.header_row or 0) + 1}"
                )
            else:
                lugar = "CSV, primera fila como encabezado"
            for etiqueta, valor in (
                ("Donde miro", lugar),
                ("Mapeo de columnas", reading.mapping_text()),
            ):
                _cell(ws, row, 1, etiqueta)
                _cell(ws, row, 2, valor, align=_WRAP)
                ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
                row += 1
            if reading.noteworthy_hygiene() and reading.hygiene is not None:
                h = reading.hygiene
                n_total = int(h.dropped.get("fila_de_total") or 0)
                n_blank = int(h.dropped.get("fila_vacia") or 0)
                n_filled = sum(h.filled.values()) if h.filled else 0
                for etiqueta, n, texto in (
                    ("Filas de TOTAL excluidas", n_total,
                     "Una fila de cierre no es una venta. Dejarla entrar infla el desvio."),
                    ("Filas en blanco excluidas", n_blank, "Separadores, no demanda."),
                    ("Celdas completadas desde combinadas", n_filled,
                     "Solo donde el libro declara una combinacion, nunca un ffill a ciegas."),
                ):
                    if n:
                        _cell(ws, row, 1, etiqueta)
                        _cell(ws, row, 2, n, fmt=FMT_COUNT)
                        _cell(ws, row, 3, texto, align=_WRAP)
                        row += 1
            row += 1

    _cell(ws, row, 1, "Archivo de recepciones", font=bold)
    row += 1
    _headers(ws, row, ["Motivo del descarte", "Filas", "Que significa"])
    row += 1
    d = value.dropped_receipts
    for motivo, n, texto in (
        ("Sin fecha usable", d.get("sin_fecha", 0),
         "La fecha de la orden o la de la recepcion no se pudo leer. Sin las dos no hay plazo que medir."),
        ("Recepcion anterior a su orden", d.get("plazo_negativo", 0),
         "Un plazo negativo es un dato roto, no un plazo de cero dias."),
        ("Sin proveedor", d.get("sin_proveedor", 0), "La fila no dice a quien se le compro."),
        (f"Proveedores con menos de {MIN_DELIVERIES} recepciones", value.n_receipts_below_min,
         f"Se leyeron bien, pero {value.n_suppliers_below_min} proveedor(es) no junta(n) "
         f"{MIN_DELIVERIES} recepciones. Con menos, la variacion es ruido."),
    ):
        _cell(ws, row, 1, motivo)
        _cell(ws, row, 2, n, fmt=FMT_COUNT)
        _cell(ws, row, 3, texto, align=_WRAP)
        row += 1
    fuera = sum(d.values()) + value.n_receipts_below_min
    for etiqueta, n in (
        ("Total fuera de los numeros", fuera),
        ("Recepciones usadas", value.n_receipts_used),
        ("Filas del archivo", fuera + value.n_receipts_used),
    ):
        _cell(ws, row, 1, etiqueta, font=bold)
        _cell(ws, row, 2, n, fmt=FMT_COUNT, font=bold)
        row += 1


def _sheet_metodo(ws, value: GapValue) -> None:
    _widths(ws, {"A": 46, "B": 92})
    q = value.intake_quality
    syn = dict(q.synthesized) if q else {}
    n_std = syn.get("lead_time_std", 0)
    if q and n_std >= q.n_canonical_rows:
        sigma = (" Tu archivo de movimientos no traia variacion del plazo, asi que el escenario de "
                 "tu sistema usa cero: es exactamente lo que hace un ERP.")
    elif n_std:
        sigma = (f" {n_std} de las filas agregadas no traian variacion del plazo y se rellenaron "
                 "con cero.")
    else:
        sigma = ""

    pct = f"{HOLDING_RATE * 100:.0f}%"
    bloques = [
        ("Como se calculo", [
            ("El metodo",
             "Se dimensiono el stock de seguridad de cada SKU dos veces con el mismo calculo: una "
             "con el plazo que declara tu sistema, otra con el plazo y la variacion que muestran "
             "tus recepciones. La diferencia, multiplicada por el costo unitario, es la columna "
             "'Plata'." + sigma),
            ("Nivel de servicio usado",
             "95%. Es el supuesto por defecto del escaneo, no una medicion de tu operacion."),
            ("Costo de mantener",
             f"{pct} anual sobre el valor del inventario: capital, bodegaje, seguro y obsolescencia. "
             "Es el supuesto por defecto del escaneo, no una medicion de tu operacion. El rango "
             "habitual es 18-28% y las dos cifras anuales escalan en proporcion directa: a la mitad "
             "de tasa, la mitad de ahorro y la mitad de costo."),
            ("Unidades",
             "Los plazos van en DIAS. El stock de seguridad va en UNIDADES. El motor trabaja "
             "internamente en semanas y ninguna cifra de este archivo esta en semanas."),
            ("Que SKU entraron",
             "Solo los que tienen las tres cosas: plazo declarado, plazo observado y variacion "
             "observada, y ademas aparecen en los dos archivos. El embudo completo esta en "
             "'Calidad de los datos'."),
        ]),
        ("Que NO se midio", [
            ("Quiebres y compras de emergencia",
             "Nadie los registra, y en repuestos suelen ser mayores que la exposicion de arriba."),
            ("SKU con dos proveedores",
             "Quedan asignados al que mas veces se los entrego. El plazo mezclado ponderado por "
             "volumen es trabajo del Diagnostico."),
            ("Homologacion de codigos",
             "El cruce entre los dos archivos es por codigo exacto. Si compras y ventas usan codigos "
             "distintos, esos SKU quedan fuera y estan contados en 'Calidad de los datos'."),
            ("Proveedores con plazo declarado en cero",
             "El campo del ERP trae un cero, que suele ser un campo sin configurar y no una promesa "
             "de entrega inmediata. Para esos SKU el escenario del sistema queda en cero stock de "
             "seguridad y la columna 'Plata' es todo el stock de seguridad, no una diferencia."),
            ("Lo que si baja el inventario total",
             "El pooling entre bodegas. Necesita la red completa y no estos dos archivos. Medir la "
             "dispersion del plazo, por si sola, casi siempre revela buffer que falta: no es un "
             "resultado sobre tu operacion, es como funciona la formula."),
            ("Estacionalidad",
             "La demanda se agrega por semana sobre el periodo que trae tu archivo. El escaneo no "
             "proyecta un ano."),
        ]),
        ("Convenciones", [
            ("Signo de 'Plata'",
             "Negativo: stock de seguridad que sobra, caja que se libera. Positivo: el que falta, "
             "capital que hay que poner."),
            ("Cambio neto", "La suma de los dos. Negativo significa que el inventario total baja."),
            ("Las dos cifras anuales",
             "El ahorro se calcula solo sobre la caja que se libera. El cambio neto del costo de "
             "mantener es ese ahorro restado del costo de cubrir lo que falta: ya lo contiene, no "
             "se suman."),
            ("SKU sin cambio",
             "Un SKU sin demanda o sin costo registrado en el periodo da el mismo stock de seguridad "
             "en los dos escenarios. Cuenta como SKU valorizado y no entra en ninguna de las dos "
             "monedas."),
            ("Moneda", NOTA_MONEDA),
            ("'no medido'", "El dato no venia en el archivo. Nunca significa cero."),
        ]),
        ("Aviso", [
            ("Solo lectura", "Ningun numero de este archivo se escribio en tus sistemas."),
            ("Siguiente paso", f"Diagnostico de Arranque: {_BASE_URL}{CTA_PATH}"),
        ]),
    ]
    row = 1
    for titulo, pares in bloques:
        _cell(ws, row, 1, titulo, font=_SECTION)
        row += 1
        for etiqueta, texto in pares:
            _cell(ws, row, 1, etiqueta)
            _cell(ws, row, 2, texto, align=_WRAP)
            ws.row_dimensions[row].height = 30
            row += 1
        row += 1


def write_gap_value_workbook(
    value: GapValue, path: str | Path, *, dataset_label: str, ts: str
) -> Path:
    """Escribe ``brecha_en_plata.xlsx``. Solo puede llamarse con ``value.ok``.

    ``ts`` entra por parametro y nunca se lee del reloj.
    """
    if not value.ok:
        raise ValueError("el libro no se construye con QA fallado: no hay titular que escribir")

    wb = Workbook()
    _sheet_resumen(wb.active, value, dataset_label=dataset_label, ts=ts)
    wb.active.title = "Resumen"
    _sheet_skus(wb.create_sheet("SKU por SKU"), value)
    _sheet_calidad(wb.create_sheet("Calidad de los datos"), value)
    _sheet_metodo(wb.create_sheet("Como leer esto"), value)

    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
