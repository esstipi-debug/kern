"""``escaneo_stock.xlsx`` -- el entregable de /api/demo-scan.

Segundo de los tres libros de /demo. Las cinco reglas que lo gobiernan estan en
el docstring de :mod:`webapp.demo_workbook`; las primitivas, en
:mod:`webapp.demo_workbook_base`.

La trampa propia de este libro, y la razon por la que casi todo el codigo de
abajo calcula dos columnas donde parecia bastar una:

    ``excess_value`` de un SKU EXCEDENTE es **solo lo que pasa del objetivo de
    cobertura**; ``on_hand_value`` es todo lo que ese SKU tiene.

De ahi salen dos consecuencias que un libro ingenuo rompe:

* ``total - dead_value - excess_value`` **no es** el valor de los SKU sanos:
  absorbe la parte dentro del objetivo de cada excedente. Medido, sobreestima en
  exactamente esa cantidad.
* poner una columna de plata al lado de una de conteo compara **poblaciones
  distintas** en la misma fila, y el lector suma dos numeros que no se suman.

Por eso la tabla del Resumen lleva las dos columnas separadas, cada una con su
propio total, y lo dice en una nota.

La segunda trampa es el balde "muerto". ``classify_sku`` marca DEAD cuando la
demanda diaria es <= 0 **o** cuando lleva mucho sin venderse, y la fecha de
ultima venta es una columna OPCIONAL del archivo. Decir "sin movimiento hace N
dias" sobre todo el balde afirma algo que el archivo puede no haber traido
nunca. La rama es recuperable: la cobertura es infinita si y solo si la demanda
es <= 0, asi que el balde se parte sin necesitar nada nuevo aguas arriba.
"""

from __future__ import annotations

import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font

from webapp.demo_scan import CTA_PATH, DemoScanResult
from webapp.demo_workbook_base import (
    _BASE_URL,
    _FILL_TARDE,
    _FILL_TEMPRANO,
    _NOTE,
    _SECTION,
    _TITLE,
    _WRAP,
    C_ALERTA,
    C_DECLARADO,
    C_NEUTRO,
    C_OBSERVADO,
    FMT_COUNT,
    FMT_MONEY,
    FMT_PCT,
    FMT_PCT_1,
    FMT_UNITS_1,
    MAX_TABLE_ROWS,
    TOP_N_SKU,
    _bar,
    _cap,
    _cell,
    _chart_block,
    _headers,
    _isnumber_rule,
    _link_row,
    _measured,
    _widths,
    rows_below,
)
from webapp.operator_profile import get_operator_profile

#: Los cortes de valor acumulado de la clasificacion ABC. Es el default del motor
#: (``jobs/abc_xyz_job.run``), que ``demo_scan`` invoca sin override. Se declara
#: aca para poder nombrarlo en la hoja; un test lo ata a la firma del motor para
#: que no se despeguen.
ABC_THRESHOLDS = (0.80, 0.95)

ESTADO_ES = {"healthy": "sano", "excess": "excedente", "dead": "muerto"}

#: Nota de moneda, identica en los tres libros: nada en /demo sabe en que moneda
#: esta el archivo del cliente.
NOTA_MONEDA = (
    "Los montos van en la misma moneda del costo unitario de tu archivo. El escaneo no "
    "convierte ni asume simbolo; el signo $ que aparezca en los parrafos de hallazgos es "
    "del texto, no una conversion."
)


class _Derivadas:
    """Todo lo que se calcula una vez y se pasa a las hojas, nunca por hoja."""

    def __init__(self, result: DemoScanResult) -> None:
        eo = result.eo
        self.v_healthy = sum(ln.on_hand_value for ln in eo.lines if ln.classification == "healthy")
        self.v_excess_total = sum(ln.on_hand_value for ln in eo.lines if ln.classification == "excess")
        dead = [ln for ln in eo.lines if ln.classification == "dead"]
        # cobertura infinita <=> demanda diaria <= 0, asi que esta particion
        # separa exactamente las dos ramas de classify_sku.
        self.dead_sin = [ln for ln in dead if not math.isfinite(ln.days_of_cover)]
        self.dead_idle = [ln for ln in dead if math.isfinite(ln.days_of_cover)]
        self.v_sin = sum(ln.on_hand_value for ln in self.dead_sin)
        self.v_idle = sum(ln.on_hand_value for ln in self.dead_idle)
        self.r_sin = sum(ln.excess_value for ln in self.dead_sin)
        self.r_idle = sum(ln.excess_value for ln in self.dead_idle)
        #: La parte de los SKU excedentes que SI cubre demanda real.
        self.v_excess_within = self.v_excess_total - eo.excess_value
        #: `on_hand` solo se comprueba finito, nunca su signo.
        self.n_negative = sum(1 for ln in eo.lines if ln.on_hand < 0)
        self.n_neg_cover = sum(
            1 for ln in eo.lines if math.isfinite(ln.days_of_cover) and ln.days_of_cover < 0
        )


def _sheet_resumen(ws, result: DemoScanResult, d: _Derivadas, *, dataset_label: str, ts: str) -> None:
    eo, abc, fin = result.eo, result.abc, result.fin
    _widths(ws, {"A": 36, "B": 22, "C": 16, "D": 22, "E": 52, "F": 3})

    _cell(ws, 1, 1, "Tu inventario, en una hoja", font=_TITLE)
    _cell(ws, 2, 1, "Archivo analizado:")
    _cell(ws, 2, 2, dataset_label)
    _cell(ws, 3, 1, "Fecha del analisis:")
    _cell(ws, 3, 2, ts)
    _cell(ws, 4, 1, "SKU analizados:")
    _cell(ws, 4, 2, eo.n_skus, fmt=FMT_COUNT)

    _cell(ws, 6, 1, "El numero", font=_SECTION)
    _cell(ws, 7, 1, "Plata detenida en stock muerto o excedente")
    _cell(ws, 7, 2, eo.eo_value, fmt=FMT_MONEY, font=Font(bold=True, size=20, color=C_ALERTA))
    _cell(ws, 8, 1, "Sobre el valor total del stock")
    _cell(ws, 8, 2, eo.eo_pct_of_value, fmt=FMT_PCT, font=Font(bold=True, size=14, color=C_ALERTA))

    _cell(ws, 10, 1, "Como se reparte el valor del stock", font=_SECTION)
    _headers(ws, 11, ["Estado", "Valor en stock", "SKU en esa clase",
                      "De ese valor, en riesgo", "Que significa"])

    filas: list[list] = [
        ["Sano", d.v_healthy, eo.n_healthy, 0.0,
         "Rota dentro de lo esperado. Nada de este valor esta marcado en riesgo."],
        ["Excedente", d.v_excess_total, eo.n_excess, eo.excess_value,
         f"Sigue moviendose, pero hay stock para mas de {eo.target_cover_days:.0f} dias. Solo se "
         f"cuenta como riesgo lo que pasa de ese objetivo; el resto si cubre demanda real y esta "
         "en la columna de la izquierda."],
    ]
    if d.dead_sin:
        filas.append(["Muerto: sin demanda en el archivo", d.v_sin, len(d.dead_sin), d.r_sin,
                      "La columna de demanda diaria viene en cero o menos, asi que no hay salida "
                      "que proyectar. Esta fila no dice hace cuanto estan detenidos."])
    if d.dead_idle:
        filas.append([f"Muerto: sin movimiento hace {eo.dead_threshold_days:.0f} dias o mas",
                      d.v_idle, len(d.dead_idle), d.r_idle,
                      "El archivo trae la fecha de ultima venta de estos SKU y esta pasada del umbral."])

    row = 12
    for f in filas:
        _cell(ws, row, 1, f[0])
        _cell(ws, row, 2, f[1], fmt=FMT_MONEY)
        _cell(ws, row, 3, f[2], fmt=FMT_COUNT)
        _cell(ws, row, 4, f[3], fmt=FMT_MONEY)
        _cell(ws, row, 5, f[4], align=_WRAP)
        row += 1
    bold = Font(bold=True)
    _cell(ws, row, 1, "Total", font=bold)
    _cell(ws, row, 2, eo.total_on_hand_value, fmt=FMT_MONEY, font=bold)
    _cell(ws, row, 3, eo.n_skus, fmt=FMT_COUNT, font=bold)
    _cell(ws, row, 4, eo.eo_value, fmt=FMT_MONEY, font=bold)
    total_row = row

    # El bloque vive lejos a la derecha para que el grafico entre en la primera
    # pantalla, debajo de la tabla que lo explica.
    header_row, last_row = _chart_block(
        ws, 1,
        "Datos del grafico de esta hoja: valor del stock y parte en riesgo, por estado",
        ["Estado", "Valor en stock", "De ese valor, en riesgo"],
        [[f[0], f[1], f[3]] for f in filas],
        start_col=14, fmt=FMT_MONEY,
    )
    ch = _bar(
        data=Reference(ws, min_col=15, max_col=16, min_row=header_row, max_row=last_row),
        cats=Reference(ws, min_col=14, min_row=header_row + 1, max_row=last_row),
        title="Donde esta la plata de tu inventario",
        y_title="Valor, en la moneda de tu archivo",
        colors=(C_NEUTRO, C_ALERTA), width=19, height=9, gap=60,
    )
    ws.add_chart(ch, f"A{total_row + 2}")
    row = total_row + 2 + rows_below(9.0)

    for nota in (
        "Las dos columnas de plata miden cosas distintas y no se suman entre si: 'Valor en stock' "
        "es todo lo que tiene esa clase y cierra en el valor total; 'en riesgo' es solo la parte "
        "que sobra y cierra en la plata detenida. Puedes comprobar las dos sumando las columnas D "
        "y F de la hoja 'Plata en riesgo'.",
        (f"Tu archivo trae {d.n_negative} SKU con unidades en stock negativas. Estan sumados tal "
         "como vinieron y estan listados en 'Plata en riesgo'; el grafico no puede dibujar una "
         "barra bajo cero.") if d.n_negative else None,
    ):
        if not nota:
            continue
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        _cell(ws, row, 1, nota, font=_NOTE, align=_WRAP)
        ws.row_dimensions[row].height = 42
        row += 1

    row += 1
    _cell(ws, row, 1, "Otros numeros", font=_SECTION)
    row += 1
    _headers(ws, row, ["Indicador", "Valor", "Que quiere decir"])
    row += 1
    otros = [
        ("Dias de inventario (DIO)", _measured(fin.dio, digits=0), FMT_COUNT,
         "Cuantos dias tarda en salir el stock de hoy, al ritmo de la demanda de tu archivo."),
        ("Rotacion de inventario (veces al ano)", _measured(fin.turns), FMT_UNITS_1,
         "Cuantas veces se renueva el inventario en un ano, al mismo ritmo."),
        ("Salida sobre stock disponible (sell-through)", _measured(fin.sell_through, digits=4),
         FMT_PCT, "De cada 100 unidades entre lo que sale en un ano y lo que hay hoy, cuantas salen."),
        ("SKU clase A", abc.n_a, FMT_COUNT,
         f"Concentran {abc.a_value_share * 100:.0f}% del valor anual del catalogo."),
    ]
    for etiqueta, valor, fmt, texto in otros:
        _cell(ws, row, 1, etiqueta)
        _cell(ws, row, 2, valor, fmt=fmt)
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
        _cell(ws, row, 3, texto, align=_WRAP)
        row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    _cell(ws, row, 1,
          "Los tres se calculan proyectando la columna de demanda diaria a 365 dias. Una foto de "
          "stock no tiene estacionalidad: describen el ritmo de hoy, no el del ano. El Diagnostico "
          "los recalcula sobre el historial real.", font=_NOTE, align=_WRAP)
    ws.row_dimensions[row].height = 30
    row += 2

    _cell(ws, row, 1, "Que dice esto", font=_SECTION)
    row += 1
    for finding in result.findings:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        _cell(ws, row, 1, finding, align=_WRAP)
        ws.row_dimensions[row].height = 42
        row += 1

    # Los hallazgos vienen del .md y se pegan tal cual para que los dos artefactos
    # no se contradigan. Cuando uno de ellos afirma de mas, se corrige AQUI en vez
    # de reescribirlo: la correccion queda al lado de la frase que corrige.
    precisiones = []
    if d.dead_sin and eo.n_dead:
        precisiones.append(
            f"De los {eo.n_dead} SKU muertos, {len(d.dead_sin)} lo son porque su demanda diaria "
            "viene en cero o menos, no porque el archivo diga hace cuanto no se venden.")
    if abc.n_cz:
        precisiones.append(
            "La sigla CZ de arriba mezcla valor bajo con variabilidad de demanda, y una foto de "
            "stock no permite medir variabilidad. Leelo como bajo valor y sin demanda registrada, "
            "no como una recomendacion de descontinuar.")
    if precisiones:
        _cell(ws, row, 1, "Precisiones sobre lo anterior", font=_SECTION)
        row += 1
        for p in precisiones:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
            _cell(ws, row, 1, p, font=_NOTE, align=_WRAP)
            ws.row_dimensions[row].height = 42
            row += 1

    row += 1
    _cell(ws, row, 1, "Siguiente paso", font=_SECTION)
    row += 1
    _link_row(ws, row, 1, "Diagnostico de Arranque (abrir)", f"{_BASE_URL}{CTA_PATH}")
    row += 1
    operator = get_operator_profile()
    if operator.email:
        _cell(ws, row, 1, "Responde a:")
        _cell(ws, row, 2, f"{operator.name} <{operator.email}>")
        row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    _cell(ws, row, 1,
          "Este archivo lo genero un escaneo automatico sobre el que subiste, en solo lectura: no "
          "se escribio nada en tus sistemas. 'no medido' significa que el dato no venia o no se "
          f"puede calcular, nunca que el valor sea cero. {NOTA_MONEDA}",
          font=_NOTE, align=_WRAP)
    ws.row_dimensions[row].height = 60


def _sheet_riesgo(ws, result: DemoScanResult) -> None:
    eo = result.eo
    _cell(ws, 1, 1, "Cada SKU del archivo, ordenado por la plata que tiene detenida", font=_SECTION)
    _cell(ws, 2, 1,
          f"'Plata en riesgo' es la parte que sobra: en un SKU excedente es solo lo que pasa de "
          f"{eo.target_cover_days:.0f} dias de cobertura; en uno muerto es todo lo que hay. 'Valor "
          "en stock' es siempre todo lo que ese SKU tiene.", font=_NOTE, align=_WRAP)

    rows = _cap(list(eo.lines), ws, 3, 1, len(eo.lines), "plata en riesgo")
    _headers(ws, 4, ["SKU", "Estado", "Unidades en stock", "Valor en stock", "Unidades de mas",
                     "Plata en riesgo", "Dias de cobertura", "Que hacer"])
    for r, line in enumerate(rows, 5):
        _cell(ws, r, 1, line.product_id)
        _cell(ws, r, 2, ESTADO_ES[line.classification])
        _cell(ws, r, 3, line.on_hand, fmt=FMT_UNITS_1)
        _cell(ws, r, 4, line.on_hand_value, fmt=FMT_MONEY)
        _cell(ws, r, 5, line.excess_units, fmt=FMT_UNITS_1)
        _cell(ws, r, 6, line.excess_value, fmt=FMT_MONEY)
        _cell(ws, r, 7, _measured(line.days_of_cover), fmt=FMT_UNITS_1)
        _cell(ws, r, 8, line.recommended_action, align=_WRAP)
    last = 4 + len(rows)
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:H{last}"

    if rows:
        ref, rule = _isnumber_rule("G", 5, last, f">{eo.target_cover_days:g}", _FILL_TEMPRANO)
        ws.conditional_formatting.add(ref, rule)
        # Sin guardia de tipo a proposito, y la razon va escrita para que nadie lo
        # "arregle": es una IGUALDAD contra un literal del desarrollador sobre un
        # rango que solo puede contener los tres valores de ESTADO_ES.
        ws.conditional_formatting.add(
            f"B5:B{last}", FormulaRule(formula=['$B5="muerto"'], fill=_FILL_TARDE)
        )

    for i, texto in enumerate((
        f"Fondo naranja en 'Dias de cobertura': el stock alcanza para mas de "
        f"{eo.target_cover_days:.0f} dias, que es el objetivo del escaneo.",
        "Fondo rojo en 'Estado': SKU muerto. Las dos razones por las que un SKU cae ahi estan en "
        "'Como leer esto'.",
        "'no medido' en 'Dias de cobertura': el SKU no tiene demanda en el archivo, asi que no hay "
        "entre que dividir. No son cero dias de stock; es lo contrario.",
        "Las columnas 'Valor en stock' y 'Plata en riesgo' miden cosas distintas: la primera suma "
        "el total de la hoja Resumen, la segunda suma la plata detenida.",
    )):
        _cell(ws, last + 2 + i, 1, texto, font=_NOTE)

    top = [ln for ln in eo.lines if ln.excess_value > 0][:TOP_N_SKU]
    if top:
        header_row, last_row = _chart_block(
            ws, 1, f"Datos del grafico: los {len(top)} SKU con mas plata detenida",
            ["SKU", "Plata en riesgo"],
            [[ln.product_id, round(ln.excess_value, 2)] for ln in top],
            start_col=10, fmt=FMT_MONEY,
        )
        ch = _bar(
            data=Reference(ws, min_col=11, min_row=header_row, max_row=last_row),
            cats=Reference(ws, min_col=10, min_row=header_row + 1, max_row=last_row),
            title=f"Los {len(top)} SKU con mas plata detenida",
            kind="bar", colors=(C_ALERTA,), width=17,
            height=min(3 + 0.6 * len(top), 12), legend=None,
        )
        ws.add_chart(ch, "M2")

    _widths(ws, {"A": 24, "B": 13, "C": 18, "D": 18, "E": 17, "F": 18, "G": 19, "H": 46})


def _sheet_abc(ws, result: DemoScanResult) -> None:
    abc = result.abc
    _cell(ws, 1, 1, "Donde se concentra el valor anual de tu catalogo", font=_SECTION)
    _cell(ws, 2, 1,
          "El valor anual de un SKU es su costo unitario por su demanda diaria proyectada a 365 "
          f"dias. Clase A: los SKU que acumulan el primer {ABC_THRESHOLDS[0]:.0%} del valor; B: "
          f"hasta {ABC_THRESHOLDS[1]:.0%}; C: el resto. Los cortes son los del motor, no una "
          "eleccion de esta planilla.", font=_NOTE, align=_WRAP)
    plano = not abc.value_basis.discriminates
    if plano:
        # Requerido, no cosmetico: src/classification.py llama a no decirlo
        # "una mentira silenciosa".
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=4)
        _cell(ws, 3, 1,
              "Todos los SKU de tu archivo traen el mismo costo unitario, asi que esta "
              "clasificacion ordena por cantidad y no por plata.", font=_NOTE, align=_WRAP)

    rows = _cap(list(abc.classifications), ws, 4, 1, abc.n_skus, "valor anual")
    _headers(ws, 5, ["SKU", "Valor anual", "Acumulado", "Clase"])
    for r, c in enumerate(rows, 6):
        _cell(ws, r, 1, c.product_id)
        _cell(ws, r, 2, c.annual_value, fmt=FMT_MONEY)
        _cell(ws, r, 3, c.cumulative_share, fmt=FMT_PCT_1)
        _cell(ws, r, 4, c.abc)
    last = 5 + len(rows)
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A5:D{last}"
    _cell(ws, last + 2, 1,
          f"Clase A, B y C se cortan sobre el valor anual acumulado en {ABC_THRESHOLDS[0]:.0%} y "
          f"{ABC_THRESHOLDS[1]:.0%}. Un SKU que cruza un corte se queda en la clase mas importante.",
          font=_NOTE)

    total_annual = sum(c.annual_value for c in abc.classifications)
    if total_annual > 0 and abc.n_skus:
        titulo = "Datos del grafico: cuanto pesa cada clase"
        if plano:
            titulo += " -- ordenado por cantidad, no por plata: todos los SKU traen el mismo costo"
        bloque = []
        for cls in ("A", "B", "C"):
            miembros = [c for c in abc.classifications if c.abc == cls]
            bloque.append([cls, len(miembros) / abc.n_skus,
                           sum(c.annual_value for c in miembros) / total_annual])
        header_row, last_row = _chart_block(
            ws, 1, titulo, ["Clase", "Parte de los SKU", "Parte del valor anual"],
            bloque, start_col=6, fmt=FMT_PCT_1,
        )
        ch = _bar(
            data=Reference(ws, min_col=7, max_col=8, min_row=header_row, max_row=last_row),
            cats=Reference(ws, min_col=6, min_row=header_row + 1, max_row=last_row),
            title="Pocos SKU concentran casi todo el valor",
            colors=(C_NEUTRO, C_OBSERVADO), width=17, height=9, gap=60,
        )
        ch.y_axis.number_format = FMT_PCT_1
        ch.y_axis.scaling.max = 1
        ws.add_chart(ch, "J2")

    _widths(ws, {"A": 24, "B": 18, "C": 13, "D": 8})


def _sheet_cobertura(ws, result: DemoScanResult, d: _Derivadas, top: list) -> None:
    eo = result.eo
    k = len(top)
    _cell(ws, 1, 1, f"Los {k} SKU con mas dias de stock", font=_SECTION)
    texto = (f"Cuantos dias alcanza el stock actual al ritmo de la demanda de tu archivo, contra el "
             f"objetivo de {eo.target_cover_days:.0f} dias que usa el escaneo. Solo entran los SKU "
             "con demanda: uno sin salidas no tiene dias de cobertura y esta en la hoja "
             "'Sin demanda'.")
    if d.n_neg_cover:
        texto += (f" Tu archivo trae ademas {d.n_neg_cover} SKU con cobertura negativa, por unidades "
                  "en stock bajo cero; tampoco entran al grafico.")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=4)
    _cell(ws, 2, 1, texto, font=_NOTE, align=_WRAP)
    ws.row_dimensions[2].height = 44

    # Esta hoja ES su propio bloque, y no viola la regla 3 por construccion: no
    # puede aparecer NO_MEDIDO (los inf estan filtrados), no puede dispararse el
    # tope (k <= TOP_N_SKU), y NO SE LE PONE AUTOFILTRO. Eso ultimo es carga:
    # agregarlo recrea el peligro que la regla 3 existe para evitar.
    header_row, last_row = _chart_block(
        ws, 4, f"Datos del grafico: los {k} SKU con mas dias de stock",
        ["SKU", "Dias de cobertura", "Objetivo (d)"],
        [[ln.product_id, round(ln.days_of_cover, 1), eo.target_cover_days] for ln in top],
        start_col=1, fmt=FMT_UNITS_1,
    )
    ch = _bar(
        data=Reference(ws, min_col=2, max_col=3, min_row=header_row, max_row=last_row),
        cats=Reference(ws, min_col=1, min_row=header_row + 1, max_row=last_row),
        title=f"Los {k} SKU con mas dias de stock, contra el objetivo de {eo.target_cover_days:.0f} d",
        y_title="Dias", colors=(C_OBSERVADO, C_DECLARADO), width=20, height=10, gap=40,
    )
    ws.add_chart(ch, "E4")
    _widths(ws, {"A": 24, "B": 20, "C": 14, "D": 3})


def _sheet_sin_demanda(ws, d: _Derivadas) -> None:
    _cell(ws, 1, 1, "SKU sin demanda en el archivo", font=_SECTION)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=5)
    _cell(ws, 2, 1,
          "Estos SKU traen cero (o menos) en la columna de demanda diaria, asi que no hay entre "
          "que dividir: no tienen dias de cobertura y no entran en el grafico de 'Cobertura'. El "
          "escaneo los cuenta como muertos por eso, no porque el archivo diga cuanto tiempo llevan "
          "detenidos.", font=_NOTE, align=_WRAP)
    ws.row_dimensions[2].height = 44

    rows = _cap(list(d.dead_sin), ws, 3, 1, len(d.dead_sin), "plata en riesgo")
    _headers(ws, 4, ["SKU", "Unidades en stock", "Valor en stock", "Estado", "Que hacer"])
    for r, line in enumerate(rows, 5):
        _cell(ws, r, 1, line.product_id)
        _cell(ws, r, 2, line.on_hand, fmt=FMT_UNITS_1)
        _cell(ws, r, 3, line.on_hand_value, fmt=FMT_MONEY)
        _cell(ws, r, 4, ESTADO_ES[line.classification])
        _cell(ws, r, 5, line.recommended_action, align=_WRAP)
    last = 4 + len(rows)
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:E{last}"
    bold = Font(bold=True)
    # Sobre TODOS, no solo los mostrados, para que una hoja topada siga cuadrando
    # contra la fila del Resumen.
    _cell(ws, last + 1, 1, f"Total de los {len(d.dead_sin)} SKU sin demanda", font=bold)
    _cell(ws, last + 1, 3, d.v_sin, fmt=FMT_MONEY, font=bold)
    _widths(ws, {"A": 24, "B": 20, "C": 18, "D": 13, "E": 46})


def _sheet_metodo(ws, result: DemoScanResult, d: _Derivadas, *, dataset_label: str, ts: str) -> None:
    eo, abc = result.eo, result.abc
    _widths(ws, {"A": 44, "B": 92})

    n_sin, n_idle = len(d.dead_sin), len(d.dead_idle)
    if eo.n_dead == 0:
        detenido = "Ningun SKU quedo marcado como muerto en este archivo."
    elif n_idle == 0:
        detenido = ("Ningun SKU quedo marcado como muerto por antiguedad. Los de esa fila estan ahi "
                    "porque su demanda diaria viene en cero o menos, no porque el archivo diga hace "
                    "cuanto no se venden.")
    elif n_sin:
        detenido = (f"Solo se pudo medir en los {n_idle} SKU cuya fila traia la fecha de ultima "
                    f"venta pasada del umbral. Los otros {n_sin} estan marcados como muertos por no "
                    "tener demanda, y de esos el archivo no dice hace cuanto estan detenidos.")
    else:
        detenido = (f"Se midio en los {n_idle} SKU cuya fila traia la fecha de ultima venta. Es una "
                    "columna opcional del archivo.")

    bloques = [
        ("Como se midio", [
            ("Origen de los datos", f"{dataset_label}. Una sola foto de stock: una fila por SKU con "
             "unidades, demanda diaria y costo unitario."),
            ("Fecha del analisis", ts),
            ("SKU analizados", f"{eo.n_skus}. Es el numero de FILAS de tu archivo: si un SKU aparece "
             "dos veces, se cuenta dos veces y aparece dos veces en las tablas."),
            ("Valor en stock", "Unidades por costo unitario, SKU por SKU."),
            ("Plata en riesgo", f"En un SKU excedente es SOLO lo que pasa del objetivo de "
             f"{eo.target_cover_days:.0f} dias de cobertura; en uno muerto es todo lo que hay. Por "
             "eso 'Valor en stock' y 'Plata en riesgo' no son la misma cifra ni miden lo mismo, y "
             "no se suman entre si."),
            ("Objetivo de cobertura", f"{eo.target_cover_days:.0f} dias. Es el supuesto por defecto "
             "del escaneo, no una medicion de tu operacion."),
            ("Umbral de stock muerto", f"{eo.dead_threshold_days:.0f} dias sin movimiento. Solo se "
             "puede aplicar a los SKU cuya fila trae la fecha de ultima venta; es una columna opcional."),
            ("Dias de inventario, rotacion y sell-through", "Se calculan proyectando la columna de "
             "demanda diaria a 365 dias. Una foto de stock no tiene estacionalidad: describen el "
             "ritmo de hoy, no el del ano."),
            ("Clase A, B y C", f"Se ordena por valor anual y se corta en {ABC_THRESHOLDS[0]:.0%} y "
             f"{ABC_THRESHOLDS[1]:.0%} de valor acumulado. Son los cortes del motor."),
        ]),
        ("Que NO se midio", [
            ("Cuanto tiempo lleva detenido cada SKU", detenido),
            ("Clasificacion XYZ (variabilidad de la demanda)", "No se calcula. Una foto de stock "
             "trae un solo punto de demanda por SKU, asi que la variabilidad no es medible con este "
             "archivo. Va en el Diagnostico."),
            ("Candidatos a descontinuar", "Dependen de la variabilidad de la demanda, que este "
             "archivo no permite medir. No se reportan aqui."),
            ("Margen y retorno sobre el inventario (GMROI)", "El archivo no trae margen. No se "
             "reporta: un cero aqui diria que tu margen es cero, y eso no es lo que sabemos."),
            ("Ventas netas, cobros, pagos, ciclo de caja", "Necesitan los libros de clientes y "
             "proveedores. Fuera del alcance del escaneo gratuito."),
            ("Rotacion por bodega", "El archivo es una sola foto consolidada."),
            ("Plazo de entrega de tus proveedores", "Ese es el otro escaneo de esta pagina, sobre "
             "tu archivo de recepciones."),
        ]),
        ("Convenciones", [
            ("'no medido'", "El dato no venia en el archivo o no se puede calcular. Nunca significa "
             "que el valor sea cero."),
            ("'no medido' en 'Dias de cobertura'", "El SKU no tiene demanda registrada: no hay entre "
             "que dividir. No son cero dias de stock, es lo contrario."),
            ("Moneda", NOTA_MONEDA),
            ("Orden de las tablas", "'Plata en riesgo' viene ordenada por plata detenida de mayor a "
             "menor; 'Concentracion ABC' por valor anual."),
            ("Filas mostradas", f"Las tablas de detalle se topan en {MAX_TABLE_ROWS:,} filas. Cuando "
             "eso pasa, la propia hoja lo dice arriba de la tabla y el resto entra completo en el "
             "Diagnostico.".replace(",", ".")),
        ]),
        ("Aviso", [
            ("Solo lectura", "Ningun numero de este archivo se escribio en tus sistemas."),
            ("Siguiente paso", f"Diagnostico de Arranque: {_BASE_URL}{CTA_PATH}"),
        ]),
    ]
    row = 1
    for titulo, pares in bloques:
        _cell(ws, row, 1, titulo, font=_SECTION)
        row += 1
        for etiqueta, texto in pares:
            _cell(ws, row, 1, etiqueta)
            _cell(ws, row, 2, texto, align=_WRAP)
            ws.row_dimensions[row].height = 30
            row += 1
        row += 1
    _ = abc  # el bloque ABC se describe arriba; se deja el binding por simetria


def write_stock_workbook(
    result: DemoScanResult, path: str | Path, *, dataset_label: str, ts: str
) -> Path:
    """Escribe ``escaneo_stock.xlsx``. Solo puede llamarse con ``result.ok``.

    ``ts`` entra por parametro y nunca se lee del reloj: dos corridas sobre el
    mismo resultado tienen que producir el mismo archivo.
    """
    if not result.ok:
        raise ValueError("el libro no se construye con QA fallado: no hay titular que escribir")

    d = _Derivadas(result)
    wb = Workbook()
    _sheet_resumen(wb.active, result, d, dataset_label=dataset_label, ts=ts)
    wb.active.title = "Resumen"
    _sheet_riesgo(wb.create_sheet("Plata en riesgo"), result)
    _sheet_abc(wb.create_sheet("Concentracion ABC"), result)

    finitos = [ln for ln in result.eo.lines
               if math.isfinite(ln.days_of_cover) and ln.days_of_cover >= 0]
    top = sorted(finitos, key=lambda ln: ln.days_of_cover, reverse=True)[:TOP_N_SKU]
    if top:
        _sheet_cobertura(wb.create_sheet("Cobertura"), result, d, top)
    if d.dead_sin:
        _sheet_sin_demanda(wb.create_sheet("Sin demanda"), d)
    _sheet_metodo(wb.create_sheet("Como leer esto"), result, d,
                  dataset_label=dataset_label, ts=ts)

    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
