"""Los jobs que leen stock escriben la memoria del Tower (D1 / N1c).

``scm_agent/monitors.py::run_all_monitors`` lee ``stock``, ``forecast`` y ``outcomes``, y
el UNICO escritor de ``stock`` era el puente del marcador
(``examples/run_world_sim.py``). En produccion nadie escribia: el bucle de sensado del
Control Tower corria sobre vacio y no podia emitir un evento aunque el quiebre estuviera
pasando -- hallazgo H1 de ``documentation/KERN_DISRUPCIONES_Y_DESAJUSTES.md``.

Lo que este archivo fija, en el orden en que importa:

1. **El Tower sensa.** Correr un tool de conector con ``state_client_id`` y despues
   ``run_all_monitors`` sobre esa particion emite un evento real. Es la prueba de punta
   a punta; todo lo demas es como se llega.
2. **El mapeo de cada job es el declarado**, y los valores SE MUEVEN cuando cambia el
   insumo -- la disciplina de ``tests/test_neteo_en_transito.py``.
3. **La identidad no se adivina.** Sin ``state_client_id`` no se escribe nada, y NUNCA
   la particion compartida.
4. **La ausencia se declara.** Una linea sin los tres numeros no se escribe con ceros.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

_REPO_ROOT = Path(__file__).resolve().parents[1]
if str(_REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(_REPO_ROOT))

from jobs import state_bridge  # noqa: E402
from nosilo_core.registry import JobRequest  # noqa: E402
from scm_agent.monitors import run_all_monitors  # noqa: E402
from scm_agent.tools import build_default_registry  # noqa: E402
from src.connectors.odoo import demo_odoo  # noqa: E402
from src.connectors.replenish import daily_demand_by_sku, demand_by_sku  # noqa: E402
from src.connectors.shopify import InMemoryShopifyCommerce, demo_shopify  # noqa: E402
from src.state import latest  # noqa: E402
from src.state.store import ClientIdError, StateStore, store_for  # noqa: E402

LOC = "gid://shopify/Location/1"


def _correr(tool_key: str, params: dict, data_path: str | None = None):
    tool = build_default_registry().get(tool_key)
    prep = tool.prepare(JobRequest(brief="x", data_path=data_path, params=params), None)
    assert prep.status == "ok", prep.messages
    return tool.run(prep.payload, params)


def _stock(base_path, client_id: str = "acme"):
    snap = latest("stock", store=store_for(client_id, base_path=str(base_path)))
    return {} if snap is None else {r["product_id"]: r for r in snap.payload.to_dict("records")}


# -- 1. la prueba que vale: el Tower deja de sensar sobre vacio --------------------


def test_una_corrida_de_conector_hace_que_el_tower_emita_un_evento(tmp_path):
    """Sin esto no habia evento posible en produccion: nadie escribia ``stock``.

    La tienda demo deja SKU-3 con 12 en estante contra un objetivo de 40 y 5 unidades
    por dia: 2,4 dias de cobertura, o sea el caso urgente.
    """
    params = {"shopify_rpc": demo_shopify(), "state_client_id": "acme",
              "state_path": str(tmp_path)}
    _correr("shopify_replenishment", params)

    store = store_for("acme", base_path=str(tmp_path))
    eventos = run_all_monitors(store=store)
    por_tipo = {e.type: e for e in eventos}
    assert "stockout_projected" in por_tipo, [e.type for e in eventos]
    assert por_tipo["stockout_projected"].sku == "SKU-3"
    # Y el evento viaja con la fila del estado, que es lo que su constructor consume.
    assert por_tipo["stockout_projected"].payload["rows"][0]["product_id"] == "SKU-3"


def test_sin_corrida_previa_el_tower_no_emite_nada_y_no_revienta(tmp_path):
    """Una particion vacia es el dia uno, no un error."""
    assert run_all_monitors(store=store_for("nueva", base_path=str(tmp_path))) == []


# -- 2. el mapeo de cada job, y el valor que se mueve ------------------------------


def _tienda(on_hand: int, *, qty: int = 5, dias: int = 8) -> InMemoryShopifyCommerce:
    variante = {
        "id": "gid://shopify/ProductVariant/P1", "sku": "P1", "displayName": "P1",
        "price": "12.00",
        "inventoryItem": {"id": "gid://shopify/InventoryItem/P1",
                          "unitCost": {"amount": "5.00"}},
    }
    ordenes = [
        {"id": f"gid://shopify/Order/{d}", "createdAt": f"2026-01-{d:02d}T10:00:00Z",
         "lineItems": {"edges": [{"node": {
             "sku": "P1", "quantity": qty,
             "originalUnitPriceSet": {"shopMoney": {"amount": "12.00"}}}}]}}
        for d in range(1, dias + 1)
    ]
    return InMemoryShopifyCommerce(variants=[variante], orders=ordenes,
                                   stock={"P1": {LOC: on_hand}}, incoming={"P1": {LOC: 0}})


def test_el_estante_escrito_es_el_que_leyo_el_job(tmp_path):
    """Dos corridas sobre la MISMA tienda, cambiando solo el stock: el valor se mueve."""
    for i, on_hand in enumerate((40, 90)):
        base = tmp_path / f"c{i}"
        _correr("shopify_replenishment", {"shopify_rpc": _tienda(on_hand),
                                          "state_client_id": "acme", "state_path": str(base)})
        assert _stock(base)["P1"]["on_hand"] == float(on_hand)


def test_la_demanda_diaria_es_por_dia_calendario_y_no_por_dia_con_pedidos(tmp_path):
    """El error que este puente NO comete, y que costaria falsas alarmas.

    ``demand_by_sku`` agrega por fecha CON pedidos: su periodo es un dia-con-pedidos.
    Sobre un SKU intermitente -- aca, 11 pedidos de 10 repartidos cada 3 dias en una
    ventana de 31 -- el promedio por punto da 10/dia y la tasa real es 110/31 = 3,55/dia.
    Escribir la primera haria que el Tower calcule casi tres veces menos cobertura de la
    que hay, y dispare un quiebre inminente sobre stock que alcanza un mes.
    """
    variante = {"id": "gid://shopify/ProductVariant/Z", "sku": "Z", "displayName": "Z",
                "price": "10.00",
                "inventoryItem": {"id": "gid://shopify/InventoryItem/Z",
                                  "unitCost": {"amount": "4.00"}}}
    ordenes = [
        {"id": f"gid://shopify/Order/{d}", "createdAt": f"2026-01-{d:02d}T10:00:00Z",
         "lineItems": {"edges": [{"node": {
             "sku": "Z", "quantity": 10,
             "originalUnitPriceSet": {"shopMoney": {"amount": "10.00"}}}}]}}
        for d in range(1, 32, 3)      # 11 dias con pedidos en una ventana de 31 dias
    ]
    rpc = InMemoryShopifyCommerce(variants=[variante], orders=ordenes,
                                  stock={"Z": {LOC: 100}}, incoming={"Z": {LOC: 0}})
    _correr("shopify_replenishment", {"shopify_rpc": rpc, "state_client_id": "acme",
                                      "state_path": str(tmp_path)})
    fila = _stock(tmp_path)["Z"]
    assert fila["avg_daily_demand"] == pytest.approx(110 / 31, rel=1e-6)
    # El otro numero, el que NO se escribe, para que la diferencia quede a la vista.
    pedidos = rpc.list_orders()
    from src.connectors import Order, OrderLine
    convertidos = [
        Order(order_id=str(o["id"]), created_at=o["createdAt"],
              lines=tuple(OrderLine(sku=e["node"]["sku"], quantity=float(e["node"]["quantity"]),
                                    price=10.0) for e in o["lineItems"]["edges"]))
        for o in pedidos
    ]
    por_punto = sum(demand_by_sku(convertidos)["Z"]) / len(demand_by_sku(convertidos)["Z"])
    assert por_punto == 10.0
    assert daily_demand_by_sku(convertidos)["Z"] == pytest.approx(110 / 31, rel=1e-6)


def test_el_disparador_escrito_es_el_objetivo_del_plan(tmp_path):
    """Para un plan de cobertura, el objetivo ES el disparador: pide cuando la posicion
    cae por debajo. Escribir 0 diria que nunca hay que pedir."""
    _correr("shopify_replenishment", {"shopify_rpc": _tienda(40), "state_client_id": "acme",
                                      "state_path": str(tmp_path)})
    fila = _stock(tmp_path)["P1"]
    # 5 unidades/dia-con-pedidos x 8 periodos de cobertura
    assert fila["reorder_point"] == pytest.approx(40.0)
    assert fila["tool"] == "shopify_replenishment"


def test_el_tool_csv_escribe_solo_los_sku_con_posicion(tmp_path):
    """``inventory_optimization`` dimensiona todos los SKU y solo algunos traen stock:
    los que no, no tienen ``on_hand`` que escribir y se omiten en vez de salir en cero."""
    import numpy as np
    import pandas as pd

    rng = np.random.default_rng(1)
    fechas = pd.date_range("2025-01-06", periods=30, freq="W-MON")
    filas = []
    for sku, stock in (("A", 30), ("B", None)):
        for f in fechas:
            fila = {"date": f, "product_id": sku, "quantity": max(0, round(rng.normal(20, 4))),
                    "unit_cost": 5.0, "lead_time_days": 14}
            if stock is not None:
                fila["on_hand"] = stock
            filas.append(fila)
    csv = tmp_path / "demanda.csv"
    pd.DataFrame(filas).to_csv(csv, index=False)

    _correr("inventory_optimization",
            {"state_client_id": "acme", "state_path": str(tmp_path)}, data_path=str(csv))
    escrito = _stock(tmp_path)
    assert set(escrito) == {"A"}, "un SKU sin columna de stock no se escribe en cero"
    assert escrito["A"]["on_hand"] == 30.0
    # La tasa diaria sale del pronostico por periodo dividido los dias del periodo
    # (52 periodos al anio -> 7,02 dias), no del pronostico crudo.
    assert 0 < escrito["A"]["avg_daily_demand"] < 20.0


def test_la_planilla_en_modo_punto_de_reorden_no_inventa_una_demanda(tmp_path):
    """El unico de los cinco que puede quedarse sin escribir por diseno, y lo dice.

    Sin columna de demanda no hay tasa diaria, y el dominio la exige. Derivarla del
    punto de reorden seria inventarla.
    """
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja"
    for col, encabezado in enumerate(("Codigo", "Stock", "Punto Reorden"), 1):
        ws.cell(row=1, column=col, value=encabezado)
    for fila, (sku, stock, rop) in enumerate((("P1", 42, 50), ("P2", 30, 40)), 2):
        for col, valor in enumerate((sku, stock, rop), 1):
            ws.cell(row=fila, column=col, value=valor)
    ruta = tmp_path / "planilla.xlsx"
    wb.save(ruta)

    produced = _correr("excel_replenishment",
                       {"state_client_id": "acme", "state_path": str(tmp_path)},
                       data_path=str(ruta))
    assert _stock(tmp_path) == {}
    assert "NOT recorded" in produced.report.state_note
    assert "daily demand" in produced.report.state_note


def test_la_planilla_con_columna_de_demanda_si_escribe(tmp_path):
    """Y el contraste: con demanda semanal, la tasa diaria es esa dividida por 7."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja"
    for col, encabezado in enumerate(("Codigo", "Stock", "Demanda Semanal"), 1):
        ws.cell(row=1, column=col, value=encabezado)
    for col, valor in enumerate(("P1", 42, 14), 1):
        ws.cell(row=2, column=col, value=valor)
    ruta = tmp_path / "planilla.xlsx"
    wb.save(ruta)

    _correr("excel_replenishment",
            {"state_client_id": "acme", "state_path": str(tmp_path)}, data_path=str(ruta))
    fila = _stock(tmp_path)["P1"]
    assert fila["avg_daily_demand"] == pytest.approx(2.0)
    assert fila["on_hand"] == 42.0


def test_el_canal_escribe_el_estante_del_erp_y_no_el_del_canal(tmp_path):
    """El ERP es el sistema de registro del stock; donde los dos discrepan, el job ya
    emite una divergencia aparte y no la resuelve escribiendo la del canal."""
    params = {"shopify_rpc": demo_shopify(), "odoo_rpc": demo_odoo(),
              "state_client_id": "acme", "state_path": str(tmp_path)}
    produced = _correr("channel_replenishment", params)
    escrito = _stock(tmp_path)
    por_sku = {ln.sku: ln for ln in produced.report.lines}
    for sku, fila in escrito.items():
        assert fila["on_hand"] == por_sku[sku].erp_on_hand
        assert fila["tool"] == "channel_replenishment"


# -- 3. la identidad no se adivina -------------------------------------------------


def test_sin_state_client_id_no_se_escribe_nada_ni_en_la_particion_compartida(tmp_path):
    produced = _correr("shopify_replenishment", {"shopify_rpc": demo_shopify(),
                                                 "state_path": str(tmp_path)})
    assert "no state_client_id" in produced.report.state_note
    assert StateStore(str(tmp_path)).count_all() == 0
    # Y la nota NO va al resumen: con la memoria apagada por configuracion, una linea
    # en cada resumen seria ruido. El reporte la lleva igual.
    assert "no state_client_id" not in produced.summary


def test_un_client_id_invalido_levanta_en_vez_de_degradar(tmp_path):
    """Dos clientes no pueden terminar sanitizados en el mismo directorio: la regla es
    de ``store_for`` y este puente la hereda sin ablandarla."""
    with pytest.raises(ClientIdError):
        _correr("shopify_replenishment", {"shopify_rpc": demo_shopify(),
                                          "state_client_id": "  ", "state_path": str(tmp_path)})


def test_cada_cliente_escribe_en_su_propia_particion(tmp_path):
    for cliente, on_hand in (("uno", 40), ("dos", 90)):
        _correr("shopify_replenishment", {"shopify_rpc": _tienda(on_hand),
                                          "state_client_id": cliente,
                                          "state_path": str(tmp_path)})
    assert _stock(tmp_path, "uno")["P1"]["on_hand"] == 40.0
    assert _stock(tmp_path, "dos")["P1"]["on_hand"] == 90.0


# -- 4. la ausencia se declara -----------------------------------------------------


def test_una_linea_sin_los_tres_numeros_no_se_escribe_en_cero():
    """El filtro vive en el puente, no en cada job."""
    class _Linea:
        def __init__(self, **kw):
            self.__dict__.update({"status": "ok", "sku": "X", "on_hand": 1.0,
                                  "target": 2.0, "avg_daily_demand": 3.0, **kw})

    assert len(state_bridge.stock_lines_from_replenishment([_Linea()])) == 1
    for campo in ("on_hand", "target", "avg_daily_demand"):
        assert state_bridge.stock_lines_from_replenishment([_Linea(**{campo: None})]) == []
    assert state_bridge.stock_lines_from_replenishment(
        [_Linea(status="insufficient_history")]) == []


def test_una_demanda_diaria_de_cero_si_se_escribe():
    """Cero medido es una medicion: el SKU no vendio. Distinto de no haberlo medido."""
    class _Linea:
        status, sku, on_hand, target, avg_daily_demand = "ok", "X", 5.0, 10.0, 0.0

    lineas = state_bridge.stock_lines_from_replenishment([_Linea()])
    assert len(lineas) == 1 and lineas[0].avg_daily_demand == 0.0


def test_sin_ninguna_linea_usable_no_se_escribe_un_snapshot_vacio(tmp_path):
    """Un snapshot vacio se leeria como "no hay stock", que es otra afirmacion."""
    store = StateStore(str(tmp_path))
    nota = state_bridge.record_stock_snapshot(store, tool="x", lines=[], n_candidates=3)
    assert nota.store_named and not nota.recorded
    assert "none of the 3 line(s)" in nota.text
    assert store.count_all() == 0


def test_la_nota_cuenta_las_lineas_omitidas(tmp_path):
    store = StateStore(str(tmp_path))
    lineas = [state_bridge.StockLine("A", 1.0, 2.0, 3.0)]
    nota = state_bridge.record_stock_snapshot(store, tool="x", lines=lineas, n_candidates=4)
    assert nota.recorded and "3 line(s) omitted" in nota.text
    assert nota.cycle_id and nota.cycle_id in nota.text


def test_una_fecha_que_no_parsea_se_descarta_y_no_se_le_inventa_un_dia():
    """Un pedido con fecha ilegible sale del calculo. La alternativa -- ponerle hoy, o el
    dia anterior -- inventaria una observacion, y la tasa diaria es justo el numero con
    el que el Tower decide si un SKU esta por quebrar."""
    from src.connectors import Order, OrderLine

    def _pedido(fecha: str) -> Order:
        return Order(order_id=fecha, created_at=fecha,
                     lines=(OrderLine(sku="X", quantity=10.0, price=1.0),))

    buenos = [_pedido("2026-01-01"), _pedido("2026-01-11")]
    assert daily_demand_by_sku(buenos)["X"] == pytest.approx(20 / 11)
    # La misma ventana MAS un pedido con fecha imposible: el resultado no se mueve.
    assert daily_demand_by_sku([*buenos, _pedido("2026-01-32")])["X"] == pytest.approx(20 / 11)
    # Y un SKU cuyas fechas son TODAS ilegibles no entra al mapa: ausente, nunca 0.0.
    assert daily_demand_by_sku([_pedido("ayer")]) == {}


def test_un_solo_dia_con_pedidos_da_una_ventana_de_un_dia_y_no_una_division_por_cero():
    from src.connectors import Order, OrderLine

    pedido = Order(order_id="1", created_at="2026-01-05",
                   lines=(OrderLine(sku="X", quantity=7.0, price=1.0),))
    assert daily_demand_by_sku([pedido])["X"] == 7.0
