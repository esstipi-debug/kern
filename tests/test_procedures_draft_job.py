"""Tests for the procedure_draft job: request intake, the approval handoff, the QA
gate, and the client-ready draft deliverables."""

from dataclasses import replace
from datetime import date

import pytest
from nosilo_core.guided import HANDOFF

from jobs import procedure_draft_job
from src.procedures.templates import (
    ISO_SENTENCE_EN,
    ISO_SENTENCE_ES,
    PROCESS_KEYS,
    TEMPLATE_VERSION,
)

AS_OF_ISO = "2026-06-01"


def _request(**overrides):
    params = {
        "process": "sop_cycle",
        "owner": "Ana",
        "approver": "Beto",
        "as_of": AS_OF_ISO,
    }
    params.update(overrides)
    return procedure_draft_job.prepare(params)


# ---- prepare() ------------------------------------------------------------------


def test_prepare_needs_process_and_lists_the_three():
    with pytest.raises(ValueError) as exc_missing:
        procedure_draft_job.prepare({})
    message_missing = str(exc_missing.value)
    for key in PROCESS_KEYS:
        assert key in message_missing

    with pytest.raises(ValueError) as exc_unknown:
        procedure_draft_job.prepare({"process": "not_a_real_process"})
    message_unknown = str(exc_unknown.value)
    for key in PROCESS_KEYS:
        assert key in message_unknown


def test_prepare_defaults_cadence_365_labelled_convention():
    request = procedure_draft_job.prepare({"process": "sop_cycle"})
    assert request.review_cadence_days == 365
    assert request.cadence_source == "convention"

    request_with_param = procedure_draft_job.prepare({"process": "sop_cycle", "review_cadence_days": 180})
    assert request_with_param.review_cadence_days == 180
    assert request_with_param.cadence_source == "param"


def test_prepare_reads_service_level_and_lead_time_from_params():
    request = procedure_draft_job.prepare({
        "process": "replenishment",
        "owner": "Ana",
        "approver": "Beto",
        "as_of": AS_OF_ISO,
        "service_level": "98%",
        "lead_time_days": 10,
    })
    draft = procedure_draft_job.run(request)
    used = {name: (value, source) for name, value, source in draft.parameters_used}
    assert used["service_level"] == ("98%", "param")
    assert used["lead_time_days"] == ("10", "param")


# ---- run() ------------------------------------------------------------------------


def test_run_outcome_is_handoff_with_approver_step_and_next_review_deadline():
    draft = procedure_draft_job.run(_request())

    assert draft.outcome.status == HANDOFF
    assert draft.outcome.status != "executed"
    (packet,) = draft.outcome.handoffs
    assert packet.title == "Aprobar el borrador"
    assert len(packet.steps) == 5
    assert any("Ana" in step for step in packet.steps)
    assert any("Beto" in step and "aprob" in step.lower() for step in packet.steps)
    assert any("sistema donde viven los sop" in step.lower() for step in packet.steps)

    assert packet.deadline == draft.control.next_review.isoformat()
    assert packet.data["owner"] == "Ana"
    assert packet.data["approver"] == "Beto"
    assert packet.data["due_date"] == draft.control.next_review.isoformat()
    assert packet.data["template_version"] == TEMPLATE_VERSION
    assert packet.risk_if_skipped.strip()

    assert draft.control.approved_at is None
    assert draft.control.status == "draft"


def test_run_flags_owner_equal_approver_as_residual_not_qa_failure():
    draft = procedure_draft_job.run(_request(owner="Ana", approver="Ana"))

    matching = [r for r in draft.outcome.residuals if "segregacion" in r.description]
    assert len(matching) == 1
    assert matching[0].risk_if_skipped.strip()
    # Not a QA failure: verify() still passes -- this is a human residual, not a
    # protocol violation the machine can block on.
    assert procedure_draft_job.verify(draft) == []


def test_run_flags_blank_owner_or_approver_as_residual():
    draft = procedure_draft_job.run(_request(owner="", approver="Beto"))
    matching = [r for r in draft.outcome.residuals if "segregacion" in r.description]
    assert len(matching) == 1


def test_run_adds_clock_residual_when_as_of_defaults():
    request = procedure_draft_job.prepare({"process": "sop_cycle", "owner": "Ana", "approver": "Beto"})
    assert request.as_of_source == "clock"
    draft = procedure_draft_job.run(request)
    clock_residuals = [r for r in draft.outcome.residuals if "reloj" in r.description]
    assert len(clock_residuals) == 1
    assert clock_residuals[0].risk_if_skipped.strip()


def test_run_rejects_malformed_as_of():
    with pytest.raises(ValueError, match="as_of"):
        procedure_draft_job.prepare({"process": "sop_cycle", "as_of": "not-a-date"})


# ---- verify() -----------------------------------------------------------------------


def test_verify_passes_on_a_normal_draft():
    draft = procedure_draft_job.run(_request())
    assert procedure_draft_job.verify(draft) == []


def test_verify_rejects_a_draft_that_claims_approval():
    draft = procedure_draft_job.run(_request())
    tampered_control = replace(draft.control, approved_at=date(2026, 6, 2), status="approved")
    tampered = replace(draft, control=tampered_control)

    issues = procedure_draft_job.verify(tampered)
    assert any("approv" in i.lower() for i in issues)


def test_verify_flags_handoff_without_risk_if_skipped():
    draft = procedure_draft_job.run(_request())
    bare_packet = replace(draft.outcome.handoffs[0], risk_if_skipped="")
    tampered_outcome = replace(draft.outcome, handoffs=[bare_packet])
    tampered = replace(draft, outcome=tampered_outcome)

    issues = procedure_draft_job.verify(tampered)
    assert any("risk_if_skipped" in i and "handoff" in i.lower() for i in issues)


# ---- write_operational() -------------------------------------------------------------


def test_write_operational_creates_md_xlsx_and_three_csvs(tmp_path):
    draft = procedure_draft_job.run(_request(process="cycle_count"))
    paths = procedure_draft_job.write_operational(draft, tmp_path / "out", client="Acme")

    assert set(paths) == {"md", "xlsx", "steps_csv", "control_plan_csv", "training_csv"}
    for path in paths.values():
        assert path.exists()

    assert paths["md"].name == "procedure_draft_cycle_count.md"
    assert paths["xlsx"].name == "procedure_draft_cycle_count.xlsx"
    assert paths["steps_csv"].name == "procedure_draft_steps.csv"
    assert paths["control_plan_csv"].name == "procedure_draft_control_plan.csv"
    assert paths["training_csv"].name == "procedure_draft_training.csv"

    steps_lines = paths["steps_csv"].read_text(encoding="utf-8").strip().splitlines()
    assert len(steps_lines) - 1 == len(draft.steps)

    control_plan_lines = paths["control_plan_csv"].read_text(encoding="utf-8").strip().splitlines()
    assert len(control_plan_lines) - 1 == len(draft.control_plan)

    training_lines = paths["training_csv"].read_text(encoding="utf-8").strip().splitlines()
    assert len(training_lines) - 1 == len(draft.training)


def test_xlsx_has_header_steps_controlplan_training_sheets(tmp_path):
    from openpyxl import load_workbook

    draft = procedure_draft_job.run(_request(process="replenishment"))
    paths = procedure_draft_job.write_operational(draft, tmp_path / "out")

    wb = load_workbook(paths["xlsx"])
    assert wb.sheetnames == ["Header", "Steps", "ControlPlan", "Training"]
    assert wb["Steps"].max_row - 1 == len(draft.steps)
    assert wb["ControlPlan"].max_row - 1 == len(draft.control_plan)
    assert wb["Training"].max_row - 1 == len(draft.training)


def test_markdown_header_carries_iso_7_5_2c_sentence_in_requested_lang(tmp_path):
    draft_es = procedure_draft_job.run(_request(lang="es"))
    paths_es = procedure_draft_job.write_operational(draft_es, tmp_path / "es")
    assert ISO_SENTENCE_ES in paths_es["md"].read_text(encoding="utf-8")

    draft_en = procedure_draft_job.run(_request(lang="en"))
    paths_en = procedure_draft_job.write_operational(draft_en, tmp_path / "en")
    assert ISO_SENTENCE_EN in paths_en["md"].read_text(encoding="utf-8")


def test_two_runs_same_params_are_byte_identical(tmp_path):
    params = {
        "process": "replenishment",
        "owner": "Ana",
        "approver": "Beto",
        "as_of": AS_OF_ISO,
        "service_level": "98%",
        "lead_time_days": 10,
    }
    draft1 = procedure_draft_job.run(procedure_draft_job.prepare(dict(params)))
    draft2 = procedure_draft_job.run(procedure_draft_job.prepare(dict(params)))

    paths1 = procedure_draft_job.write_operational(draft1, tmp_path / "run1")
    paths2 = procedure_draft_job.write_operational(draft2, tmp_path / "run2")

    assert paths1["md"].read_bytes() == paths2["md"].read_bytes()
    assert paths1["steps_csv"].read_bytes() == paths2["steps_csv"].read_bytes()
    assert paths1["control_plan_csv"].read_bytes() == paths2["control_plan_csv"].read_bytes()
    assert paths1["training_csv"].read_bytes() == paths2["training_csv"].read_bytes()


# ---- build_deck() ---------------------------------------------------------------------


def test_build_deck_reports_control_kpis_and_step_findings():
    draft = procedure_draft_job.run(_request())
    deck = procedure_draft_job.build_deck(draft, client="Acme", citations=("iso_9000",))

    assert deck.client == "Acme"
    assert deck.options == tuple(draft.outcome.options)
    assert "iso_9000" in deck.citations
    assert len(deck.findings) == len(draft.steps)

    kpi_names = {k.name for k in deck.kpis}
    assert {"Dueno", "Aprobador", "Version", "Proxima revision", "Cadencia de revision"} <= kpi_names
    assert deck.residual.strip()


def test_build_deck_writes_md_and_xlsx(tmp_path):
    draft = procedure_draft_job.run(_request())
    deck = procedure_draft_job.build_deck(draft, client="Acme")
    written = deck.write_all(tmp_path / "deck_out")
    assert written["report"].exists()
    assert written["workbook"].exists()
    assert "Acme" in written["report"].read_text(encoding="utf-8")
