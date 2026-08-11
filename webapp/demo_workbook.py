"""El entregable de /demo en Excel, con graficos nativos.

El embudo escribia solo markdown. El visitante al que va dirigido vive en Excel:
sube una planilla y recibe un `.md`, que es un formato de programador. Este modulo
cierra ese hueco para la salida #1 de 05-GTM-KERN §5 -- la que "se gana el derecho
a seguir hablando" y la unica que ese documento manda entregar si solo se pudiera
entregar una.

Los graficos son **nativos** (``openpyxl.chart``), no imagenes: el comprador hace
clic, ve el dato detras, lo copia a su propia presentacion. Una PNG no se audita.

Cinco reglas que gobiernan todo lo de abajo, y que no son estilo:

1. **El grafico entra en la primera pantalla.** Es la razon por la que el archivo
   existe. Si hay que scrollear a la derecha para verlo, el archivo fallo.
2. **Nada se pinta de un color que afirme algo que no se midio.** El formato
   condicional va con ``ISNUMBER`` adelante: en Excel el texto ordena DESPUES que
   cualquier numero, asi que ``"no medido" >= 0,1`` es TRUE y una regla ingenua
   pinta de rojo justamente las celdas sin dato.
3. **Los graficos leen bloques propios**, nunca la tabla humana. La tabla lleva
   "no medido", topes y autofiltro; un grafico apuntado ahi se vacia solo.
4. **``defuse_formula`` en cada celda, por un unico punto** (:func:`_cell`). Los
   nombres de proveedor vienen de un archivo subido por un anonimo.
5. **Lo no medido lo dice.** Nunca un blanco que se lea como cero.

Se construye SOLO bajo ``scan.ok``: con QA fallado ``findings`` viene vacio y el
titular no significa nada, asi que un libro seria un objeto con forma de resultado
sin resultado adentro.
"""


from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import Reference
from openpyxl.styles import Font  # noqa: E402  (fuentes puntuales de una sola celda)

from src.supplier_scorecard import LEAD_TIME_GAP_MATERIAL, SupplierScore
from webapp.demo_lead_time_scan import (
    CTA_PATH,
    MIN_DELIVERIES,
    LeadTimeScan,
    relative_gap,
    signed_relative_gap,
)

# Las primitivas viven en un solo lugar: ver el docstring de demo_workbook_base.
from webapp.demo_workbook_base import (
    _BASE_URL,
    _BIG,
    _FILL_FINA,
    _FILL_TARDE,
    _FILL_TEMPRANO,
    _MID,
    _NOTE,
    _SECTION,
    _TITLE,
    _WRAP,
    C_ALERTA,
    C_DECLARADO,
    C_NEUTRO,
    C_OBSERVADO,
    FMT_COUNT,
    FMT_DATE,
    FMT_DAYS_1,
    FMT_PCT_SIGNED,
    FMT_UNITS_1,
    NO_MEDIDO,
    TOP_N_PROVEEDOR,
    _bar,
    _cap,
    _cell,
    _chart_block,
    _headers,
    _isnumber_rule,
    _link_row,
    _measured,
    _widths,
    rows_below,
)
from webapp.operator_profile import get_operator_profile

#: Bajo esta cantidad de recepciones el promedio es indicativo y la variacion es
#: ruido. No descalifica al proveedor (para eso esta MIN_DELIVERIES); avisa.
EVIDENCIA_FINA = 8

#: Ancho del bin del histograma, en dias.
BIN_DIAS = 5


# --------------------------------------------------------------------------- #
# hojas
# --------------------------------------------------------------------------- #


def _comparables(scan: LeadTimeScan) -> list[SupplierScore]:
    """Proveedores con promesa y medicion a la vez, del mas desviado al menos."""
    return sorted(
        (s for s in scan.scores if relative_gap(s) is not None),
        key=lambda s: relative_gap(s) or 0.0,
        reverse=True,
    )


def _cat_label(score: SupplierScore) -> str:
    """El nombre del proveedor con su n pegado.

    El n viaja dentro de la categoria a proposito: sobrevive al grafico, a la
    captura de pantalla y al copiar-pegar. Un titular sostenido por 3 recepciones
    y uno sostenido por 180 no pueden verse igual.
    """
    return f"{score.supplier} (n={score.evaluable.lead_time})"


def _sheet_resumen(ws, scan: LeadTimeScan, *, dataset_label: str, ts: str) -> None:
    h = scan.headline
    w = scan.worst
    _widths(ws, {"A": 38, "B": 22, "C": 16, "D": 3})

    _cell(ws, 1, 1, "El plazo que promete tu sistema y el que muestran tus recepciones", font=_TITLE)
    _cell(ws, 2, 1, "Archivo analizado:")
    _cell(ws, 2, 2, dataset_label)
    _cell(ws, 3, 1, "Fecha del analisis:")
    _cell(ws, 3, 2, ts)
    _cell(ws, 4, 1, "Proveedores medidos:")
    _cell(ws, 4, 2, h["n_suppliers"], fmt=FMT_COUNT)
    _cell(ws, 5, 1, "Recepciones consideradas:")
    _cell(ws, 5, 2, h["n_deliveries"], fmt=FMT_COUNT)

    _cell(ws, 7, 1, "El numero", font=_SECTION)
    _cell(ws, 8, 1, "Proveedor con la mayor diferencia")
    _cell(ws, 8, 2, w.supplier, font=_MID)
    _cell(ws, 9, 1, "Tu sistema dice que entrega en")
    _cell(ws, 9, 2, _measured(w.declared_lead_time), fmt=FMT_DAYS_1,
          font=Font(bold=True, size=18, color=C_DECLARADO))
    _cell(ws, 10, 1, f"Sus ultimas {w.evaluable.lead_time} recepciones promedian")
    _cell(ws, 10, 2, _measured(w.avg_lead_time), fmt=FMT_DAYS_1,
          font=Font(bold=True, size=18, color=C_OBSERVADO))
    _cell(ws, 11, 1, "Diferencia")
    _cell(ws, 11, 2, _measured(w.lead_time_gap), fmt=FMT_DAYS_1,
          font=Font(bold=True, size=16, color=C_ALERTA))
    _cell(ws, 12, 1, "1 de cada 10 recepciones tarda")
    _cell(ws, 12, 2, _measured(w.lead_time_p90), fmt=FMT_DAYS_1)
    _cell(ws, 13, 1, "SKU cuyo stock de seguridad esta armado sobre el numero equivocado")
    _cell(ws, 13, 2, h["n_skus_to_recalculate"], fmt=FMT_COUNT, font=_BIG)

    # El bloque del grafico vive lejos a la derecha (col N) para que el grafico
    # pueda anclarse DENTRO de la primera pantalla, debajo del titular.
    comparables = _comparables(scan)[:TOP_N_PROVEEDOR]
    header_row, last_row = _chart_block(
        ws,
        1,
        "Datos del grafico de esta hoja: proveedores con promesa y medicion, por diferencia",
        ["Proveedor", "Lo que dice tu sistema (d)", "Lo que dicen tus recepciones (d)"],
        [[_cat_label(s), round(s.declared_lead_time, 1), round(s.avg_lead_time, 1)]
         for s in comparables],
        start_col=14,
        fmt=FMT_UNITS_1,
    )
    ch = _bar(
        data=Reference(ws, min_col=15, max_col=16, min_row=header_row, max_row=last_row),
        cats=Reference(ws, min_col=14, min_row=header_row + 1, max_row=last_row),
        title="Lo prometido y lo medido, por proveedor",
        y_title="Dias",
        colors=(C_DECLARADO, C_OBSERVADO),
        width=19,
        height=9.5,
        gap=60,
    )
    ws.add_chart(ch, "A15")

    # Se sigue debajo del ALTO RENDERIZADO del grafico, nunca debajo de su bloque
    # de datos: el ancla fija la esquina y el alto son EMU absolutos.
    row = 15 + rows_below(9.5)
    _cell(ws, row, 1, "Que dice esto", font=_SECTION)
    row += 1
    for finding in scan.findings:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        _cell(ws, row, 1, finding, align=_WRAP)
        ws.row_dimensions[row].height = 42
        row += 1

    row += 1
    operator = get_operator_profile()
    _cell(ws, row, 1, "Siguiente paso", font=_SECTION)
    row += 1
    _link_row(ws, row, 1, "Diagnostico de Arranque (abrir)", f"{_BASE_URL}{CTA_PATH}")
    row += 1
    if operator.email:
        _cell(ws, row, 1, "Responde a:")
        _cell(ws, row, 2, f"{operator.name} <{operator.email}>")
        row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    _cell(
        ws,
        row,
        1,
        "Este archivo lo genero un escaneo automatico sobre el que subiste, en solo "
        "lectura: no se escribio nada en tus sistemas. El plazo observado sale de "
        "restar fecha de recepcion menos fecha de la orden, de tu propio archivo. El "
        "declarado es el campo de plazo de tu sistema: una promesa, no una medicion. "
        "'no medido' significa que el dato no venia, nunca que el valor sea cero.",
        font=_NOTE,
        align=_WRAP,
    )
    ws.row_dimensions[row].height = 60


def _sheet_proveedores(ws, scan: LeadTimeScan) -> None:
    ordered = _comparables(scan) + [s for s in scan.scores if relative_gap(s) is None]
    _headers(ws, 1, [
        "Proveedor", "Recepciones en el archivo", "Recepciones con fecha usable",
        "Lo que dice tu sistema (d)", "Lo que dicen tus recepciones (d)",
        "Variacion entre recepciones (d)", "1 de cada 10 tarda (d)",
        "Diferencia (d)", "Diferencia (%)", "SKU a recalcular",
    ])
    for r, s in enumerate(ordered, 2):
        _cell(ws, r, 1, s.supplier)
        _cell(ws, r, 2, s.deliveries, fmt=FMT_COUNT)
        _cell(ws, r, 3, s.evaluable.lead_time, fmt=FMT_COUNT)
        _cell(ws, r, 4, _measured(s.declared_lead_time), fmt=FMT_DAYS_1)
        _cell(ws, r, 5, _measured(s.avg_lead_time), fmt=FMT_DAYS_1)
        _cell(ws, r, 6, _measured(s.lead_time_std), fmt=FMT_DAYS_1)
        _cell(ws, r, 7, _measured(s.lead_time_p90), fmt=FMT_DAYS_1)
        _cell(ws, r, 8, _measured(s.lead_time_gap), fmt=FMT_DAYS_1)
        _cell(ws, r, 9, _measured(signed_relative_gap(s), digits=4), fmt=FMT_PCT_SIGNED)
        _cell(ws, r, 10, len(s.skus_to_recalculate), fmt=FMT_COUNT)
    last = 1 + len(ordered)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{last}"

    if last >= 2:
        for ref, rule in (
            _isnumber_rule("I", 2, last, f">={LEAD_TIME_GAP_MATERIAL}", _FILL_TARDE),
            _isnumber_rule("I", 2, last, f"<=-{LEAD_TIME_GAP_MATERIAL}", _FILL_TEMPRANO),
            _isnumber_rule("C", 2, last, f"<{EVIDENCIA_FINA}", _FILL_FINA),
        ):
            ws.conditional_formatting.add(ref, rule)

    pct = f"{LEAD_TIME_GAP_MATERIAL * 100:.0f}%"
    for i, text in enumerate((
        f"Fondo rojo: el proveedor tarda al menos {pct} mas de lo que dice la ficha. "
        "Por encima de ese umbral, el stock de seguridad armado sobre el plazo declarado deja de servir.",
        f"Fondo naranja: la ficha promete al menos {pct} mas dias de los que el proveedor "
        "realmente usa. Ahi hay buffer parado.",
        f"Fondo gris en 'Recepciones con fecha usable': menos de {EVIDENCIA_FINA} recepciones. "
        "El promedio es indicativo y la variacion todavia es ruido.",
        "'no medido': el dato no venia en el archivo. Nunca significa cero.",
    )):
        _cell(ws, last + 2 + i, 1, text, font=_NOTE)

    con_cola = [s for s in scan.scores if s.avg_lead_time is not None and s.lead_time_p90 is not None]
    con_cola.sort(key=lambda s: s.lead_time_p90 or 0.0, reverse=True)
    con_cola = con_cola[:TOP_N_PROVEEDOR]
    if con_cola:
        header_row, last_row = _chart_block(
            ws,
            1,
            "Datos del grafico: promedio y cola",
            ["Proveedor", "Promedio (d)", "1 de cada 10 tarda (d)"],
            [[_cat_label(s), round(s.avg_lead_time, 1), round(s.lead_time_p90, 1)] for s in con_cola],
            start_col=13,
            fmt=FMT_UNITS_1,
        )
        ch = _bar(
            data=Reference(ws, min_col=14, max_col=15, min_row=header_row, max_row=last_row),
            cats=Reference(ws, min_col=13, min_row=header_row + 1, max_row=last_row),
            title="El promedio no es el peor caso",
            y_title="Dias",
            colors=(C_OBSERVADO, C_ALERTA),
            width=19,
            height=10,
            gap=60,
        )
        ws.add_chart(ch, "Q2")

    _widths(ws, {"A": 28, "B": 22, "C": 24, "D": 22, "E": 24, "F": 24, "G": 20,
                 "H": 15, "I": 14, "J": 16})


def _sheet_recepciones(ws, scan: LeadTimeScan) -> None:
    """Fila por recepcion: lo unico que hace verificable el titular en 5 minutos."""
    _cell(ws, 1, 1, "Cada recepcion que entro en los numeros de este archivo", font=_SECTION)
    _cell(ws, 2, 1,
          "El plazo de cada fila es la resta de sus dos fechas. Toma una al azar, buscala "
          "en tu sistema y comprobala.", font=_NOTE, align=_WRAP)

    rows = sorted(
        scan.measured,
        key=lambda d: (str(d.get("supplier", "")), str(d.get("receipt_date", ""))),
    )
    total = len(rows)
    rows = _cap(rows, ws, 3, 1, total, "proveedor y fecha de recepcion")

    _headers(ws, 4, ["Proveedor", "Orden de compra", "SKU", "Fecha de la orden",
                     "Fecha de recepcion", "Dias"])
    for r, row in enumerate(rows, 5):
        _cell(ws, r, 1, row.get("supplier"))
        _cell(ws, r, 2, row.get("po_id") or NO_MEDIDO)
        _cell(ws, r, 3, row.get("product_id") or NO_MEDIDO)
        _cell(ws, r, 4, _as_date(row.get("order_date")), fmt=FMT_DATE)
        _cell(ws, r, 5, _as_date(row.get("receipt_date")), fmt=FMT_DATE)
        _cell(ws, r, 6, _measured(row.get("lead_time_days"), digits=0), fmt=FMT_COUNT)
    ws.freeze_panes = "A5"
    if rows:
        ws.auto_filter.ref = f"A4:F{4 + len(rows)}"
    _widths(ws, {"A": 28, "B": 20, "C": 22, "D": 18, "E": 18, "F": 10})

    _histograma(ws, scan)


def _histograma(ws, scan: LeadTimeScan) -> None:
    """La distribucion del peor proveedor: el argumento del producto, dibujado.

    El buffer se dimensiona contra la dispersion, no contra el promedio. Dos barras
    de promedio no lo muestran; la forma de la distribucion si -- y lo dice sin la
    palabra sigma, que 05-GTM-KERN prohibe en la primera conversacion.
    """
    w = scan.worst
    dias = [
        float(d["lead_time_days"])
        for d in scan.measured
        if d.get("supplier") == w.supplier and d.get("lead_time_days") is not None
    ]
    if not dias:
        return
    lo = int(min(dias) // BIN_DIAS) * BIN_DIAS
    hi = int(max(dias) // BIN_DIAS) * BIN_DIAS
    bins: list[tuple[str, int]] = []
    edge = lo
    while edge <= hi:
        n = sum(1 for d in dias if edge <= d < edge + BIN_DIAS)
        bins.append((f"{edge}-{edge + BIN_DIAS - 1} d", n))
        edge += BIN_DIAS

    header_row, last_row = _chart_block(
        ws,
        1,
        f"Datos del histograma: {w.supplier}",
        ["Tramo", "Recepciones"],
        [[label, n] for label, n in bins],
        start_col=8,
        fmt=FMT_COUNT,
    )
    declarado = f"{w.declared_lead_time:.0f}" if w.declared_lead_time else NO_MEDIDO
    ch = _bar(
        data=Reference(ws, min_col=9, min_row=header_row, max_row=last_row),
        cats=Reference(ws, min_col=8, min_row=header_row + 1, max_row=last_row),
        title=f"Como se reparten las recepciones de {w.supplier} (promete {declarado} d)",
        y_title="Recepciones",
        colors=(C_OBSERVADO,),
        width=19,
        height=9,
        legend=None,
        gap=0,
    )
    # A la derecha del bloque (H-I), no encima: con muchos tramos el ancla "H8"
    # caeria dentro de los propios datos del histograma.
    ws.add_chart(ch, "K1")
    _cell(ws, last_row + 2, 8,
          "El stock de seguridad se dimensiona contra este ancho, no contra el promedio. "
          "El campo de plazo de tu sistema es un solo numero y no puede describirlo.",
          font=_NOTE, align=_WRAP)


def _sheet_skus(ws, scan: LeadTimeScan) -> None:
    h = scan.headline
    _cell(ws, 1, 1,
          "SKU cuyo stock de seguridad esta armado sobre un plazo que las recepciones contradicen",
          font=_SECTION)
    _cell(ws, 2, 1,
          f"Total de SKU distintos: {h['n_skus_to_recalculate']}. Un SKU que se le compra a dos "
          "proveedores aparece dos veces en esta tabla y una sola vez en ese total.",
          font=_NOTE, align=_WRAP)

    pares = [(s, sku) for s in _comparables(scan) for sku in s.skus_to_recalculate]
    total = len(pares)
    pares = _cap(pares, ws, 3, 1, total, "diferencia del proveedor")

    _headers(ws, 4, ["SKU", "Proveedor", "Lo que dice tu sistema (d)",
                     "Lo que dicen tus recepciones (d)", "Diferencia (%)"])
    for r, (s, sku) in enumerate(pares, 5):
        _cell(ws, r, 1, sku)
        _cell(ws, r, 2, s.supplier)
        _cell(ws, r, 3, _measured(s.declared_lead_time), fmt=FMT_DAYS_1)
        _cell(ws, r, 4, _measured(s.avg_lead_time), fmt=FMT_DAYS_1)
        _cell(ws, r, 5, _measured(signed_relative_gap(s), digits=4), fmt=FMT_PCT_SIGNED)
    ws.freeze_panes = "A5"
    _widths(ws, {"A": 24, "B": 28, "C": 24, "D": 26, "E": 16})

    por_proveedor = [(s.supplier, len(s.skus_to_recalculate))
                     for s in _comparables(scan) if s.skus_to_recalculate]
    por_proveedor.sort(key=lambda t: t[1], reverse=True)
    por_proveedor = por_proveedor[:TOP_N_PROVEEDOR]
    if por_proveedor:
        header_row, last_row = _chart_block(
            ws, 1, "Datos del grafico", ["Proveedor", "SKU a recalcular"],
            [[name, n] for name, n in por_proveedor], start_col=7, fmt=FMT_COUNT,
        )
        ch = _bar(
            data=Reference(ws, min_col=8, min_row=header_row, max_row=last_row),
            cats=Reference(ws, min_col=7, min_row=header_row + 1, max_row=last_row),
            title="SKU a recalcular, por proveedor",
            kind="bar",
            colors=(C_ALERTA,),
            width=17,
            height=min(3 + 0.6 * len(por_proveedor), 12),
            legend=None,
        )
        ws.add_chart(ch, "K2")


def _sheet_descartes(ws, scan: LeadTimeScan) -> None:
    """Siempre se escribe, aunque todo sea cero: un cero aca es una medicion."""
    d = scan.dropped
    prov_cortos, filas_cortas = scan.below_min
    _cell(ws, 1, 1, "Que filas del archivo no entraron en ningun numero", font=_SECTION)
    _cell(ws, 2, 1, "Se reporta aunque sea cero.", font=_NOTE)
    _headers(ws, 4, ["Motivo", "Filas", "Que significa"])
    filas = [
        ("Sin fecha usable", d.get("sin_fecha", 0),
         "La fecha de la orden o la de la recepcion no se pudo leer. Sin las dos no hay plazo que medir."),
        ("Recepcion anterior a su orden", d.get("plazo_negativo", 0),
         "Un plazo negativo es un dato roto, no un plazo de cero dias: promediarlo bajaria el "
         "plazo observado y taparia justo la diferencia que veniamos a mostrar."),
        ("Sin proveedor", d.get("sin_proveedor", 0),
         "La fila no dice a quien se le compro."),
        (f"Proveedores con menos de {MIN_DELIVERIES} recepciones", filas_cortas,
         f"Se leyeron bien, pero {prov_cortos} proveedor(es) no junta(n) {MIN_DELIVERIES} "
         "recepciones. Con menos, la variacion es ruido y no describe al proveedor."),
    ]
    for r, (motivo, n, texto) in enumerate(filas, 5):
        _cell(ws, r, 1, motivo)
        _cell(ws, r, 2, n, fmt=FMT_COUNT)
        _cell(ws, r, 3, texto, align=_WRAP)
    total_fuera = sum(d.values()) + filas_cortas
    _cell(ws, 9, 1, "Total fuera de los numeros", font=Font(bold=True))
    _cell(ws, 9, 2, total_fuera, fmt=FMT_COUNT, font=Font(bold=True))
    _cell(ws, 10, 1, "Recepciones usadas", font=Font(bold=True))
    _cell(ws, 10, 2, scan.headline["n_deliveries"], fmt=FMT_COUNT, font=Font(bold=True))
    _cell(ws, 11, 1, "Filas del archivo", font=Font(bold=True))
    _cell(ws, 11, 2, total_fuera + scan.headline["n_deliveries"], fmt=FMT_COUNT,
          font=Font(bold=True))
    _widths(ws, {"A": 40, "B": 12, "C": 76})

    if total_fuera:
        ch = _bar(
            data=Reference(ws, min_col=2, min_row=4, max_row=8),
            cats=Reference(ws, min_col=1, min_row=5, max_row=8),
            title="Filas fuera de los numeros, por motivo",
            kind="bar",
            colors=(C_NEUTRO,),
            width=16,
            height=7,
            legend=None,
        )
        ws.add_chart(ch, "E4")


def _sheet_metodo(ws, scan: LeadTimeScan) -> None:
    _widths(ws, {"A": 44, "B": 92})
    bloques = [
        ("Como se midio", [
            ("Plazo observado",
             "Fecha de recepcion menos fecha de la orden, de tu propio archivo. De ningun otro lugar."),
            ("Plazo declarado",
             "El campo de plazo que trae el export: lo que tu sistema o la ficha del proveedor "
             "promete. Es una promesa, no una medicion -- por eso no se usa como observacion."),
            ("Minimo para calificar",
             f"{MIN_DELIVERIES} recepciones con fechas validas. Con menos, la variacion es ruido."),
            ("Umbral de diferencia",
             f"{LEAD_TIME_GAP_MATERIAL * 100:.0f}%. Por debajo de eso, rehacer el stock de "
             "seguridad no cambia la decision."),
            ("Unidades", "Todos los plazos de este archivo estan en DIAS."),
        ]),
        ("Que NO se midio", [
            ("Cumplimiento de la fecha comprometida",
             "El archivo de recepciones no trae la fecha comprometida, asi que no se reporta. "
             "Un porcentaje inventado aqui seria peor que no tenerlo."),
            ("Entregas completas", "No venian las cantidades pedidas contra las recibidas."),
            ("Defectos", "No venian unidades ni rechazos."),
            ("Diferencias de precio contra la orden", "No venia el precio facturado."),
            ("Proveedores con plazo declarado en cero o negativo",
             "El campo del ERP trae un valor imposible. La diferencia en % sale 'no medido' "
             "en vez de un porcentaje calculado sobre un divisor roto."),
        ]),
        ("Convenciones", [
            ("'no medido'", "El dato no venia en el archivo. Nunca significa cero."),
            ("Signo de la diferencia",
             "Positivo: el proveedor tarda mas de lo que promete. Negativo: la ficha promete "
             "mas dias de los que realmente usa, y ahi hay buffer parado."),
            ("Solo lectura", "Ningun numero de este archivo se escribio en tus sistemas."),
        ]),
    ]
    row = 1
    for titulo, pares in bloques:
        _cell(ws, row, 1, titulo, font=_SECTION)
        row += 1
        for etiqueta, texto in pares:
            _cell(ws, row, 1, etiqueta)
            _cell(ws, row, 2, texto, align=_WRAP)
            ws.row_dimensions[row].height = 30
            row += 1
        row += 1


def _as_date(value: Any) -> Any:
    """Timestamp -> date para que Excel lo trate como fecha y no como texto."""
    try:
        if value is None or (hasattr(value, "__class__") and str(value) in ("NaT", "nan")):
            return NO_MEDIDO
        return value.date() if hasattr(value, "date") else value
    except (AttributeError, ValueError):
        return NO_MEDIDO


# --------------------------------------------------------------------------- #
# entrada publica
# --------------------------------------------------------------------------- #


def write_lead_time_workbook(
    scan: LeadTimeScan, path: str | Path, *, dataset_label: str, ts: str
) -> Path:
    """Escribe ``plazo_real.xlsx``. Solo puede llamarse con ``scan.ok``.

    ``ts`` entra por parametro y nunca se lee del reloj: dos corridas sobre el
    mismo scan tienen que producir el mismo archivo.
    """
    if not scan.ok:
        raise ValueError("el libro no se construye con QA fallado: no hay titular que escribir")
    if scan.worst is None:  # pragma: no cover - ok ya lo garantiza
        raise ValueError("sin proveedor comparable no hay entregable")

    wb = Workbook()
    _sheet_resumen(wb.active, scan, dataset_label=dataset_label, ts=ts)
    wb.active.title = "Resumen"
    _sheet_proveedores(wb.create_sheet("Plazo por proveedor"), scan)
    if scan.measured:
        _sheet_recepciones(wb.create_sheet("Recepciones medidas"), scan)
    if any(s.skus_to_recalculate for s in scan.scores):
        _sheet_skus(wb.create_sheet("SKU a recalcular"), scan)
    _sheet_descartes(wb.create_sheet("Recepciones descartadas"), scan)
    _sheet_metodo(wb.create_sheet("Como leer esto"), scan)

    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
