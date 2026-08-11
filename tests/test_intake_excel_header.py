"""`load_table` encuentra el encabezado y la hoja en una planilla de cliente real.

``pd.read_excel(p)`` a secas asume PRIMERA HOJA y FILA 0. Medido contra planillas
armadas como las arma un ERP, esos dos supuestos eran los dos modos de falla mas
comunes y los dos daban CERO columnas mapeadas:

* una fila de titulo arriba ("Reporte de Recepciones 2026") desplaza el encabezado;
* una portada con el logo deja los datos en la segunda o tercera hoja.

Los dos salian por el 400 de "columnas requeridas", que es honesto pero es
friccion en el unico paso que 05-GTM-KERN exige que sea de cero contacto.

Lo que estos tests fijan, en este orden de importancia:

1. los dos casos rotos ahora entran;
2. los que YA funcionaban siguen igual (no se cambio un modo de falla por otro);
3. y el archivo sin nada reconocible SIGUE siendo un 400 explicito -- la deteccion
   nunca inventa un encabezado para tener algo que devolver.
"""
from __future__ import annotations

import pytest

pytest.importorskip("openpyxl")

from openpyxl import Workbook  # noqa: E402

from jobs.intake import (  # noqa: E402
    RECEIPT_OPTIONAL,
    RECEIPT_REQUIRED,
    detect_columns,
    detect_sheet_and_header,
    load_table,
)

HEADER = ["supplier", "order_date", "receipt_date"]
ROWS = [
    ["ACME", "2026-01-01", "2026-02-04"],
    ["ACME", "2026-02-01", "2026-03-10"],
    ["GLOBEX", "2026-01-05", "2026-01-20"],
]


def _book(tmp_path, build) -> str:
    wb = Workbook()
    build(wb)
    path = tmp_path / "cliente.xlsx"
    wb.save(path)
    return str(path)


def _mapped(path: str) -> list[str]:
    df = load_table(path)
    return sorted(
        detect_columns(df, required=RECEIPT_REQUIRED, optional=RECEIPT_OPTIONAL).mapping
    )


# --- los dos casos que estaban rotos ---------------------------------------


def test_fila_de_titulo_arriba_ya_no_tapa_el_encabezado(tmp_path):
    def build(wb):
        ws = wb.active
        ws.append(["Reporte de Recepciones 2026"])
        ws.append([])
        ws.append(HEADER)
        for r in ROWS:
            ws.append(r)

    path = _book(tmp_path, build)
    assert detect_sheet_and_header(path)[1] == 2
    assert _mapped(path) == ["order_date", "receipt_date", "supplier"]
    assert len(load_table(path)) == len(ROWS)  # el titulo no entra como dato


def test_datos_en_la_segunda_hoja_detras_de_una_portada(tmp_path):
    def build(wb):
        portada = wb.active
        portada.title = "Portada"
        portada.append(["Generado por el ERP"])
        datos = wb.create_sheet("Datos")
        datos.append(HEADER)
        for r in ROWS:
            datos.append(r)

    path = _book(tmp_path, build)
    assert detect_sheet_and_header(path) == ("Datos", 0)
    assert _mapped(path) == ["order_date", "receipt_date", "supplier"]


def test_titulo_y_tercera_hoja_a_la_vez(tmp_path):
    """Los dos defectos juntos, que es como llegan de verdad."""

    def build(wb):
        wb.active.title = "Logo"
        wb.active.append(["ACME S.A."])
        wb.create_sheet("Parametros").append(["periodo", "2026"])
        detalle = wb.create_sheet("Detalle")
        detalle.append(["Informe de recepciones"])
        detalle.append([])
        detalle.append(HEADER)
        for r in ROWS:
            detalle.append(r)

    path = _book(tmp_path, build)
    assert detect_sheet_and_header(path) == ("Detalle", 2)
    assert _mapped(path) == ["order_date", "receipt_date", "supplier"]


def test_el_encabezado_desplazado_tambien_se_encuentra_en_espanol_con_tildes(tmp_path):
    """La deteccion puntua con `_norm`, asi que hereda el plegado NFKD."""

    def build(wb):
        ws = wb.active
        ws.append(["Informe mensual"])
        ws.append(["Proveedor", "Fecha Orden", "Fecha Recepción"])
        for r in ROWS:
            ws.append(r)

    path = _book(tmp_path, build)
    assert detect_sheet_and_header(path)[1] == 1
    assert _mapped(path) == ["order_date", "receipt_date", "supplier"]


# --- lo que ya funcionaba y no se puede romper ------------------------------


def test_planilla_limpia_sigue_siendo_hoja_cero_fila_cero(tmp_path):
    def build(wb):
        ws = wb.active
        ws.append(HEADER)
        for r in ROWS:
            ws.append(r)

    path = _book(tmp_path, build)
    assert detect_sheet_and_header(path)[1] == 0
    assert _mapped(path) == ["order_date", "receipt_date", "supplier"]


def test_columna_vacia_a_la_izquierda(tmp_path):
    def build(wb):
        ws = wb.active
        ws.append([""] + HEADER)
        for r in ROWS:
            ws.append([""] + r)

    path = _book(tmp_path, build)
    assert _mapped(path) == ["order_date", "receipt_date", "supplier"]


# --- la propiedad que no se puede perder ------------------------------------


def test_sin_ningun_alias_reconocible_no_se_inventa_un_encabezado(tmp_path):
    """Falla cerrado: (0, 0) es el comportamiento previo exacto, y el 400 sobrevive.

    Es la mitad importante del arreglo. Una deteccion que siempre elige "la mejor"
    fila convertiria un rechazo explicito en un mapeo plausible y equivocado, que
    es el modo de falla que este repo ya pago una vez con `_norm` y los acentos.
    """

    def build(wb):
        ws = wb.active
        ws.append(["aaa", "bbb", "ccc"])
        ws.append([1, 2, 3])

    path = _book(tmp_path, build)
    assert detect_sheet_and_header(path) == (0, 0)
    assert _mapped(path) == []


def test_empate_resuelve_a_la_hoja_mas_a_la_izquierda_y_la_fila_mas_arriba(tmp_path):
    """Determinismo: dos corridas sobre el mismo archivo tienen que dar lo mismo."""

    def build(wb):
        primera = wb.active
        primera.title = "Enero"
        primera.append(HEADER)
        for r in ROWS:
            primera.append(r)
        segunda = wb.create_sheet("Febrero")
        segunda.append(HEADER)  # mismo puntaje, exactamente
        for r in ROWS:
            segunda.append(r)

    path = _book(tmp_path, build)
    assert detect_sheet_and_header(path) == ("Enero", 0)
    assert detect_sheet_and_header(path) == ("Enero", 0)


def test_el_csv_no_cambia_de_comportamiento(tmp_path):
    """El arreglo es de Excel. En CSV un titulo arriba rompe el parseo con un
    error de columnas desparejas, que es otra clase de falla y otro arreglo."""
    path = tmp_path / "recepciones.csv"
    path.write_text(
        "supplier,order_date,receipt_date\n" + "\n".join(",".join(r) for r in ROWS) + "\n",
        encoding="utf-8",
    )
    assert _mapped(str(path)) == ["order_date", "receipt_date", "supplier"]
