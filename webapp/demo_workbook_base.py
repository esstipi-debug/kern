"""Primitivas compartidas por los tres libros de Excel de /demo.

Se extrajeron de ``webapp/demo_workbook.py`` cuando entro el segundo libro. La
razon no es el largo del archivo (aunque tambien): es que la regla 4 de aquel
modulo -- *"defuse_formula en cada celda, por un unico punto"* -- solo es
verificable si ese punto es UNO. Tres copias de ``_cell`` en tres archivos
convierten un invariante que se comprueba leyendo una funcion en una auditoria.

Lo mismo vale para ``_measured``: es la unica definicion del token de honestidad,
y por lo tanto el unico lugar donde se decide que un ``inf`` no es un numero.

Este modulo no importa nada de ``webapp.*`` -- solo ``src.*`` y openpyxl -- asi
que no hay ciclo posible con los libros que lo consumen.

Las cinco reglas que gobiernan los tres libros viven en el docstring de
``webapp/demo_workbook.py``, que es el primero que las escribio.
"""

from __future__ import annotations

import math
import os
from collections.abc import Collection, Sequence
from typing import Any

from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.deliverable import not_measured
from src.excel_export import HEADER_FILL, HEADER_FONT
from src.sanitize import defuse_formula

#: El token de honestidad, con una sola definicion para los tres libros.
NO_MEDIDO: str = not_measured("es")

#: Tope de series de un grafico. Las TABLAS nunca se topan por esto -- un grafico
#: con 200 barras no se lee, una tabla con 200 filas si.
TOP_N_PROVEEDOR = 15
TOP_N_SKU = 15

#: Tope duro de filas por hoja. fly.toml fija WEB_CONCURRENCY=1 sobre 512 MB (dos
#: workers ya murieron por OOM), y el tope de subida son 25 MB. Se topa, y se DICE
#: que se topo: callarlo seria romper "nada se descarta en silencio" justo en la
#: hoja que existe para no descartar nada.
MAX_TABLE_ROWS = 5_000

_BASE_URL = (os.environ.get("PUBLIC_BASE_URL") or "https://linchpin.fly.dev").rstrip("/")

#: Alto de una fila de Excel por defecto, en centimetros (15 pt). Es la constante
#: que traduce el alto de un grafico -- que openpyxl fija en EMU absolutos -- a
#: filas de la hoja. Ver :func:`rows_below`.
_ALTO_FILA_CM = 0.529

FMT_DAYS_1 = '#,##0.0" d"'
FMT_UNITS_1 = "#,##0.0"
FMT_COUNT = "#,##0"
FMT_PCT = "0%"
FMT_PCT_1 = "0.0%"
FMT_PCT_SIGNED = "+0%;-0%;0%"
FMT_DATE = "yyyy-mm-dd"

#: Dinero SIN simbolo, a proposito. Ningun camino de /demo sabe en que moneda esta
#: el archivo: el costo unitario es una columna cualquiera de un CSV subido. Poner
#: "$" seria afirmar una moneda que el archivo nunca trajo.
FMT_MONEY = "#,##0"
FMT_MONEY_2 = "#,##0.00"
#: Sin la seccion [Red] de costumbre: en el libro de la brecha un negativo es caja
#: que se libera, o sea la buena noticia. El color lo pone el formato condicional,
#: con leyenda al lado.
FMT_MONEY_SIGNED = "#,##0;-#,##0;0"

C_DECLARADO = "9AA5B1"  # gris: la promesa, o el supuesto del escaneo
C_OBSERVADO = "1F4E79"  # azul: la medicion
C_ALERTA = "C0392B"     # rojo: el hallazgo
C_NEUTRO = "9AA5B1"     # gris: contexto, un total
C_CAJA = "2E7D32"       # verde: caja que se libera

_TITLE = Font(bold=True, size=15)
_SECTION = Font(bold=True, size=12)
_BIG = Font(bold=True, size=20)
_MID = Font(bold=True, size=14)
_NOTE = Font(italic=True, size=9, color="5A6570")
_BLOCK = Font(bold=True, size=10, color="5A6570")
_LINK = Font(bold=True, color="1F4E79", underline="single")
_WRAP = Alignment(wrap_text=True, vertical="top")

_FILL_TARDE = PatternFill("solid", fgColor="F8D7DA")
_FILL_TEMPRANO = PatternFill("solid", fgColor="FDE2C8")
_FILL_FINA = PatternFill("solid", fgColor="EEF1F4")
_FILL_CAJA = PatternFill("solid", fgColor="D9EAD3")


def rows_below(height_cm: float, margin: int = 2) -> int:
    """Cuantas filas ocupa hacia abajo un grafico de ``height_cm`` de alto.

    El ancla de un grafico (``OneCellAnchor``) fija la esquina superior; el alto
    es EMU absolutos y **no tiene nada que ver con donde estan sus datos**. Poner
    el segundo grafico de una hoja a la altura del segundo bloque de datos lo
    superpone con el primero para cualquier cantidad de filas. Se calcula desde
    el alto renderizado, siempre.
    """
    return math.ceil(height_cm / _ALTO_FILA_CM) + margin


def _cell(
    ws,
    row: int,
    col: int,
    value: Any,
    *,
    fmt: str | None = None,
    font: Font | None = None,
    align: Alignment | None = None,
):
    """El unico camino por el que se escribe una celda en los tres libros.

    ``defuse_formula`` va aca y en ningun otro lado: un solo punto de paso es lo
    que hace que "se defusa todo" sea verificable leyendo una funcion en vez de
    auditando cada hoja.
    """
    c = ws.cell(row=row, column=col, value=defuse_formula(value))
    if fmt:
        c.number_format = fmt
    if font:
        c.font = font
    if align:
        c.alignment = align
    return c


def _headers(ws, row: int, labels: Sequence[str], *, start_col: int = 1) -> None:
    """Encabezados: literales del desarrollador, no dato del cliente."""
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=label)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT


def _measured(value: float | None, *, digits: int = 1) -> float | str:
    """El valor cuando es un numero finito; ``NO_MEDIDO`` cuando no lo es.

    Un ``None`` escrito como celda vacia se lee como cero, y en un plazo de
    entrega o en dias de cobertura, cero es una afirmacion fuerte y falsa. Un
    ``inf`` -- que el motor produce a proposito para un SKU sin demanda -- vale
    lo mismo: no es "cero dias de stock", es exactamente lo contrario.
    """
    if value is None:
        return NO_MEDIDO
    try:
        v = float(value)
    except (TypeError, ValueError):
        return NO_MEDIDO
    if not math.isfinite(v):
        return NO_MEDIDO
    return round(v, digits)


def _widths(ws, mapping: dict[str, int]) -> None:
    for letter, width in mapping.items():
        ws.column_dimensions[letter].width = width


def _bar(
    *,
    data: Reference,
    cats: Reference,
    title: str,
    kind: str = "col",
    y_title: str | None = None,
    colors: Sequence[str] = (),
    width: float = 20,
    height: float = 10,
    legend: str | None = "b",
    show_values: bool = True,
    gap: int | None = None,
) -> BarChart:
    """El unico constructor de barras. Concentra la higiene que se olvida de a una.

    ``y_axis.scaling.min = 0`` es la que hay que tener presente al armar un bloque:
    una barra negativa no se dibuja, desaparece en silencio. Por eso los bloques
    escriben **magnitudes positivas en dos series**, nunca una serie con signo.
    """
    ch = BarChart()
    ch.type = kind
    ch.grouping = "clustered"
    ch.title = title
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    ch.width = width
    ch.height = height
    # Sin esto el grafico se VACIA apenas el cliente usa el autofiltro de la hoja:
    # openpyxl escribe plotVisOnly=1 por defecto y Excel deja de graficar las filas
    # ocultas, incluidas las del bloque de datos que vive a la derecha.
    ch.visible_cells_only = False
    # Un hueco se dibuja como hueco y no como cero.
    ch.dispBlanksAs = "gap"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    if y_title:
        ch.y_axis.title = y_title
    ch.y_axis.scaling.min = 0
    if gap is not None:
        ch.gapWidth = gap
    if legend:
        ch.legend.position = legend
    else:
        ch.legend = None
    if show_values:
        ch.dLbls = DataLabelList()
        ch.dLbls.showVal = True
    for i, color in enumerate(colors):
        if i < len(ch.series):
            ch.series[i].graphicalProperties = GraphicalProperties(solidFill=color)
    return ch


def _chart_block(
    ws,
    top_row: int,
    title: str,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    *,
    start_col: int,
    fmt: str | None = None,
) -> tuple[int, int]:
    """Escribe el origen de un grafico, visible y rotulado, y devuelve sus filas.

    Visible a proposito: el cliente puede ver exactamente que filas entraron al
    grafico y cuales no. Un bloque oculto convierte el criterio de exclusion en
    algo que hay que creer.
    """
    _cell(ws, top_row, start_col, title, font=_BLOCK)
    header_row = top_row + 1
    _headers(ws, header_row, headers, start_col=start_col)
    for r, row in enumerate(rows, header_row + 1):
        for c, value in enumerate(row):
            _cell(ws, r, start_col + c, value, fmt=fmt if c > 0 else None)
    return header_row, header_row + len(rows)


def _cap(rows: list, ws, row: int, col: int, total: int, orden: str) -> list:
    """Aplica :data:`MAX_TABLE_ROWS` y, si topa, lo escribe en la hoja."""
    if total <= MAX_TABLE_ROWS:
        return rows
    _cell(
        ws,
        row,
        col,
        f"Se muestran las primeras {MAX_TABLE_ROWS:,} de {total:,} filas, ordenadas por "
        f"{orden}. El resto no se perdio: esta en el archivo que subiste y entra "
        "completo en el Diagnostico.".replace(",", "."),
        font=_NOTE,
        align=_WRAP,
    )
    return rows[:MAX_TABLE_ROWS]


def _isnumber_rule(col: str, first: int, last: int, test: str, fill: PatternFill) -> tuple:
    """Regla condicional que primero comprueba que la celda sea un numero.

    Sin el ``ISNUMBER`` esto pinta las celdas que dicen "no medido": Excel ordena
    numeros < texto, asi que cualquier ``>=`` contra un umbral da TRUE sobre
    texto, y la hoja terminaria afirmando en rojo justo donde no midio nada.

    Se usa **incluso donde la columna es demostrablemente numerica hoy**: asi es
    exactamente como se omite el guardia y como se rompe la primera vez que
    alguien mete un ``_measured()`` en esa columna.
    """
    return (
        f"{col}{first}:{col}{last}",
        FormulaRule(formula=[f"AND(ISNUMBER(${col}{first}),${col}{first}{test})"], fill=fill),
    )


def _link_row(ws, row: int, col: int, texto: str, url: str) -> None:
    """El CTA como hipervinculo absoluto y clickeable.

    Una ruta relativa en un ``.xlsx`` no lleva a ninguna parte: el archivo viaja
    fuera del navegador, normalmente adjunto en un correo al jefe de quien lo
    descargo.
    """
    c = _cell(ws, row, col, texto, font=_LINK)
    c.hyperlink = url


def _autosize(ws, *, min_w: int = 10, max_w: int = 52, skip: Collection[int] = ()) -> None:
    """Ancho por contenido. Las hojas de /demo NO lo usan y es deliberado.

    Se conserva porque es la convencion de ``src/excel_export.py``, pero las hojas
    de estos libros fijan anchos explicitos: el ``min_w`` aplasta las columnas de
    separacion de 3 caracteres, y medir el valor crudo deja las columnas de dinero
    en ``######`` porque el formato numerico agrega separadores que no se cuentan.
    """
    for col in ws.columns:
        idx = col[0].column
        if idx in skip:
            continue
        width = max((len(str(c.value or "")) for c in col), default=0) + 2
        ws.column_dimensions[get_column_letter(idx)].width = max(min_w, min(width, max_w))
