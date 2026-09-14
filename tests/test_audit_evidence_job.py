"""Job-layer tests for MUS + GL tie-out + EvidenceRecord workpaper.

The engine math lives in tests/test_audit_evidence.py. These tests prove the
playbook hashes inputs, never EXECUTED, leaves Reviewed-by blank, and does not
invent tolerable misstatement.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from jobs import audit_evidence_job
from nosilo_core.guided import ESCALATED, HANDOFF

_FIX = Path(__file__).resolve().parent / "fixtures" / "audit" / "population.csv"

_STABLE = {
    "tolerable_misstatement": 20000.0,
    "selection_seed": 0,
    "produced_at": "2026-09-07T12:00:00Z",
    "run_id": "test-audit-evidence",
}


def _run(path: Path | str = _FIX, extra: dict | None = None):
    params = {**_STABLE, **(extra or {})}
    prepared = audit_evidence_job.prepare(path, params)
    report = audit_evidence_job.run(prepared, params)
    return prepared, report


def test_prepare_splits_gl_and_subledger_from_source_column():
    prepared = audit_evidence_job.prepare(_FIX, {})
    assert prepared.has_tie_out is True
    assert len(prepared.gl_rows) == 10
    assert len(prepared.population_rows) == 10
    assert prepared.artifacts[0].sha256
    assert abs(sum(r["book_value"] for r in prepared.gl_rows) - 74000.0) < 1e-9


def test_missing_tm_is_required_auditor_judgment():
    prepared = audit_evidence_job.prepare(_FIX, {})
    try:
        audit_evidence_job.run(prepared, {})
    except ValueError as exc:
        assert "tolerable_misstatement" in str(exc)
    else:
        raise AssertionError("run() must refuse a missing TM")


def test_handoff_when_fieldwork_pending(tmp_path):
    _prepared, report = _run()
    assert report.fieldwork_pending is True
    assert report.evaluation is None
    assert report.outcome.status == HANDOFF
    assert report.evidence.prepared_by == audit_evidence_job.ENGINE_IDENTITY
    assert "human" not in report.evidence.prepared_by.lower() or "engine" in report.evidence.prepared_by
    assert audit_evidence_job.verify(report) == []
    written = audit_evidence_job.write_operational(report, tmp_path, "Acme")
    wb = load_workbook(written["workpaper"])
    assert [ws.title for ws in wb.worksheets] == [
        "W-1 Lead", "W-2 Sampling plan", "W-3 Selection", "W-4 Evaluation",
        "W-5 GL tie-out", "W-6 Lineage",
    ]
    for ws in wb.worksheets:
        reviewed = None
        for row in ws.iter_rows(min_col=1, max_col=2, values_only=True):
            if row[0] == "Reviewed by":
                reviewed = row[1]
        assert reviewed in (None, ""), f"{ws.title} Reviewed-by must stay blank"


def test_same_seed_is_reperformable():
    _, a = _run()
    _, b = _run()
    ids_a = [s.unit_id for s in a.selection.selected]
    ids_b = [s.unit_id for s in b.selection.selected]
    assert ids_a == ids_b
    assert a.evidence.output_sha256 == b.evidence.output_sha256
    assert a.random_start == b.random_start


def test_keyed_audited_values_evaluate_without_escalation(tmp_path):
    df = pd.read_csv(_FIX)
    df["audited_value"] = df["book_value"]
    path = tmp_path / "keyed.csv"
    df.to_csv(path, index=False)
    _prepared, report = _run(path)
    assert report.fieldwork_pending is False
    assert report.evaluation is not None
    assert report.evaluation.conclusion == "accept"
    assert report.outcome.status == HANDOFF
    assert audit_evidence_job.verify(report) == []


def test_tie_out_outside_tolerance_escalates(tmp_path):
    df = pd.read_csv(_FIX)
    # Break the subledger control total.
    mask = (df["source"] == "subledger") & (df["unit_id"] == "SKU-A")
    df.loc[mask, "book_value"] = 1000.0
    path = tmp_path / "broken.csv"
    df.to_csv(path, index=False)
    _prepared, report = _run(path, extra={"tolerance": 0.0})
    assert report.outcome.status == ESCALATED
    assert report.tie_out is not None
    assert report.tie_out.within_tolerance is False
    assert audit_evidence_job.verify(report) == []


def test_malformed_tm_validate_params():
    try:
        audit_evidence_job.validate_params({"tolerable_misstatement": "nope"})
    except ValueError as exc:
        assert "tolerable_misstatement" in str(exc)
    else:
        raise AssertionError("malformed TM must raise")
