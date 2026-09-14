"""Sin costo unitario en el archivo, el entregable habla de UNIDADES y no imprime "$".

La grieta de honestidad #1 del diagnostico del 2026-09-05: ``jobs/intake.py`` rellena
``unit_cost`` con 1,0 cuando el archivo no trae la columna, el motor corre igual, y
el ``.md``/``.xlsx`` imprimian "$244" donde habia 244 unidades. El aviso interno
(``IntakeQuality.synthesized``) existia y moria antes del deck. Estos tests fijan que
cada monto de inversion del entregable lleva la unidad correcta, y que la nota que lo
explica llega al cliente en el mismo documento.
"""

from __future__ import annotations

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from jobs import deliverables, inventory_deliverable, inventory_optimization
from jobs.intake import IntakeQuality


def _demand(unit_cost: float, n_skus: int = 2, periods: int = 60) -> pd.DataFrame:
    rng = np.random.default_rng(7)
    rows = []
    for s in range(n_skus):
        for t in range(periods):
            rows.append({
                "product_id": f"SKU-{s + 1}",
                "date": pd.Timestamp("2024-01-01") + pd.Timedelta(weeks=t),
                "quantity": float(max(0.0, 100 + 10 * rng.standard_normal())),
                "unit_cost": unit_cost,
                "lead_time_days": 14.0,
            })
    return pd.DataFrame(rows)


def _sin_costo():
    """Lo que intake produce cuando el archivo no trae costo: 1,0 en cada fila y el
    conteo de celdas inventadas en ``synthesized``."""
    demand = _demand(unit_cost=1.0)
    n = demand["product_id"].nunique()
    quality = IntakeQuality(
        n_raw=len(demand), n_dropped_bad_date=0, n_dropped_bad_quantity=0,
        n_dropped_negative_quantity=0, synthesized=(("unit_cost", n),), n_canonical_rows=n,
    )
    return inventory_optimization.run(demand, intake_quality=quality)


def _con_costo():
    demand = _demand(unit_cost=12.0)
    n = demand["product_id"].nunique()
    quality = IntakeQuality(
        n_raw=len(demand), n_dropped_bad_date=0, n_dropped_bad_quantity=0,
        n_dropped_negative_quantity=0, synthesized=(), n_canonical_rows=n,
    )
    return inventory_optimization.run(demand, intake_quality=quality)


def test_sin_costo_el_informe_no_imprime_dolares_y_avisa(tmp_path):
    report = _sin_costo()
    assert deliverables.cost_is_synthesized(report)

    texto = deliverables.write_report_md(report, tmp_path / "r.md").read_text(encoding="utf-8")

    resumen = texto.split("## ")[1]  # el Executive summary entero
    assert "$" not in resumen, resumen
    assert " u." in resumen
    assert deliverables.UNITS_NOTE in texto
    # la tabla por SKU tampoco: cada fila de inversion va en unidades
    filas = [ln for ln in texto.splitlines() if ln.startswith("| SKU-")]
    assert filas and all("$" not in ln and " u. |" in ln for ln in filas), filas


def test_con_costo_el_informe_sigue_hablando_en_dolares(tmp_path):
    report = _con_costo()
    assert not deliverables.cost_is_synthesized(report)

    texto = deliverables.write_report_md(report, tmp_path / "r.md").read_text(encoding="utf-8")

    assert "Recommended inventory investment is **$" in texto
    assert deliverables.UNITS_NOTE not in texto


def test_sin_costo_la_planilla_rotula_las_celdas_en_unidades(tmp_path):
    report = _sin_costo()
    ruta = deliverables.write_excel(report, tmp_path / "r.xlsx")
    wb = load_workbook(ruta)

    headers = [c.value for c in wb["Recommendations"][1]]
    assert "inventory units no cost" in headers and "inventory value" not in headers

    resumen = {row[0].value: row[1].value for row in wb["Summary & Assumptions"].iter_rows(min_row=1)}
    assert resumen["Note"] == deliverables.UNITS_NOTE
    assert "Final investment (units, no unit cost loaded)" in resumen
    assert "Final investment ($)" not in resumen


def test_con_costo_la_planilla_conserva_los_rotulos_en_dolares(tmp_path):
    report = _con_costo()
    wb = load_workbook(deliverables.write_excel(report, tmp_path / "r.xlsx"))

    headers = [c.value for c in wb["Recommendations"][1]]
    assert "inventory value" in headers
    resumen = {row[0].value: row[1].value for row in wb["Summary & Assumptions"].iter_rows(min_row=1)}
    assert "Final investment ($)" in resumen and "Note" not in resumen


def test_sin_costo_el_deck_consolidado_tambien_habla_en_unidades():
    report = _sin_costo()
    deck = inventory_deliverable.build(report)

    kpi = next(k for k in deck.kpis if k.name.startswith("Recommended investment"))
    assert "units" in kpi.name
    assert "$" not in kpi.value and kpi.value.endswith(" u.")


def test_con_costo_el_deck_consolidado_muestra_dolares():
    report = _con_costo()
    deck = inventory_deliverable.build(report)

    kpi = next(k for k in deck.kpis if k.name == "Recommended investment")
    assert kpi.value.startswith("$")
