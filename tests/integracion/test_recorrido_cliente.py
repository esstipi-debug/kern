"""El recorrido del cliente, contra un Odoo sembrado como lo carga un usuario.

``test_odoo_real.py`` prueba el CONTRATO del conector sobre productos que crea
el mismo test, con el proveedor puesto en la variante. Este archivo prueba las
PROMESAS del correo a los partners sobre un distribuidor completo, con el
proveedor en la ficha del producto (como lo deja el formulario de Odoo 17). Nace
de la prueba del 2026-09-19 (``kern-fase0/prueba-recorrido-2026-09-19/INFORME.md``),
donde de cuatro promesas se cumplia una:

=====================================================  =====================================
Promesa del correo                                     Test
=====================================================  =====================================
"lee plazos de proveedor"                              ``test_lee_el_plazo_...`` (F1)
"propone ... ordenes de compra en borrador"            ``test_crea_oc_en_borrador_...`` (F1, F4, F5)
el punto de reorden cubre el plazo del proveedor       ``test_la_propuesta_de_un_sku_...`` (METODO)
no compra de mas lo intermitente                       ``test_la_transpaleta_...`` (METODO)
entregable para un partner hispanohablante             ``test_el_entregable_...`` (F11)
=====================================================  =====================================

**Como correr todo** (desde la raiz del repo, con el python del venv):

    docker compose --env-file .env.odoo -f docker-compose.odoo.yml down -v   # base limpia
    python scripts/odoo_test_setup.py              # Odoo 17 + purchase/sale/stock + grupos
    python scripts/odoo_sembrar_distribuidor.py    # el distribuidor: 20 SKU, 26 semanas
    PYTHONPATH=. python -m pytest -m integracion tests/integracion -v

``-m integracion`` es obligatorio: ``pyproject.toml`` los deselecciona por
defecto (``addopts = "-m 'not integracion'"``). El sembrador se niega a correr
sobre una base que ya tiene el catalogo, asi que para repetir desde cero hay que
volver a empezar por el ``down -v``. Otro Odoo: ``ODOO_URL``, ``ODOO_DB``,
``ODOO_USERNAME`` y ``ODOO_API_KEY``, para el sembrador y para los tests.

**Lo que escribe en Odoo.** Solo ``test_crea_oc_en_borrador_...``: crea las OC
en borrador y al terminar las cancela y las borra, pase o falle. Los demas leen.
``test_odoo_real.py`` corre antes (orden alfabetico) y deja en la base productos
``KERN-*`` con proveedor sin plazo y OC en borrador: estos tests miran solo el
catalogo del sembrador, y un borrador no cuenta como en camino.

**Que valores espera, y de donde salen.** De la base, no de Kern: el plazo y
el precio de ``product.supplierinfo``, lo vendido de los pedidos confirmados, y
el catalogo de ``scripts/odoo_sembrar_distribuidor.py``. Comparar a Kern contra
su propia lectura no probaria nada.

**METODO tiene que cablear el job.** La politica min/max vive en ``plan_replenishment``,
pero solo la usa quien le pasa ``lead_time_days_by_sku``. Si el job
(``jobs/odoo_job.py::run``) no se los pasa, la corrida sigue con la politica de
cobertura, y los tests de SKU-006/013/014 y de la transpaleta siguen en rojo con
la politica nombrada en el mensaje.
"""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from nosilo_core.llm import RulesFallback
from nosilo_core.writeback_store import SqliteAuditLedger
from scm_agent.orchestrator import Orchestrator
from scripts.odoo_sembrar_distribuidor import (
    CATALOGO,
    DIAS_HISTORIA,
    PROVEEDORES,
    SKU_INTERMITENTE,
    SKU_PLAZO_SIN_CONFIGURAR,
    SKU_SIN_PROVEEDOR,
)
from src.connectors.odoo import OdooConnector

pytestmark = pytest.mark.integracion

APROBADOR = "Comprador de prueba (recorrido)"
CODIGOS = [item.codigo for item in CATALOGO]
#: Los SKU de venta regular del proveedor de plazo mas largo (45 dias).
PLAZO_MAS_LARGO = max(p.plazo_declarado for p in PROVEEDORES)
SKUS_PLAZO_LARGO = [
    item.codigo for item in CATALOGO
    if item.proveedor is not None and item.venta_semanal > 0
    and PROVEEDORES[item.proveedor].plazo_declarado == PLAZO_MAS_LARGO
]
#: Lo que el deck de ``odoo_replenishment`` decia en ingles el 2026-09-19. Frases
#: del propio tool (``jobs/odoo_job.py``) y del armazon (``src/deliverable.py``),
#: incluida la reversibilidad que la marcha atras no cumplia (F7). NO esta el
#: banner de escalamiento de ``nosilo_core/escalation.py`` ("... requires finance
#: sign-off before committing", en minuscula): eso lo decide D08 del plan.
FRASES_EN_INGLES = (
    "Executive summary", "Key findings", "Recommendations", "Options to act", "Coverage & handoff",
    "Prepared by", "Requires finance sign-off", "do not apply until", "Lead times the ERP does not record",
    "below target cover", "Thinnest SKU", "Products read", "SKUs to replenish", "Units short", "Cover target",
    "Apply reorder points", "fully reversible", "(impact:",
)
CONFIANZA = re.compile(r"(?:Confidence|Confianza):\s*\*\*(\d+)%\*\*")


# --------------------------------------------------------------------------
# Fixtures
# --------------------------------------------------------------------------


@pytest.fixture(scope="module")
def sembrado(cliente) -> None:
    """La base tiene el distribuidor, en unidades, y con el proveedor como lo deja el formulario.

    Falla en vez de saltarse, igual que ``cliente``. Tres comprobaciones:

    - que esten los 20 SKU;
    - que la unidad de compra de todos sea de la categoria "Unit" de Odoo
      (``uom.product_uom_categ_unit``). Es la premisa de "cantidades enteras" en
      ``test_crea_oc_...``: F4 redondea a la unidad solo en esa categoria, y un SKU
      en cajas, kilos o litros saldria con decimales a proposito, no por defecto.
      Se lee de la base y no de Kern, asi que vale antes y despues de F4;
    - que ninguna linea de proveedor tenga variante. La prueba del 2026-09-19
      asigno la variante a mano a mitad de camino (el "control" despues del paso
      04a), y sobre esa base el test de plazos pasaria sin probar el caso por el
      que existe.
    """
    filas = cliente.execute_kw("product.product", "search_read",
                               [[["default_code", "in", CODIGOS]]], {"fields": ["default_code"]})
    faltan = sorted(set(CODIGOS) - {f["default_code"] for f in filas})
    if faltan:
        pytest.fail(f"la base no tiene {len(faltan)} SKU del distribuidor ({faltan[:3]}...). "
                    "Siembrala con: python scripts/odoo_sembrar_distribuidor.py")
    (unidad,) = cliente.execute_kw("ir.model.data", "search_read", [[
        ["module", "=", "uom"], ["name", "=", "product_uom_categ_unit"]]], {"fields": ["res_id"]})
    fuera_de_unidad = cliente.execute_kw("product.product", "search_read", [[
        ["default_code", "in", CODIGOS], ["uom_po_id.category_id", "!=", unidad["res_id"]]]],
        {"fields": ["default_code"]})
    if fuera_de_unidad:
        pytest.fail(f"SKU con unidad de compra fuera de la categoria Unit: "
                    f"{sorted(f['default_code'] for f in fuera_de_unidad)}. El test de cantidades "
                    "enteras supone unidades: esta base no es la del sembrador.")
    con_variante = cliente.execute_kw("product.supplierinfo", "search_count", [[
        ["product_tmpl_id.default_code", "in", CODIGOS], ["product_id", "!=", False]]])
    if con_variante:
        pytest.fail(f"{con_variante} linea(s) de proveedor del catalogo tienen variante: esta base no es "
                    "la del formulario. Re-sembrar desde cero (docker compose ... down -v).")


@pytest.fixture(scope="module")
def corrida(sembrado, credenciales_odoo, tmp_path_factory):
    """Una corrida de ``odoo_replenishment`` como la del operador (``run_agent.py --job``).

    ``RulesFallback``: sin LLM, el texto del entregable no depende de una clave
    en la maquina. ``lang="es"``: si el deck en espanol llega por el idioma del
    orquestador, este es el camino; si llega por defecto, no molesta. Sin
    ``strict_params``, como el operador: el nivel de servicio que no viene en
    ``overrides`` sale del default y el reporte lo declara (METODO); con
    ``strict_params=True`` habria que pasarlo.
    """
    resultado = Orchestrator(provider=RulesFallback(), clients_root=None, lang="es").run(
        "reposicion desde odoo", job_type="odoo_replenishment", overrides=dict(credenciales_odoo),
        client="Distribuidora Prueba", out_dir=tmp_path_factory.mktemp("entregable"),
    )
    assert resultado.status == "ok", (resultado.status, resultado.summary, resultado.qa_issues)
    return resultado


@pytest.fixture(scope="module")
def reporte(corrida):
    return corrida.report


@pytest.fixture
def ledger() -> SqliteAuditLedger:
    libro = SqliteAuditLedger(":memory:")
    yield libro
    libro.close()


@pytest.fixture
def oc_a_limpiar(cliente):
    """Las OC que crea un test, canceladas y borradas al terminar, pase o falle.

    Odoo 17 no borra una OC sin cancelar (``purchase/models/purchase_order.py:264``,
    ``_unlink_if_cancelled``). Se cancela escribiendo el estado, que es lo que hace
    ``button_cancel`` por dentro (``purchase_order.py:522``), y no llamandolo: ese
    metodo devuelve ``None``, el XML-RPC de Odoo no lo sabe
    serializar (``base/controllers/rpc.py:128``, ``allow_none=False``) y contesta con
    un fault DESPUES de cancelar. Un borrador no tiene movimientos que deshacer.
    """
    ids: list[int] = []
    yield ids
    if ids:
        cliente.execute_kw("purchase.order", "write", [ids, {"state": "cancel"}])
        cliente.execute_kw("purchase.order", "unlink", [ids])


def _proveedor_por_sku(cliente) -> dict[str, dict]:
    """La linea de ``product.supplierinfo`` de cada SKU del catalogo, por codigo."""
    filas = cliente.execute_kw("product.supplierinfo", "search_read",
                               [[["product_tmpl_id.default_code", "in", CODIGOS]]],
                               {"fields": ["product_tmpl_id", "delay", "price"]})
    plantillas = sorted({f["product_tmpl_id"][0] for f in filas})
    codigo = {r["id"]: r["default_code"] for r in cliente.execute_kw(
        "product.template", "read", [plantillas], {"fields": ["default_code"]})}
    return {codigo[f["product_tmpl_id"][0]]: f for f in filas}


def _vendido(cliente, sku: str) -> float:
    """Unidades vendidas en pedidos confirmados, leidas de Odoo y no de Kern."""
    filas = cliente.execute_kw("sale.order.line", "search_read",
                               [[["product_id.default_code", "=", sku], ["state", "in", ["sale", "done"]]]],
                               {"fields": ["product_uom_qty"]})
    return sum(f["product_uom_qty"] for f in filas)


# --------------------------------------------------------------------------
# "lee ventas, stock y plazos de proveedor"
# --------------------------------------------------------------------------


def test_lee_el_plazo_de_cada_sku_con_el_proveedor_cargado_en_la_ficha(cliente, sembrado):
    """F1: con el proveedor en la plantilla Kern leia 0 plazos de 18."""
    declarado = {sku: float(f["delay"]) for sku, f in _proveedor_por_sku(cliente).items() if f["delay"] > 1}

    leido = OdooConnector(cliente).declared_lead_times().declared

    faltan = sorted(sku for sku, dias in declarado.items() if leido.get(sku) != pytest.approx(dias))
    assert not faltan, (f"sin el plazo del proveedor en {len(faltan)} de {len(declarado)} SKU con el "
                        f"proveedor cargado en la ficha: {faltan}")


def test_el_plazo_sin_configurar_se_declara_ausente_y_el_sku_sin_proveedor_no_aparece(cliente, sembrado):
    """``delay=1`` es el default de Odoo: se declara ausente, no se lee como 1 dia.
    Y un SKU sin linea de proveedor no tiene un plazo ausente: no tiene proveedor."""
    plazos = OdooConnector(cliente).declared_lead_times()

    del_catalogo = sorted(sku for sku in plazos.missing if sku in CODIGOS)
    assert SKU_PLAZO_SIN_CONFIGURAR in plazos.missing, f"plazos ausentes del catalogo: {del_catalogo}"
    assert SKU_PLAZO_SIN_CONFIGURAR not in plazos.declared
    assert SKU_SIN_PROVEEDOR not in plazos.declared and SKU_SIN_PROVEEDOR not in plazos.missing


# --------------------------------------------------------------------------
# "propone puntos de reorden": la cantidad cubre el plazo, y no infla lo intermitente
# --------------------------------------------------------------------------


@pytest.mark.parametrize("sku", SKUS_PLAZO_LARGO)
def test_la_propuesta_de_un_sku_con_plazo_largo_cubre_la_demanda_del_plazo(cliente, reporte, sku):
    """METODO: SKU-006 (45 dias de plazo) quedaba con ~3 semanas de venta y pedia 0.

    La demanda del plazo sale de lo vendido en la ventana sembrada (26 semanas
    de calendario) por el plazo declarado del proveedor. Lo que entra en camino
    mas lo que se pide tiene que alcanzar para esperar al proveedor; cualquier
    politica de reposicion sensata lo cumple -- es la definicion minima de un
    punto de reorden, sin stock de seguridad todavia.

    El mensaje nombra la politica de la linea (``policy``, el campo que agrega METODO)
    solo para el diagnostico: ``None`` es el codigo de antes de METODO, y ``cover`` es
    METODO sin cablear el job (el job no le paso los plazos al plan). La afirmacion no depende de
    ese campo.
    """
    item = next(i for i in CATALOGO if i.codigo == sku)
    plazo = PROVEEDORES[item.proveedor].plazo_declarado
    demanda_del_plazo = _vendido(cliente, sku) / DIAS_HISTORIA * plazo
    linea = next(ln for ln in reporte.lines if ln.sku == sku)
    pedido = reporte.restock.get(sku, 0.0)

    assert linea.position is not None, f"{sku}: la linea no se dimensiono ({linea.status})"
    assert linea.position + pedido >= demanda_del_plazo, (
        f"{sku} (plazo {plazo} d, politica {getattr(linea, 'policy', None)}): posicion "
        f"{linea.position:g} + pedido {pedido:g} = {linea.position + pedido:g} < demanda del plazo "
        f"{demanda_del_plazo:.1f}"
    )


def test_la_transpaleta_intermitente_no_pide_siete(cliente, reporte):
    """METODO: 6 ventas de 1 unidad en 26 semanas daban un pronostico de 1 por
    "periodo" (un dia con venta) y una compra de 7: el 55% del plan en plata.

    El tope es lo vendido en toda la ventana: pedir mas transpaletas de las que
    se vendieron en seis meses no lo justifica ningun plazo de 45 dias. Y la
    linea tiene que estar dimensionada: "no pide" por haberse negado a calcular
    (``position`` en ``None``) no es lo mismo que decidir no pedir.
    """
    vendidas = _vendido(cliente, SKU_INTERMITENTE)
    linea = next(ln for ln in reporte.lines if ln.sku == SKU_INTERMITENTE)
    pedido = reporte.restock.get(SKU_INTERMITENTE, 0.0)

    assert vendidas == 6, f"la base no es la sembrada: {SKU_INTERMITENTE} vendio {vendidas}"
    assert linea.position is not None, f"{SKU_INTERMITENTE} no se dimensiono ({linea.status})"
    assert pedido <= vendidas, (f"{SKU_INTERMITENTE} pide {pedido:g}: mas que las {vendidas:g} que vendio "
                                f"en 26 semanas (el 2026-09-19 pedia 7; politica {getattr(linea, 'policy', None)})")


# --------------------------------------------------------------------------
# "ordenes de compra en borrador"
# --------------------------------------------------------------------------


def test_crea_oc_en_borrador_con_precio_del_proveedor_y_cantidades_enteras(
    cliente, reporte, ledger, oc_a_limpiar
):
    """F1 + F4 + F5 sobre la propuesta del plan, tal como la aprobaria un comprador.

    El 2026-09-19: 0 OC y 9 SKU "sin proveedor" (F1); con la variante puesta a mano,
    lineas a precio 0 (F5) y con 187.6 rollos (F4). El precio esperado es el de
    ``product.supplierinfo`` y NO el costo del producto: el sembrador los separa
    a proposito, asi que este test distingue los dos.

    "Cantidades enteras" supone productos en la categoria Unit, y ``sembrado`` lo
    comprueba en la base. Si igual sale un decimal hay dos lecturas: que F4 no
    redondeo (defecto), o que Kern no pudo leer la categoria de la unidad y el SKU
    quedo en ``OdooUnitSteps.missing`` (a proposito: sin paso medido no se inventa
    uno). El mensaje de la falla dice como distinguirlas.
    """
    precio = {sku: f["price"] for sku, f in _proveedor_por_sku(cliente).items()}
    conector = OdooConnector(cliente, ledger=ledger)

    resultado = conector.create_draft_purchase_orders(dict(reporte.restock), approved_by=APROBADOR)
    oc_a_limpiar.extend(resultado.purchase_orders)

    assert set(resultado.unsourced) <= {SKU_SIN_PROVEEDOR}, (
        f"SKU con proveedor en la ficha reportados sin proveedor: {sorted(resultado.unsourced)}")
    assert resultado.n_orders >= 1, "la propuesta no genero ninguna OC"
    ordenes = cliente.execute_kw("purchase.order", "read", [list(resultado.purchase_orders)],
                                 {"fields": ["state"]})
    assert {o["state"] for o in ordenes} == {"draft"}, ordenes
    lineas = cliente.execute_kw("purchase.order.line", "search_read",
                                [[["order_id", "in", list(resultado.purchase_orders)]]],
                                {"fields": ["product_id", "product_qty", "price_unit"]})
    codigo = {r["id"]: r["default_code"] for r in cliente.execute_kw(
        "product.product", "read", [sorted({ln["product_id"][0] for ln in lineas})], {"fields": ["default_code"]})}
    por_linea = [(codigo[ln["product_id"][0]], ln["product_qty"], ln["price_unit"]) for ln in lineas]
    fraccionadas = {sku: q for sku, q, _p in por_linea if not float(q).is_integer()}
    mal_precio = {sku: (p, precio[sku]) for sku, _q, p in por_linea if p != pytest.approx(precio[sku])}
    assert not fraccionadas, (f"cantidades con decimales en Odoo: {fraccionadas}. Si esos SKU estan en "
                              "OdooConnector.unit_step_by_sku().missing, el usuario de API no lee "
                              "uom.category y el decimal es a proposito")
    assert not mal_precio, f"precio de la linea distinto del proveedor (linea, proveedor): {mal_precio}"
    claves = sorted(ledger.applied_keys())
    assert len(claves) == 1, f"una aprobacion por propuesta, y el ledger tiene {claves}"
    assert ledger.get(claves[0]).approved_by == APROBADOR


# --------------------------------------------------------------------------
# El entregable
# --------------------------------------------------------------------------


def test_el_entregable_sale_en_espanol_y_sin_confianza_de_cien(corrida):
    """F11: salia en ingles y con "Confidence: 100%", que era la confianza del
    ENRUTADOR al elegir la herramienta (1.0 cuando se fuerza), no la del plan."""
    texto = Path(corrida.deliverables["deck_report"]).read_text(encoding="utf-8")

    en_ingles = [frase for frase in FRASES_EN_INGLES if frase in texto]
    assert "Resumen ejecutivo" in texto and not en_ingles, f"el entregable sigue en ingles: {en_ingles}"
    confianza = CONFIANZA.search(texto)
    assert confianza is None or int(confianza.group(1)) < 100, confianza.group(0)
    resumen = load_workbook(corrida.deliverables["deck_workbook"]).worksheets[0]
    filas = {fila[0]: fila[1] for fila in resumen.iter_rows(values_only=True) if len(fila) > 1}
    assert resumen.title == "Resumen", f"la planilla sigue en ingles: hoja {resumen.title!r}"
    assert "100%" not in (filas.get("Confianza"), filas.get("Confidence")), filas
