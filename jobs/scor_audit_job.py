"""SCOR + ISO 9001/28000 preparation pack (E4): compose existing jobs, wrap EvidenceRecord.

Decision (b): a dedicated job, not a PackageSpec. The MCP surface is one ``rows`` CSV
(``_run_analysis_tool``); ``packages.py`` expects a directory of named intake files
and writes per-tool decks rather than one EvidenceRecord + Lineage sheet. This job
calls the same ``prepare``/``run`` functions the orchestrator already drives
(financial_kpis, cost_to_serve, sourcing, procedure_audit, risk) plus
``src.order_fulfillment`` for SCOR RL.1.1, then wraps the lot.

A missing panel is ``no_evaluable`` (like cross_impact), never "nobody objects" and
never "ISO compliant". Kern prepares evidence for a registrar's auditor; it does
not certify.
"""

from __future__ import annotations

import json
import math
import re
import tempfile
import uuid
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd

from jobs import (
    cost_to_serve_job,
    financial_kpis_job,
    procedure_audit_job,
    risk_job,
    sourcing_job,
)
from nosilo_core.guided import (
    EXECUTED,
    HandoffPacket,
    Residual,
    as_handoff,
    confidence_from_coverage,
    verify_guided,
)
from src.audit_evidence import (
    FORMULA_VERSION,
    InputArtifact,
    QaCheck,
    evidence_record,
    evidence_to_dict,
    hash_file,
)
from src.deliverable import DataSource, Deliverable, Finding, Kpi, money, pct
from src.export import write_summary_csv
from src.iso_catalog import texto_prohibido
from src.order_fulfillment import aggregate, measure_from_order_lines
from src.sanitize import defuse_formula

SCOR_ISO_VERSION = "1"
ENGINE_IDENTITY = f"linchpin/scor_iso_audit v{SCOR_ISO_VERSION} (engine)"
_DECK_BASE_CONFIDENCE = 0.75

OBSERVED = "observed"
GAP = "gap"
NO_EVALUABLE = "no_evaluable"
RAN = "ran"
SKIPPED = "skipped"

_DATASET_COL = ("dataset", "panel", "block", "source_panel", "bloque")
_DATASET_ALIASES: dict[str, str] = {}
for _canon, _names in {
    "financial": ("financial", "financial_kpis", "finanzas", "kpis", "finance"),
    "cost_to_serve": ("cost_to_serve", "cts", "pedidos", "serve", "costtoserve"),
    "sourcing": ("sourcing", "suppliers", "proveedores", "supplier", "proveedor"),
    "orders": ("orders", "fulfillment", "otif", "order", "pedidos_otif"),
    "procedures": ("procedures", "procedure", "sop", "register", "procedimientos"),
    "risk": ("risk", "riesgos", "risks", "riesgo"),
}.items():
    for _n in _names:
        _DATASET_ALIASES[_n] = _canon

_EN_FORBIDDEN = re.compile(
    r"\b(iso[\s-]*certified|certified[\s-]*iso|complies with iso|iso[\s-]*compliant|"
    r"kern certifies|certifies iso)\b",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class PanelCoverage:
    panel: str
    status: str  # ran | skipped
    reason: str
    summary: str = ""


@dataclass(frozen=True)
class ScorMetric:
    code: str
    name: str
    value: float | None
    status: str  # observed | no_evaluable
    source_panel: str
    note: str


@dataclass(frozen=True)
class IsoControlCheck:
    clause: str
    title: str
    status: str  # observed | gap | no_evaluable
    evidence: str


@dataclass(frozen=True)
class PreparedScorPack:
    frames: dict[str, pd.DataFrame]
    skipped: tuple[PanelCoverage, ...]
    artifacts: tuple[InputArtifact, ...]
    columns: tuple[str, ...]


@dataclass(frozen=True)
class ScorIsoReport:
    metrics: tuple[ScorMetric, ...]
    checks: tuple[IsoControlCheck, ...]
    coverage: tuple[PanelCoverage, ...]
    evidence: object
    outcome: object
    summary: str
    n_observed: int
    n_gap: int
    n_no_evaluable: int


def _pick_column(df: pd.DataFrame, candidates: tuple[str, ...]) -> str | None:
    return next((c for c in candidates if c in df.columns), None)


def _norm_dataset(value: object) -> str | None:
    text = str(value or "").strip().lower().replace(" ", "_").replace("-", "_")
    return _DATASET_ALIASES.get(text)


def _iso_claim_forbidden(text: str) -> bool:
    return texto_prohibido(text) or bool(_EN_FORBIDDEN.search(text or ""))


def _read_table(path: str | Path) -> pd.DataFrame:
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"{path} does not exist")
    return pd.read_excel(p) if p.suffix.lower() in {".xlsx", ".xls"} else pd.read_csv(p)


def prepare(data_path: str, params: dict | None = None) -> PreparedScorPack:
    """Split a tagged CSV. ``dataset`` is required (MCP / data-contract)."""
    del params
    df = _read_table(data_path)
    if df.empty:
        raise ValueError("input has no rows")
    ds_col = _pick_column(df, _DATASET_COL)
    frames: dict[str, pd.DataFrame] = {}
    skipped: list[PanelCoverage] = []
    if not ds_col:
        raise ValueError(
            "could not tell which SCOR/ISO panel this table is. Add a 'dataset' column "
            "with values financial / cost_to_serve / sourcing / orders / procedures / risk "
            f"(columns seen: {list(df.columns)[:12]})"
        )
    unknown = 0
    for raw, group in df.groupby(ds_col, dropna=False):
        canon = _norm_dataset(raw)
        if canon is None:
            unknown += len(group)
            continue
        frames[canon] = pd.concat([frames[canon], group], ignore_index=True) if canon in frames else group.copy()
    if unknown:
        skipped.append(PanelCoverage(
            "unmapped_dataset", SKIPPED,
            f"{unknown} row(s) had a dataset tag that is not one of: "
            + ", ".join(sorted(set(_DATASET_ALIASES.values()))),
        ))

    if not frames:
        raise ValueError(
            "no recognised dataset panel. Tag rows with dataset=financial|cost_to_serve|"
            "sourcing|orders|procedures|risk"
        )

    hashed = hash_file(data_path)
    control = None
    amount_col = _pick_column(df, ("book_value", "cogs", "sales", "Sales", "amount"))
    if amount_col:
        try:
            control = float(pd.to_numeric(df[amount_col], errors="coerce").fillna(0).sum())
        except (TypeError, ValueError):
            control = None
    artifact = InputArtifact(
        path_label=Path(data_path).name,
        sha256=hashed,
        n_rows=len(df),
        columns=tuple(str(c) for c in df.columns),
        control_total=control,
    )
    return PreparedScorPack(
        frames=frames,
        skipped=tuple(skipped),
        artifacts=(artifact,),
        columns=tuple(str(c) for c in df.columns),
    )


def _write_panel(df: pd.DataFrame, tmp: Path, name: str) -> str:
    # Drop the dataset tag so the child job's sniffer is not confused by extra columns.
    work = df.drop(columns=[c for c in df.columns if str(c).lower() in {x.lower() for x in _DATASET_COL}], errors="ignore")
    path = tmp / f"{name}.csv"
    work.to_csv(path, index=False)
    return str(path)


def _try_panel(name: str, fn) -> tuple[object | None, PanelCoverage]:
    try:
        result = fn()
        return result, PanelCoverage(name, RAN, "composed from existing job", getattr(result, "summary", "") or str(type(result).__name__))
    except (ValueError, FileNotFoundError, KeyError, TypeError) as exc:
        return None, PanelCoverage(name, SKIPPED, str(exc))


def _order_lines(df: pd.DataFrame) -> list[dict]:
    ordered = _pick_column(df, ("ordered_qty", "quantity_ordered", "qty_ordered", "ordered", "Quantity"))
    delivered = _pick_column(df, ("delivered_qty", "quantity_delivered", "qty_delivered", "delivered", "received_qty"))
    if ordered is None or delivered is None:
        raise ValueError("orders panel needs ordered_qty and delivered_qty (or aliases)")
    product = _pick_column(df, ("product_id", "sku", "SKU", "item"))
    promised = _pick_column(df, ("promised_date", "due_date", "Promised_Date"))
    shipped = _pick_column(df, ("delivered_date", "ship_date", "delivery_date", "Delivery_Date"))
    price = _pick_column(df, ("unit_price", "price", "Unit_Price"))
    lines: list[dict] = []
    for _, row in df.iterrows():
        line = {
            "product_id": str(row[product]).strip() if product else "PORTFOLIO",
            "ordered_qty": row[ordered],
            "delivered_qty": row[delivered],
        }
        if promised:
            line["promised_date"] = row[promised]
        if shipped:
            line["delivered_date"] = row[shipped]
        if price:
            line["unit_price"] = row[price]
        lines.append(line)
    return lines


def run(payload: PreparedScorPack, params: dict | None = None) -> ScorIsoReport:
    """Compose whichever panels the CSV carried; wrap in one EvidenceRecord."""
    params = params or {}
    produced_at = str(params.get("produced_at") or datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"))
    run_id = str(params.get("run_id") or uuid.uuid4())
    as_of = params.get("as_of")
    dso = params.get("dso_days", params.get("dso"))
    dpo = params.get("dpo_days", params.get("dpo"))
    dso_f = float(dso) if dso is not None else None
    dpo_f = float(dpo) if dpo is not None else None

    coverage: list[PanelCoverage] = list(payload.skipped)
    metrics: list[ScorMetric] = []
    checks: list[IsoControlCheck] = []

    with tempfile.TemporaryDirectory(prefix="scor_iso_") as tmp:
        tmp_path = Path(tmp)

        fin_report = None
        if "financial" in payload.frames:
            df = payload.frames["financial"]
            def _fin():
                records = financial_kpis_job.prepare_records(df, params)
                return financial_kpis_job.run(records, dso=dso_f, dpo=dpo_f)
            fin_report, cov = _try_panel("financial", _fin)
            coverage.append(cov)
            if fin_report is not None:
                issues = financial_kpis_job.verify(fin_report)
                if issues:
                    coverage[-1] = PanelCoverage("financial", SKIPPED, "; ".join(issues))
                    fin_report = None
                else:
                    metrics.append(ScorMetric(
                        "AM.1.1", "cash-to-cash (days)", fin_report.cash_to_cash,
                        OBSERVED if fin_report.cash_to_cash is not None else NO_EVALUABLE,
                        "financial",
                        "needs DIO from inventory plus DSO/DPO (params dso_days, dpo_days) to measure"
                        if fin_report.cash_to_cash is None else
                        f"turns={fin_report.turns} DIO={fin_report.dio}",
                    ))
                    metrics.append(ScorMetric(
                        "AM.inventory_turns", "inventory turns", fin_report.turns,
                        OBSERVED if fin_report.turns is not None else NO_EVALUABLE,
                        "financial", "from COGS / average inventory",
                    ))
        else:
            coverage.append(PanelCoverage("financial", SKIPPED, "no dataset=financial rows"))
            metrics.append(ScorMetric("AM.1.1", "cash-to-cash (days)", None, NO_EVALUABLE, "financial", "panel absent"))

        cts_report = None
        if "cost_to_serve" in payload.frames:
            path = _write_panel(payload.frames["cost_to_serve"], tmp_path, "cost_to_serve")
            def _cts():
                activities = cost_to_serve_job.prepare(path, params)
                return cost_to_serve_job.run(activities)
            cts_report, cov = _try_panel("cost_to_serve", _cts)
            coverage.append(cov)
            if cts_report is not None:
                issues = cost_to_serve_job.verify(cts_report)
                if issues:
                    coverage[-1] = PanelCoverage("cost_to_serve", SKIPPED, "; ".join(issues))
                    cts_report = None
                else:
                    port = cts_report.portfolio
                    metrics.append(ScorMetric(
                        "CO.1", "cost-to-serve net margin", port.overall_net_margin,
                        OBSERVED, "cost_to_serve",
                        f"{len(port.loss_making)} loss-making segment(s) of {len(port.segments)}",
                    ))
        else:
            coverage.append(PanelCoverage("cost_to_serve", SKIPPED, "no dataset=cost_to_serve rows"))
            metrics.append(ScorMetric("CO.1", "cost-to-serve net margin", None, NO_EVALUABLE, "cost_to_serve", "panel absent"))

        src_report = None
        if "sourcing" in payload.frames:
            path = _write_panel(payload.frames["sourcing"], tmp_path, "sourcing")
            def _src():
                prepared = sourcing_job.prepare(path, params)
                return sourcing_job.run(
                    prepared["scorecards"], prepared.get("prices"),
                    weights=params.get("weights"), notes=prepared.get("notes") or (),
                )
            src_report, cov = _try_panel("sourcing", _src)
            coverage.append(cov)
            if src_report is not None:
                issues = sourcing_job.verify(src_report)
                if issues:
                    coverage[-1] = PanelCoverage("sourcing", SKIPPED, "; ".join(issues))
                    src_report = None
                else:
                    otifs = [c.otif for c in src_report.scorecards if c.otif is not None]
                    mean_otif = (sum(otifs) / len(otifs)) if otifs else None
                    metrics.append(ScorMetric(
                        "RL.source_otif", "supplier OTIF (not Perfect Order RL.1.1)",
                        mean_otif,
                        OBSERVED if mean_otif is not None else NO_EVALUABLE,
                        "sourcing",
                        "supplier delivery reliability; customer Perfect Order is the orders panel",
                    ))
        else:
            coverage.append(PanelCoverage("sourcing", SKIPPED, "no dataset=sourcing rows"))

        otif_service = None
        if "orders" in payload.frames:
            def _otif():
                lines = _order_lines(payload.frames["orders"])
                by_sku: dict[str, list[dict]] = {}
                for line in lines:
                    by_sku.setdefault(line["product_id"], []).append(line)
                records = [
                    measure_from_order_lines(sku, sku_lines, as_of=as_of)
                    for sku, sku_lines in by_sku.items()
                ]
                return aggregate(records)
            otif_service, cov = _try_panel("orders", _otif)
            coverage.append(cov)
            if otif_service is not None:
                metrics.append(ScorMetric(
                    "RL.1.1", "Perfect Order / OTIF",
                    otif_service.otif,
                    OBSERVED if otif_service.otif is not None else NO_EVALUABLE,
                    "orders",
                    otif_service.method,
                ))
        else:
            coverage.append(PanelCoverage("orders", SKIPPED, "no dataset=orders rows"))
            if not any(m.code == "RL.1.1" for m in metrics):
                metrics.append(ScorMetric(
                    "RL.1.1", "Perfect Order / OTIF", None, NO_EVALUABLE, "orders",
                    "panel absent; supplier OTIF is not a substitute",
                ))

        proc_report = None
        if "procedures" in payload.frames:
            path = _write_panel(payload.frames["procedures"], tmp_path, "procedures")
            def _proc():
                prepared = procedure_audit_job.prepare(path, params)
                return procedure_audit_job.run(prepared, params)
            proc_report, cov = _try_panel("procedures", _proc)
            coverage.append(cov)
            if proc_report is not None:
                issues = procedure_audit_job.verify(proc_report)
                if issues:
                    coverage[-1] = PanelCoverage("procedures", SKIPPED, "; ".join(issues))
                    proc_report = None
        else:
            coverage.append(PanelCoverage("procedures", SKIPPED, "no dataset=procedures rows"))

        risk_report = None
        if "risk" in payload.frames:
            def _risk():
                records = risk_job.prepare_records(payload.frames["risk"], params)
                return risk_job.run(records)
            risk_report, cov = _try_panel("risk", _risk)
            coverage.append(cov)
            if risk_report is not None:
                issues = risk_job.verify(risk_report)
                if issues:
                    coverage[-1] = PanelCoverage("risk", SKIPPED, "; ".join(issues))
                    risk_report = None
                else:
                    gap = risk_report.resilience.gap_days
                    metrics.append(ScorMetric(
                        "AG.TTR_TTS", "TTR-TTS resilience gap (days)",
                        gap,
                        OBSERVED if gap is not None else NO_EVALUABLE,
                        "risk",
                        risk_report.resilience.note,
                    ))
        else:
            coverage.append(PanelCoverage("risk", SKIPPED, "no dataset=risk rows"))
            metrics.append(ScorMetric(
                "AG.TTR_TTS", "TTR-TTS resilience gap (days)", None, NO_EVALUABLE, "risk", "panel absent",
            ))

    ran = [c for c in coverage if c.status == RAN]
    if not ran:
        raise ValueError(
            "every SCOR/ISO panel was skipped. Need at least one recognised dataset "
            "with the columns that panel's existing job requires. Skips: "
            + "; ".join(f"{c.panel}: {c.reason}" for c in coverage)
        )

    # ISO 9001 / 28000 checklist -- observed / gap / no_evaluable. Never "compliant".
    rl = next((m for m in metrics if m.code == "RL.1.1"), None)
    if rl is None or rl.status == NO_EVALUABLE:
        checks.append(IsoControlCheck(
            "ISO 9001 8.5.1 / SCOR RL.1.1",
            "Control of production -- Perfect Order / OTIF",
            NO_EVALUABLE,
            "No dated order lines (ordered_qty, delivered_qty, promised_date, delivered_date).",
        ))
    else:
        checks.append(IsoControlCheck(
            "ISO 9001 8.5.1 / SCOR RL.1.1",
            "Control of production -- Perfect Order / OTIF",
            OBSERVED,
            f"OTIF={rl.value}; {rl.note}. A measured rate is not a conformity ruling.",
        ))

    if proc_report is None:
        checks.append(IsoControlCheck(
            "ISO 9001 7.5",
            "Documented information -- procedure register",
            NO_EVALUABLE,
            "No procedures panel. procedure_audit is the SOP-register tool for this clause.",
        ))
    elif proc_report.n_overdue > 0 or proc_report.n_without_owner > 0:
        checks.append(IsoControlCheck(
            "ISO 9001 7.5",
            "Documented information -- procedure register",
            GAP,
            f"{proc_report.n_overdue} overdue, {proc_report.n_without_owner} without owner. "
            "Gaps in the register, not a certification finding.",
        ))
    else:
        checks.append(IsoControlCheck(
            "ISO 9001 7.5",
            "Documented information -- procedure register",
            OBSERVED,
            "Register metadata reviewed (owners, versions, review dates). "
            "This does not verify that floor work matches the document.",
        ))

    if src_report is None:
        checks.append(IsoControlCheck(
            "ISO 9001 8.7",
            "Control of nonconforming outputs -- inbound quality",
            NO_EVALUABLE,
            "No sourcing panel (defect PPM). Kern's own QA gate is architectural, not client evidence.",
        ))
    else:
        ppms = [c.ppm for c in src_report.scorecards if c.ppm is not None]
        checks.append(IsoControlCheck(
            "ISO 9001 8.7",
            "Control of nonconforming outputs -- inbound quality",
            OBSERVED,
            f"Defect PPM observed on {len(ppms)} supplier(s). "
            "Kern's deliverable QA (8.7 on this engine) is separate from the client's process.",
        ))

    am = next((m for m in metrics if m.code == "AM.1.1"), None)
    if am is None or am.status == NO_EVALUABLE:
        checks.append(IsoControlCheck(
            "ISO 9001 9.1 / SCOR AM.1.1",
            "Monitoring -- cash-to-cash",
            NO_EVALUABLE,
            am.note if am else "financial panel absent",
        ))
    else:
        checks.append(IsoControlCheck(
            "ISO 9001 9.1 / SCOR AM.1.1",
            "Monitoring -- cash-to-cash",
            OBSERVED,
            f"cash-to-cash={am.value} days. {am.note}",
        ))

    if cts_report is None:
        checks.append(IsoControlCheck(
            "ISO 9001 9.1 / SCOR CO",
            "Monitoring -- cost-to-serve",
            NO_EVALUABLE,
            "cost_to_serve panel absent",
        ))
    else:
        checks.append(IsoControlCheck(
            "ISO 9001 9.1 / SCOR CO",
            "Monitoring -- cost-to-serve",
            OBSERVED,
            cts_report.summary,
        ))

    if risk_report is None:
        checks.append(IsoControlCheck(
            "ISO 28000 / SCOR AG",
            "Continuity -- TTR vs TTS",
            NO_EVALUABLE,
            "risk panel absent (needs likelihood, impact_value, time_to_recover, time_to_survive)",
        ))
    elif risk_report.resilience_verdict == "gap":
        checks.append(IsoControlCheck(
            "ISO 28000 / SCOR AG",
            "Continuity -- TTR vs TTS",
            GAP,
            risk_report.resilience.note,
        ))
    elif risk_report.resilience_verdict == "no_evaluable":
        checks.append(IsoControlCheck(
            "ISO 28000 / SCOR AG",
            "Continuity -- TTR vs TTS",
            NO_EVALUABLE,
            risk_report.resilience.note,
        ))
    else:
        checks.append(IsoControlCheck(
            "ISO 28000 / SCOR AG",
            "Continuity -- TTR vs TTS",
            OBSERVED,
            risk_report.resilience.note,
        ))

    n_observed = sum(1 for c in checks if c.status == OBSERVED)
    n_gap = sum(1 for c in checks if c.status == GAP)
    n_ne = sum(1 for c in checks if c.status == NO_EVALUABLE)

    used = {
        "as_of": as_of,
        "dso_days": dso_f,
        "dpo_days": dpo_f,
        "panels_ran": [c.panel for c in coverage if c.status == RAN],
        "n_observed": n_observed,
        "n_gap": n_gap,
        "n_no_evaluable": n_ne,
    }
    output_payload = {
        "metrics": [asdict(m) for m in metrics],
        "checks": [asdict(c) for c in checks],
        "coverage": [asdict(c) for c in coverage],
    }
    qa = (
        QaCheck("at_least_one_panel", ",".join(c.panel for c in ran), True),
        QaCheck("checks_cover_scorecard", f"n_checks={len(checks)}", len(checks) >= 1),
        QaCheck("prepared_by_is_engine", ENGINE_IDENTITY, True),
    )
    evidence = evidence_record(
        run_id=run_id,
        inputs=payload.artifacts,
        params_used={k: v for k, v in used.items() if v is not None and v != []},
        qa_attestation=qa,
        output_payload=output_payload,
        produced_at=produced_at,
        prepared_by=ENGINE_IDENTITY,
        formula_versions=(("audit_evidence", FORMULA_VERSION), ("scor_iso_audit", SCOR_ISO_VERSION)),
    )

    steps = [
        "Read the SCOR scorecard: observed numbers are measurements, not pass/fail vs a made-up target.",
        "Treat every no_evaluable row as a data gap to close, never as 'nobody objects' or 'compliant'.",
        "Hand the Lineage sheet (SHA-256, params, formula versions) to the registrar's auditor.",
        "Do not present this pack as an ISO certificate or a SCOR certification.",
    ]
    if n_gap:
        steps.insert(0, f"Close the {n_gap} named gap(s) on the ISO checklist before an external review.")
    packet = HandoffPacket(
        title="Prepare this evidence pack for the registrar's auditor (not a certificate)",
        steps=steps,
        risk_if_skipped=(
            "Without a human reviewer this pack can be misread as a conformity ruling. "
            "Kern does not certify ISO or SCOR."
        ),
    )
    residual = Residual(
        description=(
            "Checklist statuses are observed / gap / no_evaluable from the client's own data. "
            "They are not a certification decision."
        ),
        risk_if_skipped="A reader may quote a measured OTIF as a conformity ruling it is not.",
    )
    summary = (
        f"SCOR/ISO evidence pack: {len(ran)} panel(s) ran, {n_observed} control(s) observed, "
        f"{n_gap} gap(s), {n_ne} no_evaluable. Prepares evidence for a registrar's auditor; "
        "does not certify."
    )
    outcome = as_handoff(summary, [packet], residuals=[residual], confidence=0.75)

    return ScorIsoReport(
        metrics=tuple(metrics),
        checks=tuple(checks),
        coverage=tuple(coverage),
        evidence=evidence,
        outcome=outcome,
        summary=summary,
        n_observed=n_observed,
        n_gap=n_gap,
        n_no_evaluable=n_ne,
    )


def verify(report: ScorIsoReport) -> list[str]:
    """QA: never EXECUTED, never a conformity claim, evidence complete."""
    issues = list(verify_guided(report.outcome))
    if report.outcome.status == EXECUTED:
        issues.append("outcome must never be EXECUTED: scor_iso_audit prepares evidence, it does not certify")
    blobs = [report.summary, str(report.outcome.summary)]
    for residual in report.outcome.residuals:
        blobs.append(residual.description)
        blobs.append(residual.risk_if_skipped)
    for m in report.metrics:
        blobs.append(m.note)
        blobs.append(m.name)
    for c in report.checks:
        blobs.append(c.evidence)
        blobs.append(c.title)
        if c.status not in {OBSERVED, GAP, NO_EVALUABLE}:
            issues.append(f"illegal checklist status {c.status!r} on {c.clause}")
        if c.status in {"compliant", "conforming", "certified", "cumple"}:
            issues.append(f"conformity status on {c.clause}")
    for opt in report.outcome.options:
        blobs.extend([opt.label, opt.summary, opt.action])
    for packet in report.outcome.handoffs:
        blobs.append(packet.title)
        blobs.extend(packet.steps)
        blobs.append(packet.risk_if_skipped)
        if not packet.risk_if_skipped.strip():
            issues.append(f"handoff without risk_if_skipped: '{packet.title}'")
    for blob in blobs:
        if _iso_claim_forbidden(str(blob)):
            issues.append(f"forbidden conformity claim: {str(blob)[:80]}")
    if report.n_observed + report.n_gap + report.n_no_evaluable != len(report.checks):
        issues.append("check counts do not reconcile")
    if not any(c.status == RAN for c in report.coverage):
        issues.append("no panel ran")
    if report.evidence.prepared_by != ENGINE_IDENTITY:
        issues.append("prepared_by must be the engine identity, never a human name")
    if not report.evidence.inputs or not report.evidence.inputs[0].sha256:
        issues.append("evidence inputs missing sha256")
    return issues


def _write_xlsx(report: ScorIsoReport, path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    fill = PatternFill("solid", fgColor="1F2A44")
    font = Font(bold=True, color="FFFFFF")

    def _sheet(title: str, rows: list[dict], kv: list[tuple[str, object]] | None = None):
        ws = wb.active if title == "Scorecard" else wb.create_sheet(title)
        ws.title = title
        r = 1
        if kv:
            for key, value in kv:
                ws.cell(row=r, column=1, value=defuse_formula(key))
                ws.cell(row=r, column=2, value=defuse_formula(value))
                r += 1
            r += 1
        if rows:
            headers = list(rows[0].keys())
            for c, h in enumerate(headers, 1):
                cell = ws.cell(row=r, column=c, value=h)
                cell.fill = fill
                cell.font = font
            for ri, row in enumerate(rows, r + 1):
                for ci, h in enumerate(headers, 1):
                    val = row[h]
                    if isinstance(val, float) and not math.isfinite(val):
                        val = None
                    ws.cell(row=ri, column=ci, value=defuse_formula(val))
            last = r + len(rows) + 2
        else:
            ws.cell(row=r, column=1, value="(no rows)")
            last = r + 2
        ws.cell(row=last, column=1, value="Prepared by")
        ws.cell(row=last, column=2, value=defuse_formula(
            f"{report.evidence.prepared_by} on {report.evidence.produced_at}"
        ))
        ws.cell(row=last + 1, column=1, value="Reviewed by")
        ws.cell(row=last + 1, column=2, value="")
        ws.cell(row=last + 1, column=3, value="Date")
        ws.cell(row=last + 1, column=4, value="")
        ws.column_dimensions[get_column_letter(1)].width = 40
        ws.column_dimensions[get_column_letter(2)].width = 48

    _sheet("Scorecard", [
        {
            "code": m.code, "name": m.name, "value": m.value, "status": m.status,
            "source_panel": m.source_panel, "note": m.note,
        }
        for m in report.metrics
    ], kv=[
        ("aviso", "Kern prepares evidence for a registrar's auditor. It does not certify ISO or SCOR."),
        ("summary", report.summary),
    ])
    _sheet("ISO checklist", [
        {"clause": c.clause, "title": c.title, "status": c.status, "evidence": c.evidence}
        for c in report.checks
    ])
    _sheet("Coverage", [
        {"panel": c.panel, "status": c.status, "reason": c.reason, "summary": c.summary}
        for c in report.coverage
    ])
    evd = report.evidence
    _sheet("Lineage", [
        {"field": "run_id", "value": evd.run_id},
        {"field": "produced_at", "value": evd.produced_at},
        {"field": "prepared_by", "value": evd.prepared_by},
        {"field": "output_sha256", "value": evd.output_sha256},
        {"field": "formula_versions", "value": json.dumps(evd.formula_versions)},
        {"field": "params_used", "value": json.dumps(evd.params_used)},
        {"field": "qa_attestation", "value": json.dumps([asdict(c) for c in evd.qa_attestation])},
        {"field": "inputs", "value": json.dumps([asdict(a) for a in evd.inputs])},
    ])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def write_operational(report: ScorIsoReport, out_dir: str | Path, client: str = "Client") -> dict[str, Path]:
    del client
    d = Path(out_dir)
    d.mkdir(parents=True, exist_ok=True)
    evidence_path = d / "evidence.json"
    evidence_path.write_text(
        json.dumps(evidence_to_dict(report.evidence), indent=2, sort_keys=True),
        encoding="utf-8",
    )
    return {
        "scorecard_csv": write_summary_csv(
            [asdict(m) for m in report.metrics], d / "scor_scorecard.csv"
        ),
        "checklist_csv": write_summary_csv(
            [asdict(c) for c in report.checks], d / "iso_checklist.csv"
        ),
        "coverage_csv": write_summary_csv(
            [asdict(c) for c in report.coverage], d / "panel_coverage.csv"
        ),
        "evidence_json": evidence_path,
        "workbook": _write_xlsx(report, d / "scor_iso_workpaper.xlsx"),
    }


def build_deck(
    report: ScorIsoReport,
    *,
    client: str = "Client",
    prepared: str = "",
    citations: tuple[str, ...] = (),
    confidence: float | None = None,
) -> Deliverable:
    """Client-facing pack: scorecard + checklist. Always a human next step.

    ``confidence`` is derived when not supplied: a skipped panel or a
    ``no_evaluable`` control is a coverage gap, same class as a dead
    TOPSIS criterion (see nosilo_core.guided.confidence_from_coverage).
    """
    if confidence is None:
        gaps: list[str] = []
        if any(c.status != RAN for c in report.coverage):
            gaps.append("ignored_criterion")
        if report.n_no_evaluable:
            gaps.append("ignored_criterion")
        confidence = confidence_from_coverage(gaps, base=_DECK_BASE_CONFIDENCE)
    findings: list[Finding] = []
    for m in report.metrics:
        findings.append(Finding(
            f"{m.code} {m.name}: {m.status}",
            m.note,
            impact=pct(m.value) if m.value is not None and m.value <= 1.5 else money(m.value) if m.value is not None and m.value > 50 else (str(m.value) if m.value is not None else "no medido"),
        ))
    for c in report.checks:
        if c.status == GAP:
            findings.append(Finding(c.clause, c.evidence, impact="gap"))
    if not findings:
        findings.append(Finding(
            "No panel produced a metric",
            "Every dataset was skipped. Close the coverage table gaps and re-run.",
            impact="HANDOFF",
        ))

    kpis = (
        Kpi("panels ran", str(sum(1 for c in report.coverage if c.status == RAN)),
            rationale="Existing jobs that composed into this pack"),
        Kpi("controls observed", str(report.n_observed),
            rationale="ISO/SCOR rows with a measurement, not a certificate"),
        Kpi("gaps", str(report.n_gap), rationale="Named gaps the human closes before a registrar visit"),
        Kpi("no_evaluable", str(report.n_no_evaluable),
            rationale="Panels the data could not support -- never 'compliant by silence'"),
    )
    data_sources = tuple(
        DataSource(art.path_label, f"sha256={art.sha256[:16]}... n={art.n_rows}", "per run")
        for art in report.evidence.inputs
    )
    recommendations = [
        "Give this pack to the registrar's auditor with the Lineage sheet; do not call it a certificate.",
        "Close no_evaluable panels by supplying the missing dataset (financial, orders, procedures, risk).",
        "Procedure overdue rows are a document-control gap (ISO 9001 7.5), not a plant certificate.",
    ]
    return Deliverable(
        title="SCOR/ISO evidence pack (not a certificate)",
        client=client,
        summary=report.summary,
        findings=tuple(findings[:8]),
        kpis=kpis,
        data_sources=data_sources,
        recommendations=tuple(recommendations),
        citations=citations,
        confidence=confidence,
        residual=(
            "Reviewed by ______ Date ______. Kern prepared re-performable evidence. "
            "A registrar's auditor concludes. Kern does not certify ISO or SCOR."
        ),
        prepared=prepared or f"{report.evidence.prepared_by} on {report.evidence.produced_at}",
    )
